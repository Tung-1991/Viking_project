from __future__ import annotations

from datetime import datetime, timedelta
from zoneinfo import ZoneInfo

import pytest

from viking_v2.backtest.data import HistoricalDataStore
from viking_v2 import config as app_config
from viking_v2.backtest.engine import (
    BacktestEngine,
    _Position,
    _normal_trail_fill,
    _opening_fill_price,
    _profit_path_trade_metrics,
    _record_profit_path,
)
from viking_v2.backtest.models import BacktestConfig, BacktestScenario, BacktestSettings, BacktestTrade
from viking_v2.backtest.report import exit_detail, export_run_excel, round_headers, workbook_name
from viking_v2.backtest.replay import ReplayDataStore
from viking_v2.rules.business import StrategyDecision


VN_TZ = ZoneInfo("Asia/Ho_Chi_Minh")


def _payload(values: list[float], start: datetime) -> dict[str, list[float]]:
    timestamps = [int((start + timedelta(days=index)).timestamp()) for index in range(len(values))]
    return {
        "t": timestamps,
        "o": list(values),
        "h": [value * 1.01 for value in values],
        "l": [value * 0.99 for value in values],
        "c": list(values),
        "v": [1_000_000.0] * len(values),
    }


def _rules_without_buy_window(**overrides) -> dict:
    """Keep tests unrelated to the 14:00 gate independent of that live default."""

    return {"buy_window_enabled": False, **overrides}


def test_historical_store_caches_dnse_daily_data(tmp_path):
    calls: list[str] = []
    start = datetime(2022, 1, 1, tzinfo=VN_TZ)
    payload = _payload([10.0] * 700, start)

    def fetch(symbol, resolution, from_ts, to_ts):
        calls.append(symbol)
        return payload

    store = HistoricalDataStore(root=tmp_path, fetcher=fetch)
    first = store.load_daily("FPT", "2023-08-01", "2023-12-01", warmup_sessions=260)
    assert first
    assert calls
    calls.clear()
    second = store.load_daily("FPT", "2023-08-01", "2023-12-01", warmup_sessions=260)
    assert second == first
    assert calls == []


def test_historical_cache_reports_real_coverage_and_retries_missing_gap(tmp_path, monkeypatch):
    clock = [1_000.0]
    monkeypatch.setattr("viking_v2.backtest.data.time.time", lambda: clock[0])
    calls = []
    payload = _payload(
        [10.0] * 11,
        datetime(2026, 1, 10, tzinfo=VN_TZ),
    )

    def fetch(*args):
        calls.append(args)
        return payload

    store = HistoricalDataStore(root=tmp_path / "truthful-coverage", fetcher=fetch)
    store.load_bars("FPT", "2026-01-01", "2026-01-31", resolution="1D")
    cached = store._store("FPT", "1D").read()
    assert cached["coverage_start"] == "2026-01-10"
    assert cached["coverage_end"] == "2026-01-20"
    first_count = len(calls)

    store.load_bars("FPT", "2026-01-01", "2026-01-31", resolution="1D")
    assert len(calls) == first_count
    clock[0] += 301
    store.load_bars("FPT", "2026-01-01", "2026-01-31", resolution="1D")
    assert len(calls) > first_count


def test_backtest_reuses_static_rule_and_fills_next_session(tmp_path):
    start = datetime(2023, 1, 1, tzinfo=VN_TZ)
    values = (
        [10.0] * 260
        + [10.0 - 0.05 * index for index in range(1, 8)]
        + [9.65 + 0.12 * index for index in range(1, 12)]
        + [10.97 - 0.15 * index for index in range(1, 12)]
    )
    payload = _payload(values, start)
    store = HistoricalDataStore(root=tmp_path, fetcher=lambda *_args: payload)
    first_test_day = (start + timedelta(days=255)).date().isoformat()
    last_test_day = (start + timedelta(days=len(values) - 1)).date().isoformat()
    result = BacktestEngine(store).run(BacktestConfig(
        symbols=["FPT"],
        start_date=first_test_day,
        end_date=last_test_day,
        initial_capital=100_000_000,
        auto_market_phase=False,
        fixed_market_phase="UPTREND",
        fixed_exposure_pct=100,
        em_modes=["IND_EXIT"],
        loss_lock_enabled=True,
        loss_lock_hours=48,
        rule_parameters=_rules_without_buy_window(),
    ), save=False)
    sides = [event.side for event in result.events]
    assert sides == ["BUY", "SELL"]
    # Signal M occurs on index 268; execution must wait until index 269 open.
    assert result.events[0].date == (start + timedelta(days=269)).date().isoformat()
    assert result.closed_trades == 1
    assert result.total_fees > 0
    assert result.total_tax > 0
    assert result.config.rule_parameters["whipsaw_enabled"] is False
    assert result.data_quality["signal_resolution"] == "1D"
    assert result.data_quality["execution_resolution"]["FPT"] in {"5", "15", "1D"}


@pytest.mark.parametrize("slots, floor, ceiling", [(5, 150_000_000, 190_000_000),
                                                   (1, 800_000_000, 910_000_000)])
def test_capital_is_split_by_the_slot_count_exactly_as_entered(tmp_path, slots, floor, ceiling):
    """Exposure is divided by the entered slot count, never by the symbol count.

    One symbol with five slots deliberately deploys a fifth of the exposure and
    leaves the rest in cash, matching how the live bot sizes an order.
    """
    start = datetime(2023, 1, 1, tzinfo=VN_TZ)
    values = [10.0] * 260 + [9.8, 9.6, 9.7, 10.0, 10.3, 10.1, 9.8, 9.5]
    store = HistoricalDataStore(root=tmp_path / str(slots), fetcher=lambda *_a: _payload(values, start))
    result = BacktestEngine(store).run(BacktestConfig(
        symbols=["FPT"],
        start_date=(start + timedelta(days=255)).date().isoformat(),
        end_date=(start + timedelta(days=len(values) - 1)).date().isoformat(),
        initial_capital=1_000_000_000,
        fixed_market_phase="UPTREND", fixed_exposure_pct=90,
        rule_parameters=_rules_without_buy_window(max_positions=slots),
    ), save=False)
    buys = [event for event in result.events if event.side == "BUY"]
    assert buys, "kịch bản một mã phải vào được lệnh"
    assert floor <= buys[0].gross <= ceiling


def test_scenario_without_a_market_state_is_rejected():
    # A scenario is only meaningful once its state is chosen, so it is refused
    # outright rather than kept around as a disabled row nobody understands.
    with pytest.raises(ValueError):
        BacktestScenario("Sideway / Recovery", ["HSG"], "2022-11-17", "2025-04-25", "")


def test_loss_cooldown_is_calendar_hours():
    config = BacktestConfig(["FPT"], "2024-01-01", "2024-02-01", loss_lock_hours=48)
    assert config.loss_lock_hours == 48
    assert BacktestConfig(["FPT"], "2024-01-01", "2024-02-01").loss_lock_hours == 24


def test_backtest_defaults_are_neutral_and_independent():
    config = BacktestConfig(["FPT"], "2024-01-01", "2024-02-01")
    assert config.initial_capital == 1_000_000_000
    assert config.auto_market_phase is False
    assert config.fixed_market_phase == "ACCUMULATION"
    assert config.fixed_exposure_pct == 60
    assert config.loss_lock_enabled is False
    assert config.loss_lock_hours == 24
    assert config.whipsaw_enabled is False
    assert config.em_modes == ["NORMAL", "IND_EXIT"]
    assert config.execution_resolution == "AUTO"
    assert config.simulation_mode == "DAILY"
    assert config.buy_fee_rate == pytest.approx(app_config.DEFAULT_BUY_FEE_PCT / 100.0)
    assert BacktestSettings().buy_fee_pct == pytest.approx(app_config.DEFAULT_BUY_FEE_PCT)
    assert BacktestSettings().loss_lock_hours == 24
    assert BacktestSettings().simulation_mode == "AUTO_HYBRID"


def test_backtest_can_lock_replay_to_two_minute_source():
    config = BacktestConfig(
        ["VIX"], "2026-03-02", "2026-09-03",
        simulation_mode="REPLAY", execution_resolution="2",
    )
    assert config.execution_resolution == "2"


def test_backtest_settings_migrate_legacy_ema_pair_to_both_directions():
    settings = BacktestSettings(rule_parameters={"ema_fast": 5, "ema_slow": 9})

    assert settings.rule_parameters["buy_ema_fast"] == 5
    assert settings.rule_parameters["buy_ema_slow"] == 9
    assert settings.rule_parameters["sell_ema_fast"] == 5
    assert settings.rule_parameters["sell_ema_slow"] == 9


def test_multi_scenario_workbook_name_keeps_single_symbol():
    name = workbook_name(
        "MODE 2 - E ONLY", symbols=["CTS", "CTS"], start="2026-03-14",
        end="2026-08-20", runs=2, stamp="AUDIT",
    )
    assert "CTS" in name and "2 kich ban" in name


def test_continuous_fill_uses_first_candle_open_at_or_after_0915():
    day = datetime(2026, 8, 20, tzinfo=VN_TZ)
    bars = [
        {"time": int(day.replace(hour=9, minute=0).timestamp()), "open": 100.0, "close": 101.0},
        {"time": int(day.replace(hour=9, minute=15).timestamp()), "open": 102.0, "close": 109.0},
        {"time": int(day.replace(hour=9, minute=30).timestamp()), "open": 110.0, "close": 111.0},
    ]
    assert _opening_fill_price(bars, "ATO") == 100.0
    # 102 is observable at 09:15; 109 is only known after that candle closes.
    assert _opening_fill_price(bars, "CONTINUOUS") == 102.0


def test_mode2_carry_keeps_pending_order_and_does_not_replay_boundary_day(tmp_path):
    start = datetime(2023, 1, 1, tzinfo=VN_TZ)
    values = [10.0] * 260 + [10.0 - 0.05 * index for index in range(1, 8)]
    values += [9.65 + 0.12 * index for index in range(1, 12)]
    store = HistoricalDataStore(root=tmp_path, fetcher=lambda *_args: _payload(values, start))
    carry: dict = {}
    first = BacktestEngine(store).run(BacktestConfig(
        symbols=["FPT"],
        start_date=(start + timedelta(days=255)).date().isoformat(),
        end_date=(start + timedelta(days=268)).date().isoformat(),
        initial_capital=100_000_000,
        fixed_market_phase="UPTREND", fixed_exposure_pct=100,
        em_modes=["IND_EXIT"],
        rule_parameters=_rules_without_buy_window(),
    ), carry=carry, save=False)
    assert first.events == []
    assert "FPT" in carry["pending"]

    second = BacktestEngine(store).run(BacktestConfig(
        symbols=["FPT"],
        # Deliberately overlap the boundary; the account must not replay it.
        start_date=(start + timedelta(days=268)).date().isoformat(),
        end_date=(start + timedelta(days=270)).date().isoformat(),
        initial_capital=100_000_000,
        fixed_market_phase="UPTREND", fixed_exposure_pct=100,
        em_modes=["IND_EXIT"],
        rule_parameters=_rules_without_buy_window(),
    ), carry=carry, save=False)
    buys = [event for event in second.events if event.side == "BUY"]
    assert len(buys) == 1
    assert buys[0].date == (start + timedelta(days=269)).date().isoformat()
    assert second.equity_curve[0]["date"] == (start + timedelta(days=269)).date().isoformat()


def test_execution_data_auto_falls_back_and_keeps_resolution_caches_separate(tmp_path):
    start = datetime(2023, 1, 1, tzinfo=VN_TZ)
    payload = _payload([10.0] * 20, start)
    calls: list[str] = []

    def fetch(_symbol, resolution, _from_ts, _to_ts):
        calls.append(resolution)
        return payload if resolution == "15" else None

    store = HistoricalDataStore(root=tmp_path, fetcher=fetch)
    rows, resolution, warnings = store.load_execution("FPT", "2023-01-01", "2023-01-20")
    assert rows
    assert resolution == "15"
    assert warnings
    assert (tmp_path / "cache" / "FPT_15.json").exists()
    assert (tmp_path / "cache" / "FPT_5.json").exists()
    calls.clear()
    second_rows, second_resolution, _ = store.load_execution("FPT", "2023-01-01", "2023-01-20")
    assert second_rows == rows
    assert second_resolution == "15"
    assert calls == []


def test_scenario_takes_indicators_and_exposure_from_shared_settings(tmp_path):
    """A scenario pins symbol, window and state only.

    Letting a scenario carry its own EMA would test a regime-switching rule
    the live bot cannot execute.
    """
    start = datetime(2023, 1, 1, tzinfo=VN_TZ)
    values = [100 + index * 0.1 for index in range(300)]
    store = HistoricalDataStore(root=tmp_path, fetcher=lambda *_args: _payload(values, start))
    scenario = BacktestScenario(
        "HSG · UPTREND", ["HSG"],
        (start + timedelta(days=255)).date().isoformat(),
        (start + timedelta(days=299)).date().isoformat(),
        "UPTREND",
    )
    assert not hasattr(scenario, "buy_ema_fast")
    result = BacktestEngine(store).run_scenario(
        scenario, initial_capital=1_000_000_000,
        rule_parameters=_rules_without_buy_window(
            sell_ema_fast=5, sell_ema_slow=10,
            exposure={"UPTREND": 0.9},
        ),
        save=False,
    )
    assert result.config.rule_parameters["sell_ema_fast"] == 5
    assert result.config.fixed_exposure_pct == 90


def test_excel_export_gives_each_run_its_own_sheet(tmp_path):
    start = datetime(2023, 1, 1, tzinfo=VN_TZ)
    values = [10.0] * 260 + [9.8, 9.6, 9.7, 10.0, 10.3, 10.1, 9.8, 9.5]
    store = HistoricalDataStore(root=tmp_path / "data", fetcher=lambda *_a: _payload(values, start))
    result = BacktestEngine(store).run(BacktestConfig(
        symbols=["FPT"],
        start_date=(start + timedelta(days=255)).date().isoformat(),
        end_date=(start + timedelta(days=len(values) - 1)).date().isoformat(),
        fixed_market_phase="ACCUMULATION", em_modes=["IND_EXIT"], run_name="MODE 1",
        rule_parameters=_rules_without_buy_window(),
    ), save=False)
    path = export_run_excel(result, tmp_path / "exports", mode="MODE 1", stamp="0819")
    # The name has to say what is inside instead of a random run id.
    assert "MODE 1" in path.name and "FPT" in path.name and "0819" in path.name
    from openpyxl import load_workbook
    book = load_workbook(path, read_only=True)
    assert book.sheetnames[0] == "MODE 1"
    assert book.sheetnames[-1] == "THÔNG TIN"
    assert {"TÍN HIỆU", "KHỚP LỆNH", "MFE T+2", "PROTECT METRICS", "PROFIT PATH"} <= set(book.sheetnames)
    headers = [cell.value for cell in next(book["MODE 1"].iter_rows(min_row=1, max_row=1))]
    # Both modes keep the same shape so their sheets can be compared or pasted
    # together; no blended average exit price anywhere.
    # The run name is the sheet tab now, so that column is gone from the rows.
    assert headers == list(round_headers(result)[1:])
    assert headers[0] == "LƯỢT"
    assert "GIÁ RA" not in headers
    assert {"CẮT LỖ", "PHIÊN", "THOÁT BỞI", "EMA3 / EMA6 RA", "RSI RA"} <= set(headers)
    info = list(book["THÔNG TIN"].values)
    assert any(row[0] == "MFE và T+2" for row in info)


def test_exit_detail_names_every_sell_with_its_own_price():
    trade = BacktestTrade("t", "HSG", "2022-02-09", exit_fills=[
        {"event": "NORMAL_PROTECTION", "quantity": 1500, "price": 21.24},
        {"event": "INDICATOR_EXIT", "quantity": 3000, "price": 21.09},
    ])
    assert exit_detail(trade) == "PROTECT 1,500@21.24 + E 3,000@21.09"
    assert exit_detail(BacktestTrade("t", "HSG", "2022-02-09")) == "CÒN MỞ"


def test_normal_intraday_trail_does_not_use_same_bar_future_low():
    bars = [
        {"open": 106.0, "high": 108.0, "low": 104.0, "close": 105.0},
        {"open": 105.0, "high": 105.5, "low": 104.0, "close": 104.5},
    ]
    fill, peak = _normal_trail_fill(
        bars, entry_price=100.0, peak_profit_pct=0.0,
        arm_pct=7.0, giveback_pct=3.0,
    )
    assert peak == pytest.approx(8.0)
    # Giveback is 3% of the peak price: 100 x 1.08 x 0.97, not 100 x 1.05.
    assert fill == pytest.approx(104.76)


def test_profit_path_splits_mfe_at_t2_release_and_accepts_a_later_equal_high():
    position = _Position(
        trade_id="t", symbol="FPT", quantity=100, entry_quantity=100,
        avg_price=100.0, opened_date="2026-01-02", settle_date="2026-01-06",
        buy_fee=0.0, capital_principal=10_000_000.0, em_modes=["NORMAL"],
        opened_at="2026-01-02T14:00:00+07:00",
    )
    for observed_at, high in (
        ("2026-01-06T12:59:00+07:00", 110.0),
        ("2026-01-06T13:00:00+07:00", 105.0),
    ):
        _record_profit_path(
            position, observed_at,
            open_price=high, high_price=high, low_price=high, close_price=high,
            source_resolution="1",
        )
    metrics = _profit_path_trade_metrics(
        position, "2026-01-06T13:00:00+07:00", 5.0,
    )
    assert metrics["mfe_before_settlement_pct"] == pytest.approx(10.0)
    assert metrics["mfe_after_settlement_pct"] == pytest.approx(5.0)
    assert metrics["mfe_peak_phase"] == "TRUOC_T2"
    assert metrics["settlement_release_at"] == "2026-01-06T13:00:00+07:00"

    _record_profit_path(
        position, "2026-01-06T13:01:00+07:00",
        open_price=110.0, high_price=110.0, low_price=110.0, close_price=110.0,
        source_resolution="1",
    )
    metrics = _profit_path_trade_metrics(
        position, "2026-01-06T13:01:00+07:00", 10.0,
    )
    assert metrics["mfe_after_settlement_pct"] == pytest.approx(10.0)
    assert metrics["mfe_peak_phase"] == "SAU_T2"


def test_scenario_accepts_runtime_callbacks_without_putting_them_in_config(tmp_path):
    start = datetime(2023, 1, 1, tzinfo=VN_TZ)
    values = [100 + index * 0.1 for index in range(300)]
    store = HistoricalDataStore(root=tmp_path, fetcher=lambda *_args: _payload(values, start))
    scenario = BacktestScenario(
        "VA", ["HSG"],
        (start + timedelta(days=255)).date().isoformat(),
        (start + timedelta(days=299)).date().isoformat(),
        "ACCUMULATION",
    )
    progress: list[float] = []
    result = BacktestEngine(store).run_scenario(
        scenario,
        initial_capital=1_000_000_000,
        rule_parameters=_rules_without_buy_window(),
        progress=lambda value, _text: progress.append(value),
        cancelled=lambda: False,
        save=False,
    )
    assert result.config.run_name == "VA"
    assert progress


def test_mode2_replay_catches_intraday_daily_ema_cross_and_fills_next_bar(tmp_path):
    start = datetime(2025, 12, 1, tzinfo=VN_TZ)
    values = [8.0] * 221
    test_day = (start + timedelta(days=220)).date().isoformat()
    payload = _payload(values, start)

    def fetch(_symbol, resolution, _from, _to):
        return payload if resolution == "1D" else None

    store = HistoricalDataStore(root=tmp_path / "data", fetcher=fetch)
    replay = ReplayDataStore(store.root / "replay")
    source = tmp_path / "HOSE_DLY_FPT, 1.csv"
    source.write_text(
        "time,open,high,low,close,Volume\n"
        f"{test_day}T02:15:00Z,8,9,8,9,100\n"
        f"{test_day}T02:16:00Z,9,9,9,9,100\n"
        f"{test_day}T07:45:00Z,8,8,8,8,100\n",
        encoding="utf-8",
    )
    replay.import_file(source, price_scale=1)
    config = BacktestConfig(
        ["FPT"], test_day, test_day, initial_capital=100_000_000,
        fixed_market_phase="UPTREND", fixed_exposure_pct=100,
        fill_session="CONTINUOUS", simulation_mode="REPLAY",
        rule_parameters=_rules_without_buy_window(
            buy_ema_fast=2, buy_ema_slow=3, rsi_period=2,
            max_positions=1, no_compound_enabled=False,
        ),
    )
    result = BacktestEngine(store, replay).run(config, save=False)
    buys = [event for event in result.events if event.side == "BUY"]
    assert len(buys) == 1
    assert datetime.fromisoformat(buys[0].signal_time).strftime("%H:%M") == "09:15"
    assert datetime.fromisoformat(buys[0].fill_time).strftime("%H:%M") == "09:16"
    assert buys[0].source_resolution == "1"
    assert result.data_quality["signal_resolution"] == "1D_REALTIME"
    report = export_run_excel(result, tmp_path / "exports", mode="MODE 2", stamp="replay")
    from openpyxl import load_workbook
    workbook = load_workbook(report, read_only=True)
    assert {"TÍN HIỆU", "KHỚP LỆNH", "THÔNG TIN"} <= set(workbook.sheetnames)
    fill_rows = list(workbook["KHỚP LỆNH"].iter_rows(values_only=True))
    assert fill_rows[1][1].endswith("+07:00")
    assert fill_rows[1][2].endswith("+07:00")
    workbook.close()

    daily = BacktestEngine(store, replay).run(BacktestConfig(
        ["FPT"], test_day, test_day, initial_capital=100_000_000,
        fixed_market_phase="UPTREND", fixed_exposure_pct=100,
        simulation_mode="DAILY",
        rule_parameters=_rules_without_buy_window(
            buy_ema_fast=2, buy_ema_slow=3, rsi_period=2,
            max_positions=1, no_compound_enabled=False,
        ),
    ), save=False)
    assert [event for event in daily.events if event.side == "BUY"] == []


def test_replay_confirms_buy_for_five_exchange_minutes_then_fills_next_bar(tmp_path):
    start = datetime(2025, 12, 1, tzinfo=VN_TZ)
    values = [8.0] * 221
    test_day = (start + timedelta(days=220)).date().isoformat()
    payload = _payload(values, start)
    store = HistoricalDataStore(
        root=tmp_path / "confirm-replay",
        fetcher=lambda _symbol, resolution, _from, _to: payload if resolution == "1D" else None,
    )
    replay = ReplayDataStore(store.root / "replay")
    source = tmp_path / "HOSE_DLY_FPT, 1.csv"
    intraday = [
        f"{test_day}T02:{minute:02d}:00Z,9,9,9,9,100"
        for minute in range(15, 22)
    ]
    source.write_text(
        "time,open,high,low,close,Volume\n" + "\n".join([
            *intraday,
            f"{test_day}T07:45:00Z,9,9,9,9,100",
        ]) + "\n",
        encoding="utf-8",
    )
    replay.import_file(source, price_scale=1)
    result = BacktestEngine(store, replay).run(BacktestConfig(
        ["FPT"], test_day, test_day, initial_capital=100_000_000,
        fixed_market_phase="UPTREND", fixed_exposure_pct=100,
        fill_session="CONTINUOUS", simulation_mode="REPLAY",
        rule_parameters=_rules_without_buy_window(
            buy_ema_fast=2, buy_ema_slow=3, rsi_period=2,
            max_positions=1, no_compound_enabled=False,
            buy_confirmation_enabled=True, buy_confirmation_minutes=5,
            buy_confirmation_require_ema=True, buy_confirmation_require_rsi=True,
        ),
    ), save=False)
    buy = next(event for event in result.events if event.side == "BUY")
    assert datetime.fromisoformat(buy.signal_time).strftime("%H:%M") == "09:15"
    assert datetime.fromisoformat(buy.decision_time).strftime("%H:%M") == "09:20"
    assert datetime.fromisoformat(buy.fill_time).strftime("%H:%M") == "09:21"
    assert any(row["reason"] == "BUY_CONFIRMATION_WAIT" for row in result.signals)


@pytest.mark.parametrize("window_on,confirmation_on,expected_fill", [
    (False, False, "09:16"), (True, False, "14:02"), (True, True, "14:08"),
])
def test_replay_buy_window_with_two_minute_data_and_confirmation(tmp_path, window_on, confirmation_on, expected_fill):
    start = datetime(2025, 12, 1, tzinfo=VN_TZ)
    test_day = (start + timedelta(days=220)).date().isoformat()
    payload = _payload([8.0] * 221, start)
    store = HistoricalDataStore(root=tmp_path / "data", fetcher=lambda _s, res, *_: payload if res == "1D" else None)
    replay = ReplayDataStore(tmp_path / "replay")
    source = tmp_path / "HOSE_DLY_FPT, 2.csv"
    times = ["02:14", "02:16", "06:58", "07:00", "07:02", "07:04", "07:06", "07:08", "07:30", "07:44"]
    source.write_text("time,open,high,low,close,Volume\n" + "\n".join(
        f"{test_day}T{stamp}:00Z,9,9,9,9,100" for stamp in times
    ) + "\n", encoding="utf-8")
    replay.import_file(source, price_scale=1)
    result = BacktestEngine(store, replay).run(BacktestConfig(
        ["FPT"], test_day, test_day, initial_capital=100_000_000,
        fixed_market_phase="UPTREND", fixed_exposure_pct=100,
        fill_session="CONTINUOUS", simulation_mode="REPLAY",
        rule_parameters={"buy_ema_fast": 2, "buy_ema_slow": 3, "rsi_period": 2,
                         "max_positions": 1, "buy_window_enabled": window_on,
                         "buy_confirmation_enabled": confirmation_on},
    ), save=False)
    buys = [event for event in result.events if event.side == "BUY"]
    assert len(buys) == 1
    assert datetime.fromisoformat(buys[0].fill_time).strftime("%H:%M") == expected_fill
    assert datetime.fromisoformat(buys[0].signal_time).strftime("%H:%M") == "09:14"
    if window_on:
        assert any(row["reason"] == "BUY_WINDOW_WAIT" for row in result.signals)
        from openpyxl import load_workbook
        report = export_run_excel(result, tmp_path / "exports", mode="MODE 2", stamp="window")
        book = load_workbook(report, read_only=True)
        info = list(book["THÔNG TIN"].values)
        assert any(row[0] == "Khung giờ mua" and "14:00" in str(row[1]) for row in info)
        assert "KHUNG GIỜ MUA" in next(book["TÍN HIỆU"].values)
        book.close()


def test_replay_uses_previous_source_bar_not_previous_daily_close_for_ema(tmp_path):
    start = datetime(2025, 12, 1, tzinfo=VN_TZ)
    values = [100.0] * 220 + [101.0, 100.0, 100.5, 100.6]
    test_day = (start + timedelta(days=len(values) - 1)).date().isoformat()
    payload = _payload(values, start)
    store = HistoricalDataStore(
        root=tmp_path / "ema-stream",
        fetcher=lambda _symbol, resolution, _from, _to: payload if resolution == "1D" else None,
    )
    replay = ReplayDataStore(store.root / "replay")
    source = tmp_path / "HOSE_DLY_FPT, 1.csv"
    source.write_text(
        "time,open,high,low,close,Volume\n"
        f"{test_day}T02:15:00Z,95,95,95,95,100\n"
        f"{test_day}T02:16:00Z,100.6,100.6,100.6,100.6,100\n"
        f"{test_day}T02:17:00Z,100.6,100.6,100.6,100.6,100\n"
        f"{test_day}T07:45:00Z,100.6,100.6,100.6,100.6,100\n",
        encoding="utf-8",
    )
    replay.import_file(source, price_scale=1)
    result = BacktestEngine(store, replay).run(BacktestConfig(
        ["FPT"], test_day, test_day,
        initial_capital=100_000_000, fixed_market_phase="UPTREND",
        fixed_exposure_pct=100, fill_session="CONTINUOUS",
        simulation_mode="REPLAY", whipsaw_enabled=False,
        rule_parameters=_rules_without_buy_window(
            buy_ema_fast=3, buy_ema_slow=6, rsi_period=14,
            max_positions=1, no_compound_enabled=False,
        ),
    ), save=False)
    buy = next(event for event in result.events if event.side == "BUY")
    assert datetime.fromisoformat(buy.signal_time).strftime("%H:%M") == "09:16"
    assert datetime.fromisoformat(buy.fill_time).strftime("%H:%M") == "09:17"


def test_replay_strict_rejects_missing_day_but_hybrid_records_fallback(tmp_path):
    start = datetime(2025, 12, 1, tzinfo=VN_TZ)
    values = [8.0] * 221
    test_day = (start + timedelta(days=220)).date().isoformat()
    payload = _payload(values, start)
    store = HistoricalDataStore(
        root=tmp_path / "data",
        fetcher=lambda _symbol, resolution, _from, _to: payload if resolution == "1D" else None,
    )
    common = dict(
        symbols=["FPT"], start_date=test_day, end_date=test_day,
        fixed_market_phase="UPTREND", initial_capital=100_000_000,
        rule_parameters=_rules_without_buy_window(),
    )
    with pytest.raises(RuntimeError, match="thiếu dữ liệu FULL"):
        BacktestEngine(store).run(BacktestConfig(**common, simulation_mode="REPLAY"), save=False)
    result = BacktestEngine(store).run(
        BacktestConfig(**common, simulation_mode="AUTO_HYBRID"), save=False,
    )
    assert result.data_quality["fallback_count"] == 1
    assert result.data_quality["source_coverage"]["FPT"][test_day]["data_quality"] == "FALLBACK_1D"


@pytest.mark.parametrize("fill_session, expected", [("ATO", "09:15"), ("CONTINUOUS", "09:16")])
def test_replay_atc_signal_carries_to_next_eligible_session(tmp_path, fill_session, expected):
    start = datetime(2025, 12, 1, tzinfo=VN_TZ)
    values = [8.0] * 222
    signal_day = (start + timedelta(days=220)).date()
    fill_day = (start + timedelta(days=221)).date()
    payload = _payload(values, start)
    store = HistoricalDataStore(
        root=tmp_path / fill_session,
        fetcher=lambda _symbol, resolution, _from, _to: payload if resolution == "1D" else None,
    )
    replay = ReplayDataStore(store.root / "replay")
    source = tmp_path / f"HOSE_DLY_FPT, 1-{fill_session}.csv"
    # Override filename inference because the test suffix follows the resolution.
    source.write_text(
        "time,open,high,low,close,Volume\n"
        f"{signal_day}T02:15:00Z,8,8,8,8,100\n"
        f"{signal_day}T07:45:00Z,8,9,8,9,100\n"
        f"{fill_day}T02:15:00Z,9,9,9,9,100\n"
        f"{fill_day}T02:16:00Z,9,9,9,9,100\n"
        f"{fill_day}T07:45:00Z,9,9,9,9,100\n",
        encoding="utf-8",
    )
    replay.import_file(source, symbol="FPT", resolution="1", price_scale=1)
    result = BacktestEngine(store, replay).run(BacktestConfig(
        ["FPT"], signal_day.isoformat(), fill_day.isoformat(),
        initial_capital=100_000_000, fixed_market_phase="UPTREND",
        fixed_exposure_pct=100, fill_session=fill_session, simulation_mode="REPLAY",
        rule_parameters=_rules_without_buy_window(
            buy_ema_fast=2, buy_ema_slow=3, rsi_period=2,
            max_positions=1, no_compound_enabled=False,
        ),
    ), save=False)
    buy = next(event for event in result.events if event.side == "BUY")
    assert datetime.fromisoformat(buy.signal_time).strftime("%Y-%m-%d %H:%M") == f"{signal_day} 14:45"
    assert datetime.fromisoformat(buy.fill_time).strftime("%Y-%m-%d %H:%M") == f"{fill_day} {expected}"


def test_replay_rechecks_waiting_t2_sell_before_fill(tmp_path, monkeypatch):
    start = datetime(2025, 12, 1, tzinfo=VN_TZ)
    values = [10.0] * 223
    days = [(start + timedelta(days=index)).date() for index in range(220, 223)]
    payload = _payload(values, start)
    store = HistoricalDataStore(
        root=tmp_path / "data",
        fetcher=lambda _symbol, resolution, _from, _to: payload if resolution == "1D" else None,
    )
    replay = ReplayDataStore(store.root / "replay")
    source = tmp_path / "HOSE_DLY_FPT, 1.csv"
    rows = []
    for day in days:
        for utc_time in ("02:15", "02:16", "07:45"):
            rows.append(f"{day}T{utc_time}:00Z,10,10,10,10,100")
    source.write_text(
        "time,open,high,low,close,Volume\n" + "\n".join(rows) + "\n",
        encoding="utf-8",
    )
    replay.import_file(source, price_scale=1)

    class Rule:
        def __init__(self, _params):
            pass

        def evaluate(self, context, portfolio):
            stamp = datetime.fromtimestamp(context["bars"][-1]["time"], VN_TZ)
            quantity = portfolio.get("position_quantity", 0)
            if not quantity and stamp.date() == days[0] and stamp.strftime("%H:%M") == "09:15":
                return StrategyDecision("BUY", "FPT", "BUY_SIGNAL", event="ENTRY_BUY", signal="BUY")
            if quantity and stamp.date() == days[0]:
                return StrategyDecision("SELL", "FPT", "SELL_SIGNAL", event="INDICATOR_EXIT", signal="SELL")
            return StrategyDecision("WAIT", "FPT", "HOLD_POSITION", signal="")

    monkeypatch.setattr("viking_v2.backtest.engine.StaticRule", Rule)
    result = BacktestEngine(store, replay).run(BacktestConfig(
        ["FPT"], days[0].isoformat(), days[-1].isoformat(),
        initial_capital=100_000_000, fixed_market_phase="UPTREND",
        fixed_exposure_pct=100, fill_session="CONTINUOUS",
        simulation_mode="REPLAY", sell_wait_policy="RECHECK", em_modes=[],
        rule_parameters=_rules_without_buy_window(
            max_positions=1, no_compound_enabled=False,
        ),
    ), save=False)
    assert [event.side for event in result.events] == ["BUY"]
    assert result.trades[0].outcome == "OPEN"


def test_replay_waiting_t2_sell_fills_no_earlier_than_1300(tmp_path, monkeypatch):
    start = datetime(2025, 12, 1, tzinfo=VN_TZ)
    values = [10.0] * 223
    days = [(start + timedelta(days=index)).date() for index in range(220, 223)]
    payload = _payload(values, start)
    store = HistoricalDataStore(
        root=tmp_path / "data-t2-time",
        fetcher=lambda _symbol, resolution, _from, _to: payload if resolution == "1D" else None,
    )
    replay = ReplayDataStore(store.root / "replay")
    source = tmp_path / "HOSE_DLY_FPT, 1-t2-time.csv"
    rows = []
    for day in days:
        # UTC times map to 09:15, 09:16, 11:29, 13:00 and 14:45 Vietnam time.
        # 12:59 is lunch and correctly rejected by the replay importer; the
        # PAPER test above covers the exact 12:59:59 boundary.
        for utc_time in ("02:15", "02:16", "04:29", "06:00", "07:45"):
            rows.append(f"{day}T{utc_time}:00Z,10,10,10,10,100")
    source.write_text(
        "time,open,high,low,close,Volume\n" + "\n".join(rows) + "\n",
        encoding="utf-8",
    )
    replay.import_file(source, symbol="FPT", resolution="1", price_scale=1)

    class Rule:
        def __init__(self, _params):
            pass

        def evaluate(self, context, portfolio):
            stamp = datetime.fromtimestamp(context["bars"][-1]["time"], VN_TZ)
            quantity = portfolio.get("position_quantity", 0)
            if not quantity and stamp.date() == days[0] and stamp.strftime("%H:%M") == "09:15":
                return StrategyDecision("BUY", "FPT", "BUY_SIGNAL", event="ENTRY_BUY", signal="BUY")
            if quantity and stamp.date() == days[0] and stamp.strftime("%H:%M") == "14:45":
                return StrategyDecision("SELL", "FPT", "SELL_SIGNAL", event="INDICATOR_EXIT", signal="SELL")
            return StrategyDecision("WAIT", "FPT", "HOLD_POSITION", signal="")

    monkeypatch.setattr("viking_v2.backtest.engine.StaticRule", Rule)
    result = BacktestEngine(store, replay).run(BacktestConfig(
        ["FPT"], days[0].isoformat(), days[-1].isoformat(),
        initial_capital=100_000_000, fixed_market_phase="UPTREND",
        fixed_exposure_pct=100, fill_session="CONTINUOUS",
        simulation_mode="REPLAY", sell_wait_policy="KEEP", em_modes=[],
        rule_parameters=_rules_without_buy_window(
            max_positions=1, no_compound_enabled=False,
        ),
    ), save=False)
    sell = next(event for event in result.events if event.side == "SELL")
    assert datetime.fromisoformat(sell.fill_time).strftime("%Y-%m-%d %H:%M") == f"{days[2]} 13:00"


def test_daily_never_uses_morning_open_on_t2_for_waiting_sell(tmp_path, monkeypatch):
    start = datetime(2025, 12, 1, tzinfo=VN_TZ)
    values = [10.0] * 225
    days = [(start + timedelta(days=index)).date() for index in range(220, 225)]
    store = HistoricalDataStore(
        root=tmp_path / "daily-t2-time",
        fetcher=lambda _symbol, resolution, _from, _to: (
            _payload(values, start) if resolution == "1D" else None
        ),
    )

    class Rule:
        def __init__(self, _params):
            pass

        def evaluate(self, context, portfolio):
            day = datetime.fromtimestamp(context["bars"][-1]["time"], VN_TZ).date()
            quantity = portfolio.get("position_quantity", 0)
            if not quantity and day == days[0]:
                return StrategyDecision("BUY", "FPT", "BUY_SIGNAL", event="ENTRY_BUY", signal="BUY")
            if quantity:
                return StrategyDecision("SELL", "FPT", "SELL_SIGNAL", event="INDICATOR_EXIT", signal="SELL")
            return StrategyDecision("WAIT", "FPT", "HOLD_POSITION", signal="")

    monkeypatch.setattr("viking_v2.backtest.engine.StaticRule", Rule)
    result = BacktestEngine(store).run(BacktestConfig(
        ["FPT"], days[0].isoformat(), days[-1].isoformat(),
        initial_capital=100_000_000, fixed_market_phase="UPTREND",
        fixed_exposure_pct=100, simulation_mode="DAILY", em_modes=[],
        rule_parameters=_rules_without_buy_window(
            max_positions=1, no_compound_enabled=False,
        ),
    ), save=False)
    sell = next(event for event in result.events if event.side == "SELL")
    # BUY fills on days[1], so days[3] is T+2. DAILY has no afternoon price
    # and must not pretend its morning Open is sellable; it waits to days[4].
    assert sell.date == days[4].isoformat()
