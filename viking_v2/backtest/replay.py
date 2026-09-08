from __future__ import annotations

from dataclasses import asdict, dataclass
from contextlib import closing
from datetime import date, datetime
from hashlib import sha256
from pathlib import Path
import csv
import re
import sqlite3
from typing import Any, Iterable
from zoneinfo import ZoneInfo

from ..trading.market import (
    exchange_close_minute,
    exchange_open_minute,
    normalize_exchange,
)


VN_TZ = ZoneInfo("Asia/Ho_Chi_Minh")
REQUIRED_COLUMNS = ("time", "open", "high", "low", "close", "volume")


@dataclass(slots=True)
class ReplayDataset:
    symbol: str
    resolution: str
    exchange: str
    bar_count: int
    coverage_start: str
    coverage_end: str
    full_days: int
    partial_days: int
    source_name: str = ""
    source_hash: str = ""
    timezone: str = "Asia/Ho_Chi_Minh"
    price_scale: float = 1.0
    imported_at: str = ""

    def to_dict(self) -> dict[str, Any]:
        return asdict(self)


@dataclass(slots=True)
class ReplayImportPreview:
    path: str
    symbol: str
    resolution: str
    exchange: str
    timezone: str
    price_scale: float
    bar_count: int
    coverage_start: str
    coverage_end: str
    full_days: int
    partial_days: int
    partial_dates: list[str]
    source_hash: str
    inserted: int = 0
    replaced: int = 0
    unchanged: int = 0

    def to_dict(self) -> dict[str, Any]:
        return asdict(self)


def infer_symbol(path: str | Path) -> str:
    """TradingView names exports like ``HOSE_DLY_CTS, 1.csv``."""
    stem = Path(path).stem.strip()
    before_resolution = stem.rsplit(",", 1)[0].strip()
    tokens = [token for token in re.split(r"[_\s.-]+", before_resolution.upper()) if token]
    candidate = tokens[-1] if tokens else ""
    return candidate if re.fullmatch(r"[A-Z0-9]{2,12}", candidate) else ""


def infer_exchange(path: str | Path) -> str:
    """Infer HOSE/HNX/UPCOM from the TradingView file prefix."""
    stem = Path(path).stem.strip().upper()
    first = next((token for token in re.split(r"[_\s.-]+", stem) if token), "")
    return normalize_exchange(first)


def normalize_resolution(value: str | int | None) -> str:
    raw = str(value or "").strip().upper().replace("MIN", "").replace("M", "")
    if raw.endswith("H") and raw[:-1].isdigit():
        minutes = int(raw[:-1]) * 60
        return "1H" if minutes == 60 else str(minutes)
    if raw.isdigit() and int(raw) > 0:
        return str(int(raw))
    return ""


def infer_resolution(path: str | Path) -> str:
    stem = Path(path).stem
    match = re.search(r",\s*([0-9]+\s*[mMhH]?)\s*$", stem)
    return normalize_resolution(match.group(1)) if match else ""


def resolution_minutes(value: str) -> int:
    normalized = normalize_resolution(value)
    if normalized == "1H":
        return 60
    return int(normalized) if normalized.isdigit() else 10**9


def aggregate_intraday_bars(
    rows: list[dict[str, Any]],
    target_resolution: str | int,
) -> list[dict[str, Any]]:
    """Aggregate ordered intraday OHLCV rows into larger minute buckets."""
    minutes = resolution_minutes(str(target_resolution))
    if minutes <= 0 or minutes >= 24 * 60:
        raise ValueError("Resolution gộp phải là số phút intraday.")
    grouped: dict[int, dict[str, Any]] = {}
    for source in sorted(rows, key=lambda item: int(item.get("time", 0) or 0)):
        stamp = int(source.get("time", 0) or 0)
        if stamp <= 0:
            continue
        bucket = stamp - stamp % (minutes * 60)
        current = grouped.get(bucket)
        if current is None:
            grouped[bucket] = {
                "time": bucket,
                "open": float(source.get("open", 0.0) or 0.0),
                "high": float(source.get("high", 0.0) or 0.0),
                "low": float(source.get("low", 0.0) or 0.0),
                "close": float(source.get("close", 0.0) or 0.0),
                "volume": float(source.get("volume", 0.0) or 0.0),
                "closed": True,
            }
            continue
        current["high"] = max(float(current["high"]), float(source.get("high", 0.0) or 0.0))
        current["low"] = min(float(current["low"]), float(source.get("low", 0.0) or 0.0))
        current["close"] = float(source.get("close", 0.0) or 0.0)
        current["volume"] = float(current["volume"]) + float(source.get("volume", 0.0) or 0.0)
    return list(grouped.values())


def _column_map(headers: Iterable[Any]) -> dict[str, str]:
    available = {str(value or "").strip().casefold(): str(value or "") for value in headers}
    missing = [name for name in REQUIRED_COLUMNS if name not in available]
    if missing:
        raise ValueError("Thiếu cột bắt buộc: " + ", ".join(missing))
    return {name: available[name] for name in REQUIRED_COLUMNS}


def _raw_rows(path: Path) -> list[dict[str, Any]]:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            reader = csv.DictReader(handle)
            if not reader.fieldnames:
                raise ValueError("File CSV không có dòng tiêu đề.")
            columns = _column_map(reader.fieldnames)
            return [{key: row.get(source) for key, source in columns.items()} for row in reader]
    if suffix == ".xlsx":
        try:
            from openpyxl import load_workbook
        except ImportError as exc:  # pragma: no cover - dependency ships with the app
            raise RuntimeError("Thiếu thư viện openpyxl để đọc Excel.") from exc
        book = load_workbook(path, read_only=True, data_only=True)
        try:
            for sheet in book.worksheets:
                iterator = sheet.iter_rows(values_only=True)
                headers = next(iterator, None)
                if not headers:
                    continue
                try:
                    columns = _column_map(headers)
                except ValueError:
                    continue
                positions = {str(value or ""): index for index, value in enumerate(headers)}
                return [
                    {key: values[positions[source]] if positions[source] < len(values) else None
                     for key, source in columns.items()}
                    for values in iterator
                    if any(value is not None for value in values)
                ]
        finally:
            book.close()
        raise ValueError("Không có sheet nào chứa đủ cột OHLCV.")
    raise ValueError("Chỉ hỗ trợ file .csv hoặc .xlsx.")


def _timestamp(raw: Any, timezone_name: str) -> int:
    zone = ZoneInfo(timezone_name)
    if isinstance(raw, datetime):
        value = raw
    elif isinstance(raw, (int, float)):
        number = float(raw)
        if number > 10**12:
            number /= 1000.0
        return int(number)
    else:
        text = str(raw or "").strip()
        if not text:
            raise ValueError("Timestamp trống.")
        if re.fullmatch(r"\d+(?:\.\d+)?", text):
            number = float(text)
            if number > 10**12:
                number /= 1000.0
            return int(number)
        value = datetime.fromisoformat(text[:-1] + "+00:00" if text.endswith("Z") else text)
    if value.tzinfo is None:
        value = value.replace(tzinfo=zone)
    return int(value.timestamp())


def _minute_of_day(value: datetime) -> int:
    return value.hour * 60 + value.minute


def _bucket_start(minute: int, resolution: str) -> int:
    size = max(1, resolution_minutes(resolution))
    return minute - minute % size


def _inside_session(stamp: datetime, resolution: str, exchange: str = "HOSE") -> bool:
    """Accept TradingView's bar-open timestamp for any intraday resolution.

    A two-minute bar containing the 09:15 ATO print is labelled 09:14 and the
    bar containing the 14:45 ATC print is labelled 14:44.  Requiring literal
    09:15/14:45 therefore rejected valid non-1m exports.
    """
    current = _minute_of_day(stamp)
    opening = _bucket_start(9 * 60, resolution)
    closing = _bucket_start(exchange_close_minute(exchange), resolution)
    return (
        opening <= current <= 11 * 60 + 30
        or 13 * 60 <= current <= closing
    )


def _day_is_full(first: datetime, last: datetime, resolution: str, exchange: str = "HOSE") -> bool:
    size = max(1, resolution_minutes(resolution))
    first_minute = _minute_of_day(first)
    last_minute = _minute_of_day(last)
    opening = _bucket_start(9 * 60, resolution)
    # TradingView only emits a candle when the symbol actually trades.  A
    # complete file can therefore start a few minutes after the continuous
    # session opens (CTS did so at 09:17/09:18) without missing market data.
    # Keep the tolerance narrow so genuinely mid-session first days remain
    # PARTIAL.
    opening_grace = max(size, 15) if exchange == "HOSE" else size
    latest_opening = _bucket_start(exchange_open_minute(exchange), resolution) + opening_grace
    closing = exchange_close_minute(exchange)
    return (
        # A file can legitimately omit the opening bucket when the symbol has
        # no ATO print.  The immediately following bucket is still complete
        # market data; execution already waits for the first real bar.
        opening <= first_minute <= latest_opening
        and last_minute <= closing < last_minute + size
    )


def _normalized_rows(
    path: Path,
    *,
    resolution: str,
    exchange: str,
    timezone_name: str,
    price_scale: float | None,
) -> tuple[list[dict[str, Any]], float]:
    source = _raw_rows(path)
    if not source:
        raise ValueError("File không có dữ liệu.")
    closes: list[float] = []
    for row in source[: min(500, len(source))]:
        try:
            closes.append(abs(float(row.get("close") or 0.0)))
        except (TypeError, ValueError):
            continue
    scale = float(price_scale or 0.0)
    if scale <= 0:
        useful = sorted(value for value in closes if value > 0)
        median = useful[len(useful) // 2] if useful else 0.0
        scale = 1000.0 if median >= 1000.0 else 1.0
    if scale <= 0:
        raise ValueError("Hệ số giá phải lớn hơn 0.")

    normalized: dict[int, dict[str, Any]] = {}
    for line, raw in enumerate(source, start=2):
        try:
            stamp = _timestamp(raw.get("time"), timezone_name)
            local = datetime.fromtimestamp(stamp, VN_TZ)
            if not _inside_session(local, resolution, exchange):
                raise ValueError(f"ngoài giờ giao dịch ({local:%H:%M})")
            opened = float(raw.get("open")) / scale
            high = float(raw.get("high")) / scale
            low = float(raw.get("low")) / scale
            close = float(raw.get("close")) / scale
            volume = float(raw.get("volume") or 0.0)
        except (TypeError, ValueError, OverflowError) as exc:
            raise ValueError(f"Dòng {line} không hợp lệ: {exc}") from exc
        if min(opened, high, low, close) <= 0:
            raise ValueError(f"Dòng {line}: giá phải lớn hơn 0.")
        if high < max(opened, close, low) or low > min(opened, close, high):
            raise ValueError(f"Dòng {line}: OHLC không hợp lệ.")
        if volume < 0:
            raise ValueError(f"Dòng {line}: volume âm.")
        row = {
            "time": stamp,
            "open": opened,
            "high": high,
            "low": low,
            "close": close,
            "volume": volume,
            "closed": True,
        }
        previous = normalized.get(stamp)
        if previous is not None and previous != row:
            raise ValueError(f"Timestamp {local.isoformat()} bị trùng với OHLCV khác nhau trong cùng file.")
        normalized[stamp] = row
    return [normalized[key] for key in sorted(normalized)], scale


def _day_summary(
    rows: list[dict[str, Any]], resolution: str, exchange: str,
) -> tuple[str, str, int, int, list[str]]:
    grouped: dict[str, list[datetime]] = {}
    for row in rows:
        local = datetime.fromtimestamp(int(row["time"]), VN_TZ)
        grouped.setdefault(local.date().isoformat(), []).append(local)
    partial: list[str] = []
    for day, stamps in grouped.items():
        first, last = min(stamps), max(stamps)
        if not _day_is_full(first, last, resolution, exchange):
            partial.append(day)
    days = sorted(grouped)
    return days[0], days[-1], len(days) - len(partial), len(partial), partial


class ReplayDataStore:
    """Transactional, resolution-aware storage for user-imported replay bars."""

    def __init__(self, root: str | Path):
        self.root = Path(root)
        self.path = self.root / "replay.sqlite3"

    def _connect(self) -> sqlite3.Connection:
        self.root.mkdir(parents=True, exist_ok=True)
        connection = sqlite3.connect(self.path)
        connection.row_factory = sqlite3.Row
        connection.execute("PRAGMA foreign_keys = ON")
        connection.executescript(
            """
            CREATE TABLE IF NOT EXISTS replay_bars (
                symbol TEXT NOT NULL,
                resolution TEXT NOT NULL,
                timestamp INTEGER NOT NULL,
                open REAL NOT NULL,
                high REAL NOT NULL,
                low REAL NOT NULL,
                close REAL NOT NULL,
                volume REAL NOT NULL,
                PRIMARY KEY (symbol, resolution, timestamp)
            );
            CREATE INDEX IF NOT EXISTS idx_replay_bars_range
                ON replay_bars(symbol, timestamp);
            CREATE TABLE IF NOT EXISTS replay_days (
                symbol TEXT NOT NULL,
                resolution TEXT NOT NULL,
                day TEXT NOT NULL,
                status TEXT NOT NULL,
                bar_count INTEGER NOT NULL,
                first_timestamp INTEGER NOT NULL,
                last_timestamp INTEGER NOT NULL,
                PRIMARY KEY (symbol, resolution, day)
            );
            CREATE TABLE IF NOT EXISTS replay_datasets (
                symbol TEXT NOT NULL,
                resolution TEXT NOT NULL,
                exchange TEXT NOT NULL DEFAULT '',
                source_name TEXT NOT NULL,
                source_hash TEXT NOT NULL,
                timezone TEXT NOT NULL,
                price_scale REAL NOT NULL,
                imported_at TEXT NOT NULL,
                PRIMARY KEY (symbol, resolution)
            );
            """
        )
        columns = {str(row[1]) for row in connection.execute("PRAGMA table_info(replay_datasets)")}
        migrated = False
        if "exchange" not in columns:
            connection.execute("ALTER TABLE replay_datasets ADD COLUMN exchange TEXT NOT NULL DEFAULT ''")
            migrated = True
        changed: list[tuple[str, str, str]] = []
        for row in connection.execute(
            "SELECT symbol,resolution,source_name,exchange FROM replay_datasets"
        ):
            exchange = normalize_exchange(row["exchange"]) or infer_exchange(row["source_name"])
            if exchange and exchange != str(row["exchange"] or ""):
                connection.execute(
                    "UPDATE replay_datasets SET exchange=? WHERE symbol=? AND resolution=?",
                    (exchange, row["symbol"], row["resolution"]),
                )
                changed.append((row["symbol"], row["resolution"], exchange))
        if migrated or changed:
            for symbol, resolution, exchange in changed:
                self._refresh_days(connection, symbol, resolution, exchange)
        return connection

    def inspect_file(
        self,
        path: str | Path,
        *,
        symbol: str | None = None,
        resolution: str | None = None,
        exchange: str | None = None,
        timezone_name: str = "Asia/Ho_Chi_Minh",
        price_scale: float | None = None,
    ) -> ReplayImportPreview:
        source = Path(path)
        if not source.is_file():
            raise ValueError("Không tìm thấy file import.")
        chosen_symbol = str(symbol or infer_symbol(source)).strip().upper()
        if not re.fullmatch(r"[A-Z0-9]{2,12}", chosen_symbol):
            raise ValueError("Không xác định được mã; hãy nhập mã trước khi lưu.")
        chosen_resolution = normalize_resolution(resolution or infer_resolution(source))
        if not chosen_resolution or resolution_minutes(chosen_resolution) >= 24 * 60:
            raise ValueError("Không xác định được resolution intraday.")
        chosen_exchange = normalize_exchange(exchange or infer_exchange(source))
        if not chosen_exchange:
            raise ValueError("Không xác định được sàn; hãy chọn HOSE, HNX hoặc UPCOM trước khi lưu.")
        rows, scale = _normalized_rows(
            source,
            resolution=chosen_resolution,
            exchange=chosen_exchange,
            timezone_name=timezone_name,
            price_scale=price_scale,
        )
        first, last, full, partial_count, partial = _day_summary(
            rows, chosen_resolution, chosen_exchange,
        )
        digest = sha256(source.read_bytes()).hexdigest()
        return ReplayImportPreview(
            str(source), chosen_symbol, chosen_resolution, chosen_exchange, timezone_name, scale,
            len(rows), first, last, full, partial_count, partial, digest,
        )

    def import_file(self, path: str | Path, **kwargs: Any) -> ReplayImportPreview:
        preview = self.inspect_file(path, **kwargs)
        rows, _ = _normalized_rows(
            Path(path), resolution=preview.resolution,
            exchange=preview.exchange,
            timezone_name=preview.timezone, price_scale=preview.price_scale,
        )
        now = datetime.now().astimezone().isoformat()
        with closing(self._connect()) as connection, connection:
            existing = {
                int(row["timestamp"]): tuple(float(row[key]) for key in ("open", "high", "low", "close", "volume"))
                for row in connection.execute(
                    "SELECT timestamp, open, high, low, close, volume FROM replay_bars "
                    "WHERE symbol=? AND resolution=? AND timestamp BETWEEN ? AND ?",
                    (preview.symbol, preview.resolution, rows[0]["time"], rows[-1]["time"]),
                )
            }
            inserted = replaced = unchanged = 0
            values = []
            for row in rows:
                current = tuple(float(row[key]) for key in ("open", "high", "low", "close", "volume"))
                old = existing.get(int(row["time"]))
                if old is None:
                    inserted += 1
                elif old == current:
                    unchanged += 1
                else:
                    replaced += 1
                values.append((preview.symbol, preview.resolution, int(row["time"]), *current))
            connection.executemany(
                "INSERT INTO replay_bars(symbol,resolution,timestamp,open,high,low,close,volume) "
                "VALUES(?,?,?,?,?,?,?,?) ON CONFLICT(symbol,resolution,timestamp) DO UPDATE SET "
                "open=excluded.open,high=excluded.high,low=excluded.low,close=excluded.close,volume=excluded.volume",
                values,
            )
            connection.execute(
                "INSERT INTO replay_datasets(symbol,resolution,exchange,source_name,source_hash,timezone,price_scale,imported_at) "
                "VALUES(?,?,?,?,?,?,?,?) ON CONFLICT(symbol,resolution) DO UPDATE SET "
                "exchange=excluded.exchange,source_name=excluded.source_name,source_hash=excluded.source_hash,timezone=excluded.timezone,"
                "price_scale=excluded.price_scale,imported_at=excluded.imported_at",
                (preview.symbol, preview.resolution, preview.exchange, Path(path).name, preview.source_hash,
                 preview.timezone, preview.price_scale, now),
            )
            self._refresh_days(connection, preview.symbol, preview.resolution, preview.exchange)
        preview.inserted, preview.replaced, preview.unchanged = inserted, replaced, unchanged
        return preview

    @staticmethod
    def _refresh_days(
        connection: sqlite3.Connection, symbol: str, resolution: str, exchange: str = "HOSE",
    ) -> None:
        connection.execute("DELETE FROM replay_days WHERE symbol=? AND resolution=?", (symbol, resolution))
        grouped: dict[str, list[int]] = {}
        for row in connection.execute(
            "SELECT timestamp FROM replay_bars WHERE symbol=? AND resolution=? ORDER BY timestamp",
            (symbol, resolution),
        ):
            stamp = int(row["timestamp"])
            grouped.setdefault(datetime.fromtimestamp(stamp, VN_TZ).date().isoformat(), []).append(stamp)
        values = []
        for day, stamps in grouped.items():
            first = datetime.fromtimestamp(stamps[0], VN_TZ)
            last = datetime.fromtimestamp(stamps[-1], VN_TZ)
            status = "FULL" if _day_is_full(first, last, resolution, exchange) else "PARTIAL"
            values.append((symbol, resolution, day, status, len(stamps), stamps[0], stamps[-1]))
        connection.executemany(
            "INSERT INTO replay_days(symbol,resolution,day,status,bar_count,first_timestamp,last_timestamp) "
            "VALUES(?,?,?,?,?,?,?)",
            values,
        )

    def list_datasets(self) -> list[ReplayDataset]:
        if not self.path.exists():
            return []
        with closing(self._connect()) as connection, connection:
            rows = connection.execute(
                "SELECT symbol,resolution,exchange,source_name,source_hash,timezone,price_scale,imported_at "
                "FROM replay_datasets ORDER BY symbol,resolution"
            ).fetchall()
        result: list[ReplayDataset] = []
        with closing(self._connect()) as connection, connection:
            for row in rows:
                bar_count = connection.execute(
                    "SELECT COUNT(*) FROM replay_bars WHERE symbol=? AND resolution=?",
                    (row["symbol"], row["resolution"]),
                ).fetchone()[0]
                day_counts = connection.execute(
                    "SELECT MIN(day),MAX(day),SUM(status='FULL'),SUM(status='PARTIAL') FROM replay_days "
                    "WHERE symbol=? AND resolution=?",
                    (row["symbol"], row["resolution"]),
                ).fetchone()
                result.append(ReplayDataset(
                    row["symbol"], row["resolution"], str(row["exchange"] or ""), int(bar_count or 0),
                    str(day_counts[0] or ""), str(day_counts[1] or ""),
                    int(day_counts[2] or 0), int(day_counts[3] or 0),
                    row["source_name"], row["source_hash"], row["timezone"],
                    float(row["price_scale"]), row["imported_at"],
                ))
        return result

    def exchange_for(self, symbol: str, resolution: str | None = None) -> str:
        """Return one unambiguous exchange for a replay symbol."""
        if not self.path.exists():
            return ""
        symbol = str(symbol or "").strip().upper()
        with closing(self._connect()) as connection, connection:
            if resolution:
                rows = connection.execute(
                    "SELECT exchange FROM replay_datasets WHERE symbol=? AND resolution=?",
                    (symbol, normalize_resolution(resolution)),
                ).fetchall()
            else:
                rows = connection.execute(
                    "SELECT DISTINCT exchange FROM replay_datasets WHERE symbol=?", (symbol,),
                ).fetchall()
        values = {normalize_exchange(row[0]) for row in rows}
        values.discard("")
        return next(iter(values)) if len(values) == 1 else ""

    def delete_dataset(self, symbol: str, resolution: str) -> int:
        symbol = str(symbol).strip().upper()
        resolution = normalize_resolution(resolution)
        if not self.path.exists() or not symbol or not resolution:
            return 0
        with closing(self._connect()) as connection, connection:
            count = connection.execute(
                "SELECT COUNT(*) FROM replay_bars WHERE symbol=? AND resolution=?", (symbol, resolution),
            ).fetchone()[0]
            connection.execute("DELETE FROM replay_bars WHERE symbol=? AND resolution=?", (symbol, resolution))
            connection.execute("DELETE FROM replay_days WHERE symbol=? AND resolution=?", (symbol, resolution))
            connection.execute("DELETE FROM replay_datasets WHERE symbol=? AND resolution=?", (symbol, resolution))
        return int(count or 0)

    def set_exchange(self, symbol: str, resolution: str, exchange: str) -> None:
        """Resolve exchange metadata for an already imported legacy dataset."""
        symbol = str(symbol or "").strip().upper()
        resolution = normalize_resolution(resolution)
        market = normalize_exchange(exchange)
        if not symbol or not resolution or not market:
            raise ValueError("Cần chọn mã, resolution và sàn hợp lệ.")
        with closing(self._connect()) as connection, connection:
            found = connection.execute(
                "SELECT 1 FROM replay_datasets WHERE symbol=? AND resolution=?",
                (symbol, resolution),
            ).fetchone()
            if not found:
                raise ValueError("Không tìm thấy bộ dữ liệu cần cập nhật sàn.")
            connection.execute(
                "UPDATE replay_datasets SET exchange=? WHERE symbol=? AND resolution=?",
                (market, symbol, resolution),
            )
            self._refresh_days(connection, symbol, resolution, market)

    def day_status(self, symbol: str, resolution: str, day: str | date) -> str:
        if not self.path.exists():
            return "MISSING"
        value = day.isoformat() if isinstance(day, date) else str(day)[:10]
        with closing(self._connect()) as connection, connection:
            row = connection.execute(
                "SELECT status FROM replay_days WHERE symbol=? AND resolution=? AND day=?",
                (str(symbol).upper(), normalize_resolution(resolution), value),
            ).fetchone()
        return str(row[0]) if row else "MISSING"

    def load_day(
        self,
        symbol: str,
        day: str | date,
        *,
        resolution: str | None = None,
        complete_only: bool = True,
    ) -> tuple[list[dict[str, Any]], str, str]:
        if not self.path.exists():
            return [], "", "MISSING"
        day_text = day.isoformat() if isinstance(day, date) else str(day)[:10]
        symbol = str(symbol).strip().upper()
        with closing(self._connect()) as connection, connection:
            if resolution:
                candidates = [normalize_resolution(resolution)]
            else:
                candidates = [row[0] for row in connection.execute(
                    "SELECT resolution FROM replay_days WHERE symbol=? AND day=? ORDER BY resolution",
                    (symbol, day_text),
                )]
            candidates = sorted(set(filter(None, candidates)), key=resolution_minutes)
            for candidate in candidates:
                status_row = connection.execute(
                    "SELECT status,first_timestamp,last_timestamp FROM replay_days "
                    "WHERE symbol=? AND resolution=? AND day=?",
                    (symbol, candidate, day_text),
                ).fetchone()
                if not status_row:
                    continue
                status = str(status_row["status"])
                if complete_only and status != "FULL":
                    continue
                rows = [dict(row) for row in connection.execute(
                    "SELECT timestamp AS time,open,high,low,close,volume FROM replay_bars "
                    "WHERE symbol=? AND resolution=? AND timestamp BETWEEN ? AND ? ORDER BY timestamp",
                    (symbol, candidate, int(status_row["first_timestamp"]), int(status_row["last_timestamp"])),
                )]
                for row in rows:
                    row["closed"] = True
                return rows, candidate, status
            # A requested larger bucket may be derived from a complete finer
            # dataset. This is deliberately one-way: never invent 1m detail
            # from a 2m file. CTS therefore supports requested 2m from its 1m
            # import while VIX keeps using its exact 2m source.
            requested = normalize_resolution(resolution) if resolution else ""
            target_minutes = resolution_minutes(requested)
            if requested and target_minutes < 24 * 60:
                source_rows = connection.execute(
                    "SELECT resolution,status,first_timestamp,last_timestamp FROM replay_days "
                    "WHERE symbol=? AND day=? AND status='FULL'",
                    (symbol, day_text),
                ).fetchall()
                finer = sorted(
                    (
                        row for row in source_rows
                        if resolution_minutes(str(row["resolution"])) < target_minutes
                        and target_minutes % resolution_minutes(str(row["resolution"])) == 0
                    ),
                    key=lambda row: resolution_minutes(str(row["resolution"])),
                    reverse=True,
                )
                for source in finer:
                    source_resolution = str(source["resolution"])
                    raw_rows = [dict(row) for row in connection.execute(
                        "SELECT timestamp AS time,open,high,low,close,volume FROM replay_bars "
                        "WHERE symbol=? AND resolution=? AND timestamp BETWEEN ? AND ? ORDER BY timestamp",
                        (
                            symbol, source_resolution,
                            int(source["first_timestamp"]), int(source["last_timestamp"]),
                        ),
                    )]
                    aggregated = aggregate_intraday_bars(raw_rows, requested)
                    if aggregated:
                        return (
                            aggregated,
                            requested,
                            f"FULL_AGGREGATED_{source_resolution}_TO_{requested}",
                        )
        return [], "", "MISSING"

    def load_daily_aggregates(
        self,
        symbol: str,
        *,
        before: str | date | None = None,
        resolution: str | None = None,
        complete_only: bool = True,
    ) -> list[dict[str, Any]]:
        """Aggregate imported intraday bars into daily candles.

        This is primarily used to warm up EMA/RSI in strict replay mode.  It
        keeps the warm-up and the simulated session on the same price basis;
        daily bars from another vendor may use different corporate-action
        adjustments.
        """
        if not self.path.exists():
            return []
        symbol = str(symbol or "").strip().upper()
        before_text = before.isoformat() if isinstance(before, date) else str(before or "")[:10]
        with closing(self._connect()) as connection, connection:
            query = "SELECT day,resolution,status FROM replay_days WHERE symbol=?"
            values: list[Any] = [symbol]
            if before_text:
                query += " AND day<?"
                values.append(before_text)
            if resolution:
                query += " AND resolution=?"
                values.append(normalize_resolution(resolution))
            if complete_only:
                query += " AND status='FULL'"
            day_rows = connection.execute(query + " ORDER BY day,resolution", values).fetchall()

            selected: dict[str, str] = {}
            for row in day_rows:
                candidate = str(row["resolution"])
                current = selected.get(str(row["day"]))
                if current is None or resolution_minutes(candidate) < resolution_minutes(current):
                    selected[str(row["day"])] = candidate

            result: list[dict[str, Any]] = []
            for day_text, chosen in sorted(selected.items()):
                bars = connection.execute(
                    "SELECT timestamp,open,high,low,close,volume FROM replay_bars "
                    "WHERE symbol=? AND resolution=? AND date(timestamp,'unixepoch','+7 hours')=? "
                    "ORDER BY timestamp",
                    (symbol, chosen, day_text),
                ).fetchall()
                if not bars:
                    continue
                result.append({
                    "time": int(bars[-1]["timestamp"]),
                    "open": float(bars[0]["open"]),
                    "high": max(float(row["high"]) for row in bars),
                    "low": min(float(row["low"]) for row in bars),
                    "close": float(bars[-1]["close"]),
                    "volume": sum(float(row["volume"]) for row in bars),
                    "closed": True,
                })
        return result
