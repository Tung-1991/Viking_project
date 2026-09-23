from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any, Callable, Iterable
from zoneinfo import ZoneInfo

from .. import config


VN_TZ = ZoneInfo("Asia/Ho_Chi_Minh")
VN100_AS_OF = "2026-09-22"
VN100_SOURCE = (
    "https://staticfile.hsx.vn/Uploads/UploadDocuments/2487402/"
    "Form_Factsheet_MCIndices_VN_T08.2026.pdf"
)
VN100_ORDER_SOURCE = "https://vn.tradingview.com/symbols/HOSE-VN100/components/"

# Current VN100 members ordered by market capitalization.  The order is data,
# not strategy: the popup simply takes the first N symbols requested by the
# operator.  Refresh this tuple and VN100_AS_OF after an index rebalance.
VN100_SYMBOLS_BY_MARKET_CAP: tuple[str, ...] = (
    "VIC", "VHM", "VCB", "BID", "CTG", "TCB", "VPB", "GAS", "MBB", "MCH",
    "HPG", "BSR", "VPL", "STB", "HDB", "LPB", "ACB", "GVR", "VNM", "FPT",
    "VJC", "MWG", "MSN", "TCX", "SSB", "VCK", "SSI", "SHB", "SAB", "VRE",
    "BVH", "VIB", "MSB", "PLX", "VPX", "BCM", "GEE", "TPB", "POW", "HCM",
    "GMD", "VIX", "OCB", "GEX", "NVL", "EIB", "REE", "NAB", "KBC", "FRT",
    "VCI", "VND", "VPI", "SBT", "PNJ", "PVD", "VGC", "DCM", "KDH", "HAG",
    "SJS", "DPM", "KDC", "DXG", "SIP", "PVT", "BMP", "TCH", "NLG", "BAF",
    "PDR", "TAL", "VHC", "BWE", "VCG", "DGW", "VTP", "DSE", "CTR", "EVF",
    "CII", "DIG", "PC1", "HSG", "FTS", "PHR", "BSI", "DBC", "CTD", "HDG",
    "KOS", "NT2", "CTS", "HHV", "CMG", "NKG", "VSC", "PAN", "HT1", "ANV",
)


@dataclass(frozen=True, slots=True)
class VolumeScanOptions:
    scan_count: int = 100
    result_count: int = 20
    sessions: int = 5
    threshold_pct: float = 20.0
    direction: str = "CẢ HAI"

    def validate(self) -> "VolumeScanOptions":
        if not 1 <= int(self.scan_count) <= len(VN100_SYMBOLS_BY_MARKET_CAP):
            raise ValueError("Số mã quét phải từ 1 đến 100.")
        if not 1 <= int(self.result_count) <= int(self.scan_count):
            raise ValueError("Số mã lấy phải từ 1 đến số mã quét.")
        if int(self.sessions) not in {5, 10}:
            raise ValueError("Chu kỳ chỉ hỗ trợ 5 hoặc 10 phiên.")
        if float(self.threshold_pct) <= 0:
            raise ValueError("Ngưỡng phải lớn hơn 0%.")
        if str(self.direction).upper() not in {"TĂNG", "GIẢM", "CẢ HAI"}:
            raise ValueError("Hướng lọc không hợp lệ.")
        return self


@dataclass(frozen=True, slots=True)
class VolumeScanRow:
    symbol: str
    previous_average: float
    recent_average: float
    change_pct: float
    status: str


@dataclass(frozen=True, slots=True)
class VolumeScanSummary:
    scanned: int
    matched: int
    returned: int
    insufficient: int
    zero_base: int
    api_errors: int


@dataclass(frozen=True, slots=True)
class VolumeScanResult:
    rows: tuple[VolumeScanRow, ...]
    summary: VolumeScanSummary


def evaluate_volume_change(
    symbol: str,
    bars: Iterable[dict[str, Any]],
    sessions: int,
) -> tuple[VolumeScanRow | None, str]:
    """Compare two adjacent blocks of completed daily volume bars."""
    sessions = int(sessions)
    if sessions not in {5, 10}:
        raise ValueError("Chu kỳ chỉ hỗ trợ 5 hoặc 10 phiên.")

    closed = sorted(
        (dict(row) for row in bars if isinstance(row, dict) and bool(row.get("closed", False))),
        key=lambda row: float(row.get("time", 0.0) or 0.0),
    )
    required = sessions * 2
    if len(closed) < required:
        return None, "INSUFFICIENT_DATA"

    volumes: list[float] = []
    for row in closed[-required:]:
        try:
            volume = float(row.get("volume"))
        except (TypeError, ValueError):
            return None, "INSUFFICIENT_DATA"
        if volume < 0:
            return None, "INSUFFICIENT_DATA"
        volumes.append(volume)

    previous_average = sum(volumes[:sessions]) / sessions
    recent_average = sum(volumes[sessions:]) / sessions
    if previous_average == 0:
        return None, "ZERO_BASE"
    # Normalize harmless binary floating-point noise so an exact business
    # boundary such as +20% or -20% remains inclusive.
    change_pct = round((recent_average / previous_average - 1.0) * 100.0, 10)
    return (
        VolumeScanRow(
            symbol=str(symbol or "").strip().upper(),
            previous_average=previous_average,
            recent_average=recent_average,
            change_pct=change_pct,
            status="TĂNG" if change_pct >= 0 else "GIẢM",
        ),
        "",
    )


def _matches(row: VolumeScanRow, direction: str, threshold_pct: float) -> bool:
    direction = str(direction).upper()
    threshold = float(threshold_pct)
    if direction == "TĂNG":
        return row.change_pct >= threshold
    if direction == "GIẢM":
        return row.change_pct <= -threshold
    return abs(row.change_pct) >= threshold


def _sort_rows(rows: list[VolumeScanRow], direction: str) -> list[VolumeScanRow]:
    direction = str(direction).upper()
    if direction == "TĂNG":
        return sorted(rows, key=lambda row: (-row.change_pct, row.symbol))
    if direction == "GIẢM":
        return sorted(rows, key=lambda row: (row.change_pct, row.symbol))
    return sorted(rows, key=lambda row: (-abs(row.change_pct), row.symbol))


class VolumeScanner:
    """On-demand VN100 volume scanner; it owns no trading or watchlist state."""

    def __init__(
        self,
        bars_loader: Callable[..., list[dict[str, Any]]],
        *,
        api_health: Callable[[], dict[str, Any]] | None = None,
    ):
        self.bars_loader = bars_loader
        self.api_health = api_health

    def _last_request_failed(self) -> bool:
        if self.api_health is None:
            return False
        try:
            status = self.api_health().get("last_status")
        except Exception:
            return False
        return isinstance(status, (int, float)) and not 200 <= int(status) < 300

    def scan(
        self,
        options: VolumeScanOptions,
        *,
        progress: Callable[[int, int, str], None] | None = None,
    ) -> VolumeScanResult:
        options.validate()
        symbols = VN100_SYMBOLS_BY_MARKET_CAP[: int(options.scan_count)]
        matched: list[VolumeScanRow] = []
        insufficient = zero_base = api_errors = 0

        for index, symbol in enumerate(symbols, start=1):
            try:
                bars = self.bars_loader(symbol, count=40, exchange="HOSE")
            except Exception:
                bars = []
                api_errors += 1
            else:
                if not bars and self._last_request_failed():
                    api_errors += 1
                else:
                    row, reason = evaluate_volume_change(symbol, bars, options.sessions)
                    if reason == "INSUFFICIENT_DATA":
                        insufficient += 1
                    elif reason == "ZERO_BASE":
                        zero_base += 1
                    elif row is not None and _matches(row, options.direction, options.threshold_pct):
                        matched.append(row)
            if progress:
                progress(index, len(symbols), symbol)

        ordered = _sort_rows(matched, options.direction)
        rows = tuple(ordered[: int(options.result_count)])
        return VolumeScanResult(
            rows=rows,
            summary=VolumeScanSummary(
                scanned=len(symbols),
                matched=len(matched),
                returned=len(rows),
                insufficient=insufficient,
                zero_base=zero_base,
                api_errors=api_errors,
            ),
        )


def export_volume_scan(
    rows: Iterable[VolumeScanRow],
    *,
    output_dir: str | Path | None = None,
    now: datetime | None = None,
) -> Path:
    values = list(rows)
    if not values:
        raise ValueError("Không có kết quả để xuất Excel.")
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
    except ImportError as exc:  # pragma: no cover - dependency ships with Viking
        raise RuntimeError("Thiếu thư viện openpyxl để xuất Excel.") from exc

    root = Path(output_dir or (config.RUNTIME_ROOT / "exports"))
    root.mkdir(parents=True, exist_ok=True)
    timestamp = now or datetime.now(VN_TZ)
    stem = f"vn100_volume_{timestamp:%Y%m%d_%H%M%S}"
    path = root / f"{stem}.xlsx"
    suffix = 2
    while path.exists():
        path = root / f"{stem}_{suffix}.xlsx"
        suffix += 1

    book = Workbook()
    sheet = book.active
    sheet.title = "VOLUME"
    headers = ("MÃ CK", "TB KỲ TRƯỚC", "TB KỲ GẦN NHẤT", "% THAY ĐỔI", "TRẠNG THÁI")
    sheet.append(headers)
    for row in values:
        sheet.append((
            row.symbol,
            row.previous_average,
            row.recent_average,
            row.change_pct / 100.0,
            row.status,
        ))

    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
    for row_index in range(2, sheet.max_row + 1):
        sheet.cell(row_index, 2).number_format = "#,##0"
        sheet.cell(row_index, 3).number_format = "#,##0"
        sheet.cell(row_index, 4).number_format = "0.00%;[Red]-0.00%"
    for column, width in zip("ABCDE", (12, 20, 22, 18, 16)):
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:E{sheet.max_row}"
    book.save(path)
    book.close()
    return path
