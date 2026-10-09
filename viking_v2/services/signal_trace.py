"""Bounded, per-account observations. No network, order or strategy mutations."""
from __future__ import annotations

from datetime import datetime, timedelta
from contextlib import closing
import hashlib
import json
import math
from pathlib import Path
import sqlite3
from typing import Any, Iterable

from ..config import AppSettings
from ..connections.dnse.snapshot_health import safe_detail
from ..rules.business import StaticRuleParameters
from ..trading.market import VN_TZ
from ..trading.validation import decision_is_fresh, quote_diagnostics, timestamp_value
from .indicator_comparison import number_comparison


def finite(value: Any) -> float | None:
    try:
        result = float(value)
        return result if math.isfinite(result) else None
    except (TypeError, ValueError, OverflowError):
        return None


def compare_entry(marks: dict, params: dict) -> tuple[bool | None, bool | None, bool | None]:
    fast = finite(marks.get("buy_ema_fast", marks.get("ema_fast")))
    slow = finite(marks.get("buy_ema_slow", marks.get("ema_slow")))
    current, prior = finite(marks.get("rsi")), finite(marks.get("rsi_previous"))
    ema = fast > slow if fast is not None and slow is not None else None
    rsi = current > prior if current is not None and prior is not None else None
    chosen = [value for enabled, value in ((params.get("buy_signal_use_ema", True), ema),
                                          (params.get("buy_signal_use_rsi", True), rsi)) if enabled]
    entry = None if not chosen or any(value is None for value in chosen) else all(chosen)
    return ema, rsi, entry


def export_trace(rows: list[dict], path: str | Path) -> None:
    """Export the selected audit rows and only their allowlisted settings."""
    from openpyxl import Workbook
    book = Workbook()
    sheet = book.active
    replay = bool(rows) and all(row.get("execution_mode") == "REPLAY" for row in rows)
    sheet.title = "ENTRY REPLAY" if replay else "TRACE"
    headers = (
        ("timestamp", "Ghi nhận"), ("scheduled_at", "Mốc lấy mẫu"), ("symbol", "Mã"),
        ("execution_mode", "Chế độ"), ("priority", "Ưu tiên"), ("price_vnd", "Giá VND"),
        ("price_source", "Nguồn giá"), ("quote_age_seconds", "Tuổi giá giây"),
        ("quote_issue", "Lỗi giá"), ("decision_time", "Giờ đánh giá"),
        ("system_issue", "Lỗi hệ thống lúc ghi"),
        ("decision_fresh", "Đánh giá mới"), ("ema_fast", "EMA nhanh"), ("ema_slow", "EMA chậm"),
        ("ema_comparison", "So EMA"), ("rsi", "RSI hiện tại"), ("rsi_previous", "RSI phiên trước"),
        ("ema_fast_period", "Chu kỳ EMA nhanh"), ("ema_slow_period", "Chu kỳ EMA chậm"), ("rsi_period", "Chu kỳ RSI"),
        ("rsi_previous_date", "Phiên RSI trước"), ("rsi_comparison", "So RSI"),
        ("entry", "ENTRY EMA/RSI"), ("ema_cross_required", "Cần vừa cắt EMA"),
        ("fresh_cross_entry", "ENTRY nếu cần vừa cắt"), ("baseline_source", "Nguồn nền 1D"),
        ("exit_e", "EXIT E"), ("rule_action", "Xử lý APP"), ("reason", "Mã lý do"),
        ("bot_on", "BOT ON"), ("otp_ok", "OTP còn hạn"), ("pause_active", "Tạm khóa BUY"),
        ("market_phase", "Phiên"), ("buy_window_state", "Giờ mua"),
        ("whipsaw_count", "WHIPSAW đếm"), ("whipsaw_limit", "WHIPSAW ngưỡng"),
        ("whipsaw_window", "WHIPSAW phiên"), ("whipsaw_on", "WHIPSAW ON"),
        ("loss_streak", "Chuỗi lỗ"), ("loss_limit", "Ngưỡng khóa lỗ"), ("loss_blocked", "Khóa lỗ"),
        ("slot_usage", "Slot"), ("entry_orders_used", "Lệnh mã đã dùng"),
        ("entry_orders_max", "Lệnh mã tối đa"), ("available_cash", "Tiền khả dụng"),
        ("available_capital", "Vốn được mua"), ("order_budget", "Vốn/lệnh"),
        ("priority_limit", "Hạn mức Priority"), ("priority_reserved", "Tiền giữ mã khác"),
        ("budget_price_vnd", "Giá dự trù VND"), ("queue_summary", "Lệnh liên quan đã ghi"),
        ("settings_hash", "Phiên bản setting"),
    )
    if replay:
        # Replay cannot know cash, OTP or actual broker orders. Omit those
        # empty columns entirely instead of presenting an apparent live trace.
        keys = {"timestamp", "symbol", "price_vnd", "ema_comparison", "rsi_comparison",
                "rsi_previous_date", "entry", "ema_cross_required", "fresh_cross_entry",
                "baseline_source", "whipsaw_count", "whipsaw_limit", "whipsaw_window",
                "buy_window_state", "settings_hash"}
        headers = tuple(item for item in headers if item[0] in keys) + (("reason_text", "Kết quả kỹ thuật"),)
    sheet.append([title for _, title in headers])
    for row in rows:
        values = []
        for key, _ in headers:
            value = row.get(key, "")
            if replay and key in {"entry", "fresh_cross_entry"}:
                value = "ĐẠT" if value is True else "CHƯA ĐẠT" if value is False else "—"
            values.append(value)
        sheet.append(values)
    settings_sheet = book.create_sheet("SETTING")
    settings_sheet.append(["Phiên bản", "Setting dùng lúc ghi (không chứa token)"])
    unique = {row.get("settings_hash", ""): row.get("settings", {}) for row in rows}
    for fingerprint, settings in unique.items():
        settings_sheet.append([fingerprint, json.dumps(settings, ensure_ascii=False, sort_keys=True)])
    for ws in book.worksheets:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for cells in ws.iter_rows():
            for cell in cells:
                if isinstance(cell.value, str) and cell.value.startswith(("=", "+", "-", "@")):
                    cell.data_type = "s"
        for column in ws.columns:
            ws.column_dimensions[column[0].column_letter].width = min(55, max(16, max(len(str(c.value or "")) for c in column) + 2))
    try:
        book.save(path)
    finally:
        book.close()


class SignalTraceStore:
    RETENTION_DAYS = 30

    def __init__(self, path: str | Path):
        self.path = Path(path)
        self._last_capture: tuple | None = None

    @staticmethod
    def schedule(settings: AppSettings, now: datetime) -> tuple[str, str] | None:
        if not settings.signal_trace_enabled:
            return None
        now = now.astimezone(VN_TZ)
        start = datetime.strptime(settings.signal_trace_start, "%H:%M").time()
        end = datetime.strptime(settings.signal_trace_end, "%H:%M").time()
        boundary = now.replace(hour=start.hour, minute=start.minute, second=0, microsecond=0)
        end_boundary = now.replace(hour=end.hour, minute=end.minute, second=0, microsecond=0) + timedelta(minutes=1)
        if now.weekday() >= 5 or not boundary <= now < end_boundary:
            return None
        interval = settings.signal_trace_interval_minutes * 60
        bucket = int((now - boundary).total_seconds()) // interval
        scheduled = boundary + timedelta(seconds=bucket * interval)
        return now.date().isoformat(), scheduled.isoformat()

    @staticmethod
    def settings_snapshot(settings: AppSettings) -> dict:
        # Never serialize AppSettings wholesale: it also has private chat info.
        keys = ("signal_mode", "realtime_indicator_interval", "rule_parameters", "watchlist",
                "priority_symbols", "priority_capital_enabled", "priority_total_capital",
                "priority_allocations", "bot_order_mode", "allow_ato", "allow_atc",
                "buy_fee_pct", "signal_trace_interval_minutes", "signal_trace_start", "signal_trace_end",
                "market_phase_override_enabled", "market_phase_override", "market_phase_override_exposure_pct")
        snapshot = {key: getattr(settings, key) for key in keys}
        snapshot["rule_parameters"] = StaticRuleParameters.from_dict(settings.rule_parameters).to_dict()
        return snapshot

    def capture(self, status: dict, settings: AppSettings, *, active_mode: str,
                intents: Iterable[Any] = (), otp_ok: bool = False, pause_by_mode: dict | None = None,
                now: datetime | None = None) -> int:
        observed = (now or datetime.now(VN_TZ)).astimezone(VN_TZ)
        scheduled = self.schedule(settings, observed)
        if scheduled is None or observed.date().isoformat() in settings.trading_holidays:
            return 0
        books = status.get("decisions_by_mode") or {active_mode: status.get("decisions", {})}
        modes = [active_mode, *[m for m, rows in books.items() if m != active_mode and rows]]
        symbols = list(dict.fromkeys([*settings.watchlist, *status.get("active_symbols", [])]))
        signature = (*scheduled, settings.signal_trace_interval_minutes, tuple(modes), tuple(symbols))
        if self._last_capture == signature:
            return 0
        snapshot = self.settings_snapshot(settings)
        fingerprint = hashlib.sha256(json.dumps(snapshot, sort_keys=True).encode()).hexdigest()[:16]
        saved_intents = [item.to_dict() if hasattr(item, "to_dict") else item for item in intents]
        ticks = status.get("ticks") or {}
        rows = []
        for mode in modes:
            decisions = books.get(mode) or {}
            for symbol in symbols:
                raw = decisions.get(symbol) or {}
                details = raw.get("details") or {}
                fresh = decision_is_fresh(raw, symbol, mode, now=observed.timestamp())
                marks = (details.get("indicators") or {}) if fresh else {}
                checks = (details.get("entry_checks") or {}) if fresh else {}
                tick = ticks.get(symbol) or {}
                quote = quote_diagnostics(tick, symbol, now=observed.timestamp())
                _, _, entry = compare_entry(marks, settings.rule_parameters)
                if not quote["valid"]:
                    entry = None  # Keep recorded numeric evidence, never label stale data as a live ENTRY.
                prior_date = details.get("rsi_previous_date", "")
                prior_time = timestamp_value(marks.get("rsi_previous_time"))
                if prior_time and math.isfinite(prior_time):
                    prior_date = datetime.fromtimestamp(prior_time, VN_TZ).date().isoformat()
                priority = checks.get("priority_capital") or {}
                related = [r for r in saved_intents if isinstance(r, dict) and r.get("symbol") == symbol
                           and r.get("execution_mode") == mode
                           and datetime.fromtimestamp(float(r.get("created_at") or 0), VN_TZ).date() == observed.date()]
                row = {
                    "timestamp": observed.isoformat(), "scheduled_at": scheduled[1], "symbol": symbol,
                    "execution_mode": mode, "priority": settings.priority_symbols.index(symbol) + 1 if symbol in settings.priority_symbols else "",
                    "price_vnd": finite(tick.get("price")), "price_source": quote["source"],
                    "quote_age_seconds": quote["age_seconds"], "quote_issue": quote["reason"],
                    "system_issue": safe_detail(status.get("error", "")),
                    "decision_time": details.get("updated_at", ""), "decision_fresh": fresh,
                    "ema_fast": marks.get("buy_ema_fast", marks.get("ema_fast")),
                    "ema_slow": marks.get("buy_ema_slow", marks.get("ema_slow")),
                    "rsi": marks.get("rsi"), "rsi_previous": marks.get("rsi_previous"),
                    "ema_fast_period": marks.get("buy_ema_fast_period", marks.get("ema_fast_period")),
                    "ema_slow_period": marks.get("buy_ema_slow_period", marks.get("ema_slow_period")),
                    "rsi_period": marks.get("rsi_period"),
                    "rsi_previous_date": prior_date, "entry": entry,
                    "ema_cross_required": bool(settings.rule_parameters.get("buy_signal_require_ema_cross", True)),
                    "exit_e": raw.get("signal") == "SELL" if fresh and quote["valid"] else None,
                    "rule_action": raw.get("action", "WAIT") if fresh else "WAIT",
                    "reason": raw.get("reason", "") if fresh else "NO_FRESH_DECISION",
                    "bot_on": bool(status.get("bot_enabled")), "otp_ok": bool(otp_ok) if mode == "REAL" else None,
                    "pause_active": bool((pause_by_mode or {}).get(mode, {}).get("active")),
                    "market_phase": (status.get("symbol_phases") or {}).get(symbol, status.get("market_status", "")),
                    "buy_window_state": (details.get("buy_window") or {}).get("state", ""),
                    "whipsaw_count": checks.get("whipsaw_crossovers"), "whipsaw_limit": checks.get("whipsaw_limit"),
                    "whipsaw_window": checks.get("whipsaw_window"), "whipsaw_on": checks.get("whipsaw_enabled"),
                    "loss_streak": checks.get("loss_streak"), "loss_limit": checks.get("loss_lock_count"),
                    "loss_blocked": checks.get("loss_blocked"),
                    "slot_usage": f"{checks.get('open_positions', '—')}/{checks.get('max_positions', '—')}",
                    "entry_orders_used": checks.get("entry_orders_used"), "entry_orders_max": checks.get("entry_orders_max"),
                    "available_cash": checks.get("available_cash"), "available_capital": checks.get("available_capital"),
                    "order_budget": checks.get("order_budget"), "priority_limit": priority.get("limit_vnd"),
                    "priority_reserved": priority.get("reserved_cash"), "budget_price_vnd": checks.get("buy_budget_price"),
                    "queue_summary": " · ".join(f"{r.get('side')} {r.get('quantity')} #{r.get('id', '')[:8]} {r.get('status')} khớp {r.get('filled_quantity', 0)}" for r in related),
                    "settings_hash": fingerprint, "settings": snapshot,
                }
                # One broken indicator must not poison the whole sample batch.
                for key, value in row.items():
                    if isinstance(value, float) and not math.isfinite(value):
                        row[key] = None
                for key in ("price_vnd", "budget_price_vnd"):
                    if row.get(key) is not None:
                        row[key] *= 1000
                row["ema_comparison"] = number_comparison(row["ema_fast"], row["ema_slow"], 4)
                row["rsi_comparison"] = number_comparison(row["rsi"], row["rsi_previous"])
                rows.append(row)
        self.path.parent.mkdir(parents=True, exist_ok=True)
        with closing(sqlite3.connect(self.path, timeout=0.5)) as db, db:
            db.execute("CREATE TABLE IF NOT EXISTS samples (day TEXT, scheduled TEXT, mode TEXT, symbol TEXT, payload TEXT, PRIMARY KEY(day, scheduled, mode, symbol))")
            before = db.total_changes
            db.executemany("INSERT OR IGNORE INTO samples VALUES (?, ?, ?, ?, ?)",
                           [(scheduled[0], scheduled[1], row["execution_mode"], row["symbol"],
                             json.dumps(row, ensure_ascii=False, allow_nan=False)) for row in rows])
            inserted = db.total_changes - before
            cutoff = (observed.date() - timedelta(days=self.RETENTION_DAYS - 1)).isoformat()
            db.execute("DELETE FROM samples WHERE day < ?", (cutoff,))
        self._last_capture = signature
        return inserted

    def read(self, *, day: str = "", symbol: str = "", mode: str = "", limit: int = 2000) -> list[dict]:
        if not self.path.exists():
            return []
        filters, args = [], []
        for key, value in (("day", day), ("symbol", symbol), ("mode", mode)):
            if value:
                filters.append(f"{key} = ?")
                args.append(value)
        where = " WHERE " + " AND ".join(filters) if filters else ""
        with closing(sqlite3.connect(self.path.resolve().as_uri() + "?mode=ro", uri=True, timeout=0.5)) as db:
            return [json.loads(row[0]) for row in db.execute(
                f"SELECT payload FROM samples{where} ORDER BY day DESC, scheduled DESC, symbol, mode LIMIT ?",
                [*args, max(1, min(100000, int(limit)))])]

    def days(self) -> list[str]:
        if not self.path.exists():
            return []
        with closing(sqlite3.connect(self.path.resolve().as_uri() + "?mode=ro", uri=True, timeout=0.5)) as db:
            return [row[0] for row in db.execute("SELECT DISTINCT day FROM samples ORDER BY day DESC")]
