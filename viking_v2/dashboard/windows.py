from __future__ import annotations

from datetime import datetime
from concurrent.futures import ThreadPoolExecutor
from math import ceil, isfinite
import re
from pathlib import Path
import sqlite3
import tkinter as tk
from tkinter import font as tkfont
from tkinter import ttk, filedialog, messagebox
from typing import Any, Callable

import customtkinter as ctk

from ..branding import window_title
from ..services.indicator_comparison import DNSEIndicatorNormalizer, number_comparison, rsi_observation_display
from ..services.signal_trace import SignalTraceStore, trace_workbook
from ..services.signal_history import (
    SignalHistoryTrash, history_sort_key, observation_id, periodic_history_row, recording_options,
)
from ..rules.observations import ema_cross_caption
from ..trading.market import VN_TZ

# One palette for every popup, so the app cannot drift into four colour schemes.
PALETTE = {
    "BG": "#111318",        # window background
    "PANEL": "#15181D",     # tab body
    "SURFACE": "#22262D",   # cards and table rows
    "SURFACE_2": "#1B1F25", # inputs, headers, striped rows
    "BORDER": "#343A43",    # every border and chip
    "SLATE": "#3A3F47",     # neutral buttons
    "SLATE_HOVER": "#4B515B",
    "TEXT": "#E8EBEF",
    "TITLE": "#BABEC5",     # neutral headings/keys; readable without white glare
    "MUTED": "#C5CBD4",
    "DIM": "#98A2B3",
    "BLUE": "#2B6CB0",
    "BLUE_HOVER": "#245C92",
    "GREEN": "#22C55E",
    "GREEN_HOVER": "#16A34A",
    "RED": "#EF4444",
    "WARN": "#F59E0B",
}

FONT_KEY = ("Segoe UI", 14, "bold", "italic")
FONT_VALUE = ("Segoe UI", 14)
FONT_MONO_VALUE = ("Cascadia Mono", 14)
FONT_TABLE_HEADING = ("Segoe UI", 16, "bold", "italic")
FONT_TABLE_VALUE = ("Segoe UI", 14)
HINT_FONT = ("Segoe UI", 20)


def fit_entry_text(
    entry: ctk.CTkEntry,
    text: str,
    *,
    base_font: tuple[Any, ...] = FONT_MONO_VALUE,
    minimum_size: int = 11,
    horizontal_padding: int = 22,
) -> int:
    """Keep a dynamic entry value readable without letting it be clipped.

    Normal values stay at the standard 14-point size.  Only values wider than
    the actual input area are reduced, and never below the UI minimum of 11.
    """
    family = str(base_font[0])
    base_size = max(minimum_size, int(base_font[1]))
    styles = tuple(str(item) for item in base_font[2:])
    inner = getattr(entry, "_entry", None)
    inner_width = int(inner.winfo_width()) if inner is not None else 0
    available = max(
        0,
        inner_width - 8 if inner_width > 1
        else int(entry.winfo_width()) - horizontal_padding,
    )
    selected = base_size
    if available > 0 and text:
        top = entry.winfo_toplevel()
        for size in range(base_size, minimum_size - 1, -1):
            options: dict[str, Any] = {"root": top, "family": family, "size": size}
            if "bold" in styles:
                options["weight"] = "bold"
            if "italic" in styles:
                options["slant"] = "italic"
            if tkfont.Font(**options).measure(text) <= available:
                selected = size
                break
        else:
            selected = minimum_size
    if getattr(entry, "_viking_fit_font_size", None) != selected:
        entry.configure(font=(family, selected, *styles))
        entry._viking_fit_font_size = selected
    return selected


def fit_label_text(
    label: ctk.CTkLabel,
    *,
    base_font: tuple[Any, ...],
    minimum_size: int = 10,
) -> None:
    """Fit a compact, single-line value using its actual scaled Tk font."""
    if not isinstance(label, ctk.CTkLabel) or label.winfo_width() <= 1:
        return
    text = str(label.cget("text") or "")
    available = max(0, label.winfo_width() - 4 * label._get_widget_scaling())
    signature = (text, available, base_font, minimum_size, label._get_widget_scaling())
    if getattr(label, "_viking_fit_label_signature", None) == signature:
        return
    label._viking_fit_label_signature = signature
    family, base_size, *styles = base_font
    selected = minimum_size
    for size in range(int(base_size), minimum_size - 1, -1):
        font = tkfont.Font(root=label, font=label._apply_font_scaling((family, size, *styles)))
        if font.measure(text) <= available:
            selected = size
            break
    if getattr(label, "_viking_fit_font_size", None) != selected:
        label.configure(font=(family, selected, *styles))
        label._viking_fit_font_size = selected


_SIGNAL_REASONS = {
    "WHIPSAW_LOCK": "EMA nhiễu, khóa mua",
    "MAX_POSITIONS": "Đã đủ số vị thế",
    "MAX_SYMBOL_ORDERS": "Đã đủ số lần BUY của mã trong vị thế này",
    "POSITION_NOT_READY_FOR_ADD": "Chưa được mua thêm: vị thế/thoát lệnh cần kiểm tra",
    "BOT_OFF": "BOT đang tắt",
    "MANUAL_SELL_PAUSE": "Tạm khóa BUY sau khi operator bán tay",
    "LOCKED_AFTER_3_LOSSES": "Khóa sau chuỗi 3 lệnh lỗ",
    "LOCKED_AFTER_LOSSES": "Khóa sau chuỗi lệnh lỗ",
    "NO_AVAILABLE_CAPITAL": "Không còn vốn khả dụng",
    "BUY_ALREADY_PENDING": "Đã có lệnh mua chờ",
    "BROKER_REJECTED": "Broker từ chối lệnh BUY",
    "BROKER_FAILED": "Gửi lệnh BUY tới broker thất bại",
    "CORPORATE_ACTION_BLOCK": "Đang chặn vì sự kiện quyền",
    "MARKET_STATE_UNKNOWN": "Chưa xác nhận trạng thái thị trường",
    "NO_NEW_BUY_SIGNAL": "Chưa có tín hiệu BUY mới",
    "NO_FRESH_DECISION": "Chưa có quyết định mới",
    "BUY_SIGNAL": "Đạt ENTRY của rule",
    "ENTRY_CONDITIONS_LOST": "Điều kiện ENTRY không còn đạt",
    "BUY_CONFIRMATION_WAIT": "Đang giữ điều kiện BUY đủ số phút đã đặt",
    "BUY_WINDOW_WAIT": "Tín hiệu đang chờ đến khung giờ mua",
    "BUY_WINDOW_BROKEN": "Điều kiện BUY mất trong khi chờ giờ mua",
    "BUY_WINDOW_EXPIRED": "Đã hết khung giờ mua trong ngày",
    "BUY_WINDOW_MARKET_CLOSED": "Ngoài phiên giao dịch của sàn",
    "BUY_FILTER_NEEDS_REALTIME": "Bộ lọc BUY theo giờ cần chế độ REALTIME",
    "BUY_CONFIRMATION_BROKEN": "Tín hiệu BUY không giữ đủ thời gian",
    "BUY_CONFIRMATION_NEEDS_REALTIME": "Xác nhận theo phút cần chế độ REALTIME",
    "UNKNOWN_EXCHANGE": "Chưa xác định sàn của mã",
    "HOLD_POSITION": "Tiếp tục giữ vị thế",
    "MANUAL_OR_EXTERNAL_POSITION": "Vị thế ngoài bot quản lý",
}


def signal_advice(row: dict[str, Any]) -> tuple[str, str]:
    """Translate raw rule fields into a short operator-facing suggestion."""
    signal = str(row.get("signal", "") or "").upper()
    acted = str(row.get("acted", "") or "").upper()
    blocked = str(row.get("blocked_by", "") or "").upper()
    if blocked == "BUY_CONFIRMATION_WAIT":
        held = str(row.get("confirmation_minutes", "0") or "0")
        required = str(row.get("confirmation_required", "") or "")
        return "CHỜ BUY", f"Xác nhận {held}/{required} phút"
    if blocked == "BUY_WINDOW_WAIT":
        window = str(row.get("buy_window", "") or "")
        return "CHỜ GIỜ MUA", f"Chưa xếp lệnh; chờ khung {window}" if window else "Chưa xếp lệnh; chờ khung giờ mua"
    if blocked == "ENTRY_CONDITIONS_LOST":
        return "MẤT ENTRY", "EMA/RSI không còn đạt; không phải lệnh bán"
    if blocked in {"BUY_WINDOW_BROKEN", "BUY_CONFIRMATION_BROKEN"}:
        return "HỦY CHỜ BUY", _SIGNAL_REASONS[blocked] + "; không phải BUY mới"
    if blocked == "BUY_WINDOW_EXPIRED":
        return "HẾT GIỜ MUA", _SIGNAL_REASONS[blocked]
    if signal == "SELL" and acted == "WAIT" and blocked == "NO_NEW_BUY_SIGNAL":
        return "CHỈ TÍN HIỆU SELL", "Ghi nhận tín hiệu SELL; không tạo lệnh bán"
    if blocked or acted == "WAIT":
        suggestion = "BUY BỊ CHẶN" if signal == "BUY" else "CHƯA XẾP SELL" if signal == "SELL" else "CHỈ THÔNG BÁO"
        return suggestion, _SIGNAL_REASONS.get(blocked, blocked.replace("_", " ") or "Rule chưa cho phép")
    if signal == "BUY" and acted == "BUY":
        return "ĐÃ XẾP BUY", "Đã tạo yêu cầu BUY; chưa xác nhận broker đã nhận/khớp"
    if signal == "SELL" and acted == "SELL":
        return "ĐÃ XẾP SELL", "Đã tạo yêu cầu SELL; chưa xác nhận broker đã nhận/khớp"
    return "THEO DÕI", "Có tín hiệu nhưng chưa hành động"


def compact_signal_events(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """Collapse consecutive legacy repeats for display only, preserving raw audit.

    Modern event IDs and distinct ENTRY cycles are retained. Different prices
    alone do not constitute a new legacy E event. Original IDs remain recoverable.
    """
    output: list[dict[str, Any]] = []
    last: dict[tuple[str, str, str], tuple[tuple, int]] = {}
    for raw in sorted((row for row in rows if isinstance(row, dict)),
                      key=lambda row: str(row.get("timestamp", ""))):
        row = dict(raw)
        row["_history_ids"] = list(raw.get("_history_ids") or [observation_id(raw)])
        signal = str(row.get("signal", "")).upper()
        if signal not in {"BUY", "SELL", "E ALERT"}:
            continue
        key = (str(row.get("timestamp", ""))[:10], str(row.get("execution_mode", "")).upper(),
               str(row.get("symbol", "")).upper())
        signature = (signal, row.get("acted", ""), row.get("blocked_by", ""),
                     row.get("buy_window_state", ""), row.get("signal_cycle", ""),
                     row.get("signal_event", ""))
        previous = last.get(key)
        if row.get("record_kind") != "SIGNAL_EVENT" and previous and previous[0] == signature:
            original = output[previous[1]]
            original["legacy_repeat_count"] = original.get("legacy_repeat_count", 1) + 1
            original["legacy_last_seen"] = row.get("timestamp", "")
            original["_history_ids"].extend(row["_history_ids"])
            continue
        last[key] = (signature, len(output))
        output.append(row)
    return output


def signal_rows_by_day(rows: list[dict[str, Any]], *, compact: bool = False) -> list[dict[str, Any]]:
    """Group events without hiding distinct times that happen to share prices."""
    days: dict[str, list[dict[str, Any]]] = {}
    seen: dict[str, dict[tuple[Any, ...], int]] = {}
    for raw in compact_signal_events(rows or []) if compact else rows or []:
        if not isinstance(raw, dict):
            continue
        timestamp = str(raw.get("timestamp", "") or "")
        signal = str(raw.get("signal", "") or "").upper()
        if len(timestamp) < 10 or not signal:
            continue
        day = timestamp[:10]
        signature = (
            timestamp,
            str(raw.get("symbol", "") or "").upper(), signal,
            str(raw.get("execution_mode", "") or "").upper(),
            str(raw.get("price", "") or ""), str(raw.get("ema_fast", "") or ""),
            str(raw.get("ema_slow", "") or ""), str(raw.get("rsi", "") or ""),
            str(raw.get("market_state", "") or "").upper(),
            str(raw.get("acted", "") or "").upper(),
            str(raw.get("blocked_by", "") or "").upper(),
            str(raw.get("rsi_previous", "") or ""),
            str(raw.get("signal_cycle", "") or ""), str(raw.get("candle_key", "") or ""),
        )
        suggestion, reason = signal_advice(raw)
        legacy = not (raw.get("record_kind") or raw.get("signal_event"))
        rsi_text, rsi_issue = rsi_observation_display(raw.get("rsi"), raw.get("rsi_previous"), legacy=legacy)
        detail = {
            **raw,
            "timestamp": timestamp,
            "time": timestamp[11:19] if len(timestamp) >= 19 else timestamp,
            "first_recorded_at": timestamp,
            "symbol": str(raw.get("symbol", "") or "").upper(),
            "signal": signal,
            "display_signal": (
                "LỖI LỆNH" if raw.get("signal_event") == "ORDER_RESULT"
                else "MẤT ENTRY" if raw.get("signal_event") == "ENTRY_LOST"
                or str(raw.get("blocked_by", "") or "").upper()
                in {"BUY_WINDOW_BROKEN", "BUY_CONFIRMATION_BROKEN", "ENTRY_CONDITIONS_LOST"}
                else "EXIT · E" if signal in {"SELL", "E ALERT"}
                else "ENTRY" if signal == "BUY" else signal
            ),
            "suggestion": suggestion,
            "reason": reason,
            "repeat_count": 1,
            "ema_comparison": number_comparison(raw.get("ema_fast"), raw.get("ema_slow"), 4),
            "rsi_comparison": rsi_text,
            "_rsi_issue": rsi_issue,
            "_legacy_record": legacy,
            "rsi_previous_date": raw.get("rsi_previous_date") or ("Bản cũ chưa lưu" if legacy else "—"),
            "display_price": _signal_price(raw.get("price")),
            "ema_cross_display": (ema_cross_caption({
                "state": ("CROSSED_UP" if raw.get("ema_cross_at")
                          and str(raw.get("ema_cross_required", "")).lower() == "true"
                          and raw.get("buy_window_state") in {"WAITING", "ALLOWED"}
                          else raw.get("ema_cross_state", "UNKNOWN")),
                "required": str(raw.get("ema_cross_required", "")).lower() == "true",
                "cross_at": raw.get("ema_cross_at", ""),
            })[0].removeprefix("CẮT EMA · ") if raw.get("ema_cross_state")
                else "Bản cũ chưa lưu" if legacy else "Chưa có dữ liệu"),
        }
        previous = seen.setdefault(day, {}).get(signature)
        if previous is not None:
            existing = days[day][previous]
            existing["first_recorded_at"] = min(existing["first_recorded_at"], timestamp)
            existing["timestamp"] = max(existing["timestamp"], timestamp)
            existing["time"] = existing["timestamp"][11:19]
            existing["repeat_count"] += 1
            existing.setdefault("_history_ids", []).extend(raw.get("_history_ids") or [observation_id(raw)])
            continue
        seen[day][signature] = len(days.setdefault(day, []))
        days[day].append(detail)

    output: list[dict[str, Any]] = []
    for day, details in days.items():
        details.sort(key=lambda row: str(row.get("timestamp", "")), reverse=True)
        output.append({
            "date": day,
            "rows": details,
            "symbols": sorted({row["symbol"] for row in details}),
            "waiting_count": sum(row.get("blocked_by") in {"BUY_WINDOW_WAIT", "BUY_CONFIRMATION_WAIT"} for row in details),
            "lost_count": sum(row.get("display_signal") == "MẤT ENTRY" for row in details),
            "buy_count": sum(str(row.get("signal", "")).upper() == "BUY" for row in details),
            "sell_count": sum(str(row.get("signal", "")).upper() in {"SELL", "E ALERT"} for row in details),
            "allowed_count": sum(
                str(row.get("signal", "") or "").upper() == "BUY"
                and
                str(row.get("acted", "") or "").upper() == "BUY"
                and not str(row.get("blocked_by", "") or "")
                for row in details
            ),
            "blocked_count": sum(
                str(row.get("signal", "") or "").upper() == "BUY"
                and (
                    str(row.get("acted", "") or "").upper() == "WAIT"
                    or bool(str(row.get("blocked_by", "") or ""))
                )
                and row.get("blocked_by") not in {
                    "BUY_WINDOW_WAIT", "BUY_CONFIRMATION_WAIT", "BUY_WINDOW_BROKEN", "BUY_CONFIRMATION_BROKEN",
                    "ENTRY_CONDITIONS_LOST",
                }
                for row in details
            ),
        })
    return sorted(output, key=lambda group: str(group.get("date", "")), reverse=True)


def _signal_price(value: Any) -> str:
    try:
        number = float(value)
        return f"{number * 1000:,.0f} đ" if isfinite(number) and number > 0 else "—"
    except (TypeError, ValueError, OverflowError):
        return "—"


def install_fast_scroll(root: ctk.CTk, pixels: int = 300) -> None:
    """Replace CustomTkinter's mouse wheel with a pixel-sized step.

    The stock handler scrolls in canvas "units", which on these long panels is
    a few pixels per notch and makes every popup feel stuck.  One binding here
    covers every scrollable frame in the app because the wheel is bound to the
    application-wide "all" tag.
    """
    root.unbind_all("<MouseWheel>")

    def wheel(event: Any) -> str | None:
        widget = event.widget
        while widget is not None:
            if isinstance(widget, tk.Canvas):
                first, last = widget.yview()
                if (first, last) == (0.0, 1.0):
                    return None
                box = widget.bbox("all")
                total = max(1, (box[3] - box[1]) if box else 1)
                notches = int(event.delta / 120) or (1 if event.delta > 0 else -1)
                widget.yview_moveto(first - notches * pixels / total)
                return "break"
            widget = getattr(widget, "master", None)
        return None

    root.bind_all("<MouseWheel>", wheel, add="+")


def parse_symbols(raw: str) -> tuple[list[str], list[str]]:
    """Split a free-text ticker list into accepted and rejected symbols."""
    values = list(dict.fromkeys(
        item.strip().upper() for item in re.split(r"[,;\s]+", str(raw or "")) if item.strip()
    ))
    invalid = [item for item in values if not re.fullmatch(r"[A-Z][A-Z0-9]{2,11}", item)]
    return [item for item in values if item not in invalid], invalid


class SymbolPicker(ctk.CTkFrame):
    """Shared ticker picker: one input plus removable chips.

    Both the connection popup's watchlist and the backtest's symbol list use
    this, so the two screens can never drift apart in look or in validation.
    """

    SURFACE_2 = PALETTE["SURFACE_2"]
    BORDER = PALETTE["BORDER"]
    MUTED = PALETTE["MUTED"]
    BLUE = PALETTE["BLUE"]
    RED = PALETTE["RED"]
    WARN = PALETTE["WARN"]

    def __init__(
        self,
        parent: Any,
        symbols: list[str] | None = None,
        *,
        columns: int = 8,
        placeholder: str = "FPT, SSI, VCB",
        on_change: Callable[[list[str]], None] | None = None,
        on_configure: Callable[[str], None] | None = None,
        compact: bool = False,
        **kwargs: Any,
    ):
        super().__init__(parent, fg_color="transparent", **kwargs)
        # Compact drops the separate status line and shrinks the chips, for
        # forms where vertical space is scarce.
        self.compact = bool(compact)
        self.columns = max(1, int(columns))
        self.on_change = on_change
        self.on_configure = on_configure
        self._symbols = list(dict.fromkeys(str(v).strip().upper() for v in (symbols or []) if str(v).strip()))
        self.grid_columnconfigure(0, weight=1)

        add_row = ctk.CTkFrame(self, fg_color="transparent")
        add_row.grid(row=0, column=0, sticky="ew", pady=(0, 6))
        add_row.grid_columnconfigure(0, weight=1)
        self.entry = ctk.CTkEntry(
            add_row, height=38, placeholder_text=placeholder,
            fg_color=self.SURFACE_2, border_color=self.BORDER,
            font=("Cascadia Mono", 12),
        )
        self.entry.grid(row=0, column=0, sticky="ew", padx=(0, 8))
        self.entry.bind("<Return>", lambda _event: self.add_from_entry(), add="+")
        ctk.CTkButton(
            add_row, text="+ THÊM", width=100, height=38, font=("Segoe UI", 11, "bold"),
            fg_color=self.BLUE, hover_color=PALETTE["BLUE_HOVER"], command=self.add_from_entry,
        ).grid(row=0, column=1)

        self.chips = ctk.CTkFrame(self, fg_color=self.SURFACE_2, corner_radius=8)
        self.chips.grid(row=1, column=0, sticky="ew")
        self.chips.grid_columnconfigure(tuple(range(self.columns)), weight=1)
        self.status = ctk.CTkLabel(
            self, text="", font=("Segoe UI", 11), text_color=self.MUTED, anchor="w",
        )
        if not self.compact:
            self.status.grid(row=2, column=0, sticky="w", pady=(3, 0))
        self.render()

    def get(self) -> list[str]:
        return list(self._symbols)

    def set(self, symbols: list[str]) -> None:
        self._symbols = list(dict.fromkeys(str(v).strip().upper() for v in symbols if str(v).strip()))
        self.render()

    def add_from_entry(self) -> None:
        symbols, invalid = parse_symbols(self.entry.get())
        if invalid:
            self.status.configure(text=f"MÃ KHÔNG HỢP LỆ: {', '.join(invalid)}", text_color=self.RED)
            return
        if not symbols:
            self.status.configure(text="HÃY NHẬP MÃ CKCS", text_color=self.WARN)
            return
        self._symbols = list(dict.fromkeys(self._symbols + symbols))
        self.entry.delete(0, "end")
        self.render()

    def remove(self, symbol: str) -> None:
        self._symbols = [value for value in self._symbols if value != symbol]
        self.render()

    def render(self) -> None:
        for child in self.chips.winfo_children():
            child.destroy()
        if not self._symbols:
            ctk.CTkLabel(
                self.chips, text="CHƯA CÓ MÃ", font=("Segoe UI", 11, "bold"), text_color=self.WARN,
            ).grid(row=0, column=0, sticky="w", padx=10, pady=10)
        for index, symbol in enumerate(self._symbols):
            if self.on_configure:
                chip = ctk.CTkFrame(self.chips, fg_color="transparent")
                chip.grid(row=index // self.columns, column=index % self.columns, sticky="ew", padx=4, pady=2)
                ctk.CTkButton(
                    chip, text=f"{symbol} ×", width=70, height=30,
                    font=("Cascadia Mono", 11, "bold"),
                    command=lambda value=symbol: self.remove(value),
                ).pack(side="left", fill="x", expand=True)
                ctk.CTkButton(
                    chip, text="⚙", width=30, height=30,
                    command=lambda value=symbol: self.on_configure(value),
                ).pack(side="left", padx=(2, 0))
                continue
            ctk.CTkButton(
                self.chips, text=f"{symbol}  ×", width=98, height=30 if self.compact else 36,
                fg_color=PALETTE["BORDER"], hover_color=PALETTE["SLATE_HOVER"],
                font=("Cascadia Mono", 11, "bold"),
                command=lambda value=symbol: self.remove(value),
            ).grid(row=index // self.columns, column=index % self.columns,
                   sticky="ew", padx=4, pady=2 if self.compact else 5)
        self.status.configure(
            text=f"{len(self._symbols)} MÃ · CLICK × ĐỂ BỎ", text_color=self.MUTED,
        )
        if self.on_change:
            self.on_change(list(self._symbols))

class _StableToplevel(ctk.CTkToplevel):
    # CustomTkinter redraws the Windows title bar by withdrawing the window.
    # That conflicts with the app's own hide/show popup behaviour.
    _deactivate_windows_window_header_manipulation = True


class _HistoryTabview(ctk.CTkTabview):
    def _grid_forget_all_tabs(self, exclude_name=None):
        # CTk queues this 100 ms after set(). A stale callback must not hide
        # the new current tab if the user switches again before it executes.
        if exclude_name is not None and exclude_name != self.get():
            return
        super()._grid_forget_all_tabs(exclude_name)

def _window(parent: ctk.CTk, title: str, geometry: str = "560x420") -> ctk.CTkToplevel:
    top = _StableToplevel(parent)
    top.title(title)
    top.geometry(geometry)
    top.grid_columnconfigure(0, weight=1)
    return top


def minimize_popup(popup: Any) -> bool:
    """Minimize a persistent popup without destroying or withdrawing it."""
    top = getattr(popup, "top", None)
    if top is None or not top.winfo_exists():
        return False
    if str(top.state()).lower() in {"iconic", "withdrawn"}:
        return False
    top.iconify()
    on_visibility_changed = getattr(popup, "on_visibility_changed", None)
    if callable(on_visibility_changed):
        on_visibility_changed(False)
    return True


class _HoverHint:
    def __init__(self, widget: Any, text: str | Callable[[], str], placement: str = "side"):
        self.widget, self.text, self.popup = widget, text, None
        self.placement = placement
        widget.bind("<Enter>", self._show, add="+")
        widget.bind("<Leave>", self._hide, add="+")

    def _show(self, _event: Any = None) -> None:
        text = self.text() if callable(self.text) else self.text
        if self.popup or not text:
            return
        # Keep the hint inside the same Tk window. On Windows with DPI scaling,
        # mixing winfo_root* coordinates with a new Toplevel can place the hint
        # hundreds of pixels away from its icon.
        host = self.widget.winfo_toplevel()
        self.popup = tk.Label(
            host, text=text, justify="left", wraplength=min(780, max(240, host.winfo_width() - 56)),
            bg="#252A31", fg="#F5F7FA", padx=16, pady=14,
            font=HINT_FONT, relief="solid", borderwidth=2,
        )
        self.popup.update_idletasks()
        width, height = self.popup.winfo_reqwidth(), self.popup.winfo_reqheight()
        host.update_idletasks()
        host_width = max(1, host.winfo_width())
        host_height = max(1, host.winfo_height())
        widget_x = self.widget.winfo_rootx() - host.winfo_rootx()
        widget_y = self.widget.winfo_rooty() - host.winfo_rooty()
        widget_width = self.widget.winfo_width()
        widget_height = self.widget.winfo_height()

        if self.placement in {"below", "inside"}:
            x = widget_x
            y = widget_y + widget_height + 8
            if y + height > host_height - 12:
                y = widget_y - height - 8
        else:
            x = widget_x + widget_width + 8
            if x + width > host_width - 12:
                # Keep the tooltip edge next to the icon instead of pinning the
                # tooltip to a distant screen corner.
                x = widget_x - width - 8
            y = widget_y - 4
            if y + height > host_height - 12:
                y = widget_y - height - 8

        x = min(max(8, x), max(8, host_width - width - 12))
        y = min(max(8, y), max(8, host_height - height - 12))
        self.popup.place(x=x, y=y)
        self.popup.lift()

    def _hide(self, _event: Any = None) -> None:
        if self.popup:
            self.popup.destroy()
            self.popup = None

class DataTablePopup:
    """Reusable, non-modal CKCS viewer used by the main toolbar.

    The window is deliberately detached from Tk's transient-window lifecycle.
    Clicking the dashboard minimizes it, and opening it again restores the same
    instance. This matches the RULE/CONNECTION popup behaviour.
    """

    def __init__(
        self,
        parent: ctk.CTk,
        *,
        title: str,
        columns: tuple[tuple[str, str, int, str], ...],
        rows_provider: Callable[[str], list[tuple[str, tuple[Any, ...], tuple[str, ...]]]],
        initial_mode: str = "PAPER",
        summary_provider: Callable[[str], list[tuple[str, str, str]]] | None = None,
        on_visibility_changed: Callable[[bool], None] | None = None,
    ):
        self.parent = parent
        self.rows_provider = rows_provider
        self.summary_provider = summary_provider
        self.on_visibility_changed = on_visibility_changed
        self.columns = columns
        parent.update_idletasks()
        screen_w = max(1100, int(parent.winfo_screenwidth() or 1100))
        screen_h = max(700, int(parent.winfo_screenheight() or 700))
        parent_w = max(0, int(parent.winfo_width() or 0))
        parent_h = max(0, int(parent.winfo_height() or 0))
        width = min(screen_w - 48, max(1080, parent_w - 36))
        height = min(screen_h - 80, max(650, parent_h - 52))
        x = max(16, int(parent.winfo_rootx()) + max(0, (parent_w - width) // 2))
        y = max(16, int(parent.winfo_rooty()) + max(0, (parent_h - height) // 2))
        self.top = _window(parent, window_title(title.upper()), f"{width}x{height}+{x}+{y}")
        try:
            self.top.grab_release()
        except tk.TclError:
            pass
        self.top.tk.call("wm", "transient", self.top._w, "")
        self.top.resizable(True, True)
        self.top.minsize(900, 480)
        self.top.configure(fg_color=PALETTE["BG"])
        self.top.grid_columnconfigure(0, weight=1)
        self.top.grid_rowconfigure(2, weight=1)
        self.top.protocol("WM_DELETE_WINDOW", self.close)
        self.top.bind("<Escape>", lambda _event: self.hide(), add="+")

        header = ctk.CTkFrame(self.top, fg_color="transparent")
        header.grid(row=0, column=0, sticky="ew", padx=18, pady=(14, 5))
        header.grid_columnconfigure(0, weight=1)
        ctk.CTkLabel(
            header, text=title.upper(), font=("Segoe UI", 18, "bold"),
            text_color=PALETTE["TITLE"], anchor="w",
        ).grid(row=0, column=0, sticky="w")
        ctk.CTkButton(
            header, text="↻  LÀM MỚI", width=126, height=36,
            font=("Segoe UI", 12, "bold"), fg_color=PALETTE["BLUE"],
            hover_color=PALETTE["BLUE_HOVER"], command=self.refresh,
        ).grid(row=0, column=1, sticky="e")

        self.summary = ctk.CTkFrame(
            self.top, height=78, fg_color=PALETTE["SURFACE"],
            corner_radius=10, border_width=1, border_color=PALETTE["BORDER"],
        )
        self.summary.grid(row=1, column=0, sticky="ew", padx=14, pady=(1, 5))
        self.summary.grid_propagate(False)
        self.summary_labels: list[tuple[ctk.CTkLabel, ctk.CTkLabel]] = []
        for column in range(6):
            self.summary.grid_columnconfigure(column, weight=1, uniform="portfolio-summary")
            cell = ctk.CTkFrame(self.summary, fg_color="transparent")
            cell.grid(row=0, column=column, sticky="nsew", padx=8, pady=8)
            label = ctk.CTkLabel(
                cell, text="--", font=FONT_KEY,
                text_color=PALETTE["TITLE"], anchor="center",
            )
            label.pack(fill="x")
            value = ctk.CTkLabel(
                cell, text="--", font=FONT_MONO_VALUE,
                text_color=PALETTE["TEXT"], anchor="center",
            )
            value.pack(fill="x", pady=(2, 0))
            self.summary_labels.append((label, value))

        self.tabs = ctk.CTkTabview(
            self.top,
            command=self._refresh_summary,
            fg_color=PALETTE["PANEL"], border_width=1,
            border_color=PALETTE["BORDER"],
            segmented_button_selected_color=PALETTE["GREEN"],
            segmented_button_selected_hover_color=PALETTE["GREEN_HOVER"],
            segmented_button_unselected_color=PALETTE["SLATE"],
            segmented_button_unselected_hover_color=PALETTE["SLATE_HOVER"],
        )
        self.tabs.grid(row=2, column=0, sticky="nsew", padx=14, pady=(2, 14))
        try:
            self.tabs._segmented_button.configure(font=("Segoe UI", 12, "bold"))
        except AttributeError:
            pass

        style = ttk.Style()
        style.theme_use("clam")
        style.configure(
            "V2Popup.Treeview", background=PALETTE["SURFACE"], foreground=PALETTE["TEXT"],
            fieldbackground=PALETTE["SURFACE"], rowheight=48, font=FONT_TABLE_VALUE,
            borderwidth=0, relief="flat",
        )
        style.layout("V2Popup.Treeview", [("V2Popup.Treeview.treearea", {"sticky": "nswe"})])
        style.configure(
            "V2Popup.Treeview.Heading", background=PALETTE["SURFACE_2"],
            foreground=PALETTE["TITLE"], font=FONT_TABLE_HEADING,
            relief="flat", padding=(10, 9),
        )
        style.map(
            "V2Popup.Treeview",
            background=[("selected", "#2B6CB0")], foreground=[("selected", "#FFFFFF")],
        )
        for orientation in ("Vertical", "Horizontal"):
            style.configure(
                f"V2Popup.{orientation}.TScrollbar",
                background="#30353D", troughcolor="#181B20",
                bordercolor="#181B20", arrowcolor="#98A2B3",
                lightcolor="#30353D", darkcolor="#30353D",
            )

        self.trees: dict[str, ttk.Treeview] = {}
        self.empty_labels: dict[str, ctk.CTkLabel] = {}
        keys = tuple(item[0] for item in columns)
        for mode in ("REAL", "PAPER"):
            frame = self.tabs.add(f"CKCS {mode}")
            frame.grid_columnconfigure(0, weight=1)
            frame.grid_rowconfigure(0, weight=1)
            tree = ttk.Treeview(
                frame, columns=keys, show="headings", selectmode="extended",
                style="V2Popup.Treeview",
            )
            for key, heading, width_px, anchor in columns:
                tree.heading(key, text=heading)
                tree.column(
                    key, width=width_px, minwidth=min(width_px, 72),
                    anchor=anchor, stretch=True,
                )
            tree.tag_configure("profit", background="#193524", foreground="#EAFBF0")
            tree.tag_configure("loss", background="#3A2024", foreground="#FFF1F2")
            tree.tag_configure("pending", background="#42351B", foreground="#FEF3C7")
            tree.grid(row=0, column=0, sticky="nsew", padx=(4, 0), pady=(4, 0))
            yscroll = ttk.Scrollbar(
                frame, orient="vertical", command=tree.yview,
                style="V2Popup.Vertical.TScrollbar",
            )
            xscroll = ttk.Scrollbar(
                frame, orient="horizontal", command=tree.xview,
                style="V2Popup.Horizontal.TScrollbar",
            )
            yscroll.grid(row=0, column=1, sticky="ns", pady=(4, 0))
            xscroll.grid(row=1, column=0, sticky="ew", padx=(4, 0))
            tree.configure(yscrollcommand=yscroll.set, xscrollcommand=xscroll.set)
            self.trees[mode] = tree
            empty = ctk.CTkLabel(
                frame, text="CHƯA CÓ DỮ LIỆU", font=("Segoe UI", 12, "bold"),
                text_color="#667085", fg_color="#22262D",
            )
            self.empty_labels[mode] = empty
        self.tabs.set("CKCS REAL" if str(initial_mode).upper() == "REAL" else "CKCS PAPER")
        self.refresh()
        self.show()

    def show(self) -> None:
        self.top.deiconify()
        self.top.lift()
        self.top.focus_force()
        if self.on_visibility_changed:
            self.on_visibility_changed(True)

    def hide(self) -> None:
        if self.top.winfo_exists():
            self.top.withdraw()
        if self.on_visibility_changed:
            self.on_visibility_changed(False)

    def close(self) -> None:
        if self.on_visibility_changed:
            self.on_visibility_changed(False)
        if self.top.winfo_exists():
            self.top.destroy()

    def _active_mode(self) -> str:
        try:
            return "REAL" if self.tabs.get().upper().endswith("REAL") else "PAPER"
        except (AttributeError, tk.TclError):
            return "PAPER"

    def _refresh_summary(self) -> None:
        values = self.summary_provider(self._active_mode()) if self.summary_provider else []
        values = list(values[:6])
        while len(values) < 6:
            values.append(("", "", PALETTE["TEXT"]))
        for (label_widget, value_widget), (label, value, color) in zip(self.summary_labels, values):
            label_widget.configure(text=str(label), text_color=PALETTE["TITLE"])
            value_widget.configure(text=str(value), text_color=str(color or PALETTE["TEXT"]))

    def refresh(self) -> None:
        for mode, tree in self.trees.items():
            tree.delete(*tree.get_children())
            rows = self.rows_provider(mode)
            for iid, values, tags in rows:
                safe_iid = str(iid or "")
                kwargs: dict[str, Any] = {"values": values, "tags": tags}
                if safe_iid and not tree.exists(safe_iid):
                    kwargs["iid"] = safe_iid
                tree.insert("", "end", **kwargs)
            empty = self.empty_labels[mode]
            if rows:
                empty.place_forget()
            else:
                empty.place(relx=0.5, rely=0.5, anchor="center")
        self._refresh_summary()


def history_counts(rows: list[dict[str, Any]]) -> str:
    samples = sum(row.get("record_kind") == "PERIODIC" for row in rows)
    return f"{len(rows) - samples} sự kiện · {samples} mẫu"


def short_signal_reason(code: str, row: dict[str, Any]) -> str:
    """Keep the grid concise; original details remain available on right click."""
    reasons = {
        "BUY_WINDOW_WAIT": "Chờ " + str(row.get("buy_window") or "giờ mua"),
        "BUY_WINDOW_BROKEN": "EMA/RSI mất", "ENTRY_CONDITIONS_LOST": "EMA/RSI mất",
        "BUY_CONFIRMATION_BROKEN": "Mất xác nhận", "BUY_CONFIRMATION_WAIT": "Chờ xác nhận",
        "BUY_WINDOW_EXPIRED": "Hết giờ mua", "BUY_WINDOW_MARKET_CLOSED": "Ngoài phiên",
        "WHIPSAW_LOCK": "Khóa WHIPSAW", "LOCKED_AFTER_LOSSES": "Khóa chuỗi lỗ",
        "LOCKED_AFTER_3_LOSSES": "Khóa chuỗi lỗ", "NO_AVAILABLE_CAPITAL": "Thiếu vốn",
        "MAX_POSITIONS": "Hết slot", "MAX_SYMBOL_ORDERS": "Đủ lệnh/mã",
        "BOT_OFF": "BOT OFF", "MANUAL_SELL_PAUSE": "Tạm khóa BUY",
        "BUY_ALREADY_PENDING": "Đã có lệnh chờ", "BUY_SIGNAL": "EMA/RSI đạt",
        "BROKER_REJECTED": "DNSE từ chối", "BROKER_FAILED": "Lỗi gửi DNSE",
    }
    if row.get("quote_issue") or row.get("system_issue"):
        return "Dữ liệu cần kiểm tra"
    if row.get("display_signal") == "EXIT · E":
        return "E xuất hiện" if row.get("acted") != "SELL" else "E · Đã xếp yêu cầu"
    if row.get("acted") == "BUY" and not row.get("blocked_by"):
        return "Đã xếp yêu cầu"
    if code == "NO_NEW_BUY_SIGNAL":
        return "Chờ cắt EMA" if row.get("ema_cross_required") is True and row.get("entry") else "Chưa có ENTRY"
    return reasons.get(code, code.replace("_", " ")[:65] or "Chỉ ghi nhận")


class SignalRecordingPopup:
    """One small form; never writes trading rules or Telegram settings."""

    def __init__(self, parent, settings, on_save, *, on_saved=None):
        self.on_save, self.on_saved = on_save, on_saved
        self.top = _window(parent, "GHI TÍN HIỆU", "500x340")
        self.top.resizable(False, False)
        self.top.configure(fg_color=PALETTE["PANEL"])
        self.enabled = tk.BooleanVar(value=settings.signal_trace_enabled)
        body = ctk.CTkFrame(self.top, fg_color="transparent")
        body.pack(fill="both", expand=True, padx=22, pady=18)
        body.grid_columnconfigure(1, weight=1)
        self.enabled_switch = ctk.CTkSwitch(body, text="Lưu chỉ báo định kỳ", variable=self.enabled,
                                           text_color=PALETTE["TEXT"], progress_color=PALETTE["GREEN"])
        self.enabled_switch.grid(row=0, column=0, columnspan=2, sticky="w", pady=(0, 16))
        self.hints = {}
        self.hint_buttons = {}
        hints = {
            "enabled": "Bật: lưu giá, EMA/RSI, vốn/slot và các khóa theo lịch, kể cả mã chưa có ENTRY.\n"
                       "Tắt: chỉ ngừng ghi định kỳ; ENTRY / mất ENTRY / EXIT E vẫn ghi khi đổi.\n"
                       "Không gửi Telegram, không tạo lệnh, không bù mẫu quá khứ.",
            "start": "Giờ Việt Nam bắt đầu ghi định kỳ, không phải giờ được phép mua.\n"
                     "Ví dụ 14:00. App phải đang mở; ngày nghỉ không ghi.",
            "end": "Giờ Việt Nam kết thúc ghi định kỳ. Mốc cuối được ghi nếu đúng nhịp.\n"
                   "14:00–14:30, mỗi 2 phút: có mẫu 14:30. Không đổi giờ mua của BOT.",
            "interval": "Khoảng cách giữa các mẫu, từ 1 đến 30 phút.\n"
                        "2 phút: 14:00, 14:02, 14:04…14:30. Không đổi nhịp tính EMA/RSI.\n"
                        "Mỗi mốc chỉ ghi một lần/mã/sổ; không gửi Telegram.",
        }

        def add_hint(key, row, widget):
            button = ctk.CTkButton(body, text="?", width=28, height=28,
                                  font=("Segoe UI", 13, "bold"), fg_color=PALETTE["BLUE"])
            button.grid(row=row, column=2, padx=(10, 0), sticky="e")
            hint = _HoverHint(button, hints[key], placement="below")
            button.configure(command=hint._show)
            self.hints[key], self.hint_buttons[key] = hint, button
            _HoverHint(widget, hints[key], placement="below")

        add_hint("enabled", 0, self.enabled_switch)
        self.entries = {}
        for index, (key, label, value) in enumerate((
            ("start", "Từ giờ (HH:MM)", settings.signal_trace_start),
            ("end", "Đến giờ (HH:MM)", settings.signal_trace_end),
            ("interval", "Mỗi (phút)", settings.signal_trace_interval_minutes),
        ), start=1):
            ctk.CTkLabel(body, text=label, text_color=PALETTE["TEXT"]).grid(
                row=index, column=0, sticky="w", pady=4)
            entry = ctk.CTkEntry(body, width=120, fg_color=PALETTE["SURFACE_2"],
                                border_color=PALETTE["BORDER"], text_color=PALETTE["TEXT"])
            entry.insert(0, str(value))
            entry.grid(row=index, column=1, sticky="e", pady=4)
            self.entries[key] = entry
            add_hint(key, index, entry)
        ctk.CTkLabel(body, text="Mẫu: lưu cả mã chưa có ENTRY.\n"
                                "Không gửi Telegram, không tạo lệnh.\n"
                                "Sự kiện vẫn ghi khi trạng thái đổi.",
                     justify="left", anchor="w", text_color=PALETTE["MUTED"]).grid(
            row=4, column=0, columnspan=3, sticky="w", pady=12)
        self.save_button = ctk.CTkButton(body, text="LƯU", command=self.save,
                                       fg_color=PALETTE["GREEN"], hover_color=PALETTE["GREEN_HOVER"],
                                       state="normal" if on_save else "disabled")
        self.save_button.grid(row=5, column=1, columnspan=2, sticky="e")
        self.top.update_idletasks()
        scale = self.top._get_window_scaling()
        widget_scale = body._get_widget_scaling()
        self.top.geometry(f"{max(500, ceil((body.winfo_reqwidth() + 44 * widget_scale) / scale))}x"
                          f"{max(340, ceil((body.winfo_reqheight() + 36 * widget_scale) / scale))}")
        self.top.protocol("WM_DELETE_WINDOW", self.top.destroy)

    def save(self) -> None:
        if self.on_save is None:
            return
        try:
            options = recording_options(self.enabled.get(), self.entries["interval"].get(),
                                        self.entries["start"].get(), self.entries["end"].get())
            self.on_save(options)
        except (OSError, ValueError, RuntimeError) as exc:
            messagebox.showerror("Ghi tín hiệu", str(exc), parent=self.top)
            return
        self.top.destroy()
        if self.on_saved:
            self.on_saved()


class HistoryPopup:
    """Trade history grouped as day -> trade -> broker/local events.

    A BUY, a cancelled waiting SELL and the eventual filled SELL are three
    events of one trade, not three independent trades.  The hierarchy makes
    that relationship explicit while still preserving the audit trail.
    """

    COLUMNS = (
        ("time", "THỜI GIAN", 155, "center"),
        ("order", "ORDER ID", 145, "center"),
        ("price_qty", "GIÁ / KHỐI LƯỢNG", 210, "center"),
        ("cost", "PHÍ / THUẾ", 145, "center"),
        ("pnl", "PNL", 145, "center"),
        ("status", "TRẠNG THÁI", 170, "center"),
        ("note", "GHI CHÚ", 390, "w"),
    )

    def __init__(
        self,
        parent: ctk.CTk,
        rows_provider: Callable[[str], list[dict[str, Any]]],
        *,
        initial_mode: str = "PAPER",
        on_visibility_changed: Callable[[bool], None] | None = None,
        signals_provider: Callable[[], list[dict[str, Any]]] | None = None,
        indicator_normalizer: DNSEIndicatorNormalizer | None = None,
        trace_store: SignalTraceStore | None = None,
        trace_settings_provider: Callable[[], Any] | None = None,
        on_trace_settings: Callable[[dict[str, Any]], None] | None = None,
        history_trash: SignalHistoryTrash | None = None,
    ):
        self.parent = parent
        self.rows_provider = rows_provider
        self.signals_provider = signals_provider
        if indicator_normalizer is None:
            from .. import config
            root = trace_store.path.parent if trace_store else config.RUNTIME_ROOT
            indicator_normalizer = DNSEIndicatorNormalizer(root / "market_bars.json")
        self.indicator_normalizer = indicator_normalizer
        self.trace_store = trace_store or SignalTraceStore(indicator_normalizer.cache_path.parent / "signal_trace.sqlite3")
        self.trace_settings_provider = trace_settings_provider
        self.on_trace_settings = on_trace_settings
        self.history_trash = history_trash or SignalHistoryTrash(self.trace_store.path.parent / "signal_history_trash.json")
        self.normalization_enabled = False
        self._normalization_rows: dict[str, dict[str, Any]] = {}
        self._normalization_request_key = None
        self._normalization_future = None
        self._normalization_after = None
        self._normalization_executor = None
        self._normalization_generation = 0
        self.on_visibility_changed = on_visibility_changed
        parent.update_idletasks()
        screen_w = max(1100, int(parent.winfo_screenwidth() or 1100))
        screen_h = max(700, int(parent.winfo_screenheight() or 700))
        width = min(screen_w - 48, max(1120, int(parent.winfo_width() or 0) - 36))
        height = min(screen_h - 80, max(680, int(parent.winfo_height() or 0) - 52))
        x = max(16, int(parent.winfo_rootx()) + max(0, (int(parent.winfo_width() or 0) - width) // 2))
        y = max(16, int(parent.winfo_rooty()) + max(0, (int(parent.winfo_height() or 0) - height) // 2))
        self.top = _window(parent, window_title("LỊCH SỬ"), f"{width}x{height}+{x}+{y}")
        try:
            self.top.grab_release()
        except tk.TclError:
            pass
        self.top.tk.call("wm", "transient", self.top._w, "")
        self.top.resizable(True, True)
        self.top.minsize(980, 560)
        self.top.configure(fg_color=PALETTE["BG"])
        self.top.grid_columnconfigure(0, weight=1)
        self.top.grid_rowconfigure(1, weight=1)
        self.top.protocol("WM_DELETE_WINDOW", self.close)
        self.top.bind("<Escape>", lambda _event: self.hide(), add="+")

        header = ctk.CTkFrame(self.top, fg_color="transparent")
        header.grid(row=0, column=0, sticky="ew", padx=18, pady=(14, 4))
        header.grid_columnconfigure(0, weight=1)
        ctk.CTkLabel(
            header, text="LỊCH SỬ GIAO DỊCH", font=("Segoe UI", 18, "bold"),
            text_color=PALETTE["TITLE"], anchor="w",
        ).grid(row=0, column=0, sticky="w")
        self.subtitle = ctk.CTkLabel(
            header, text="Mỗi ngày → mỗi giao dịch → các lần đặt/hủy/khớp",
            font=("Segoe UI", 12), text_color=PALETTE["MUTED"], anchor="w",
        )
        self.subtitle.grid(row=1, column=0, sticky="w", pady=(2, 0))
        ctk.CTkButton(
            header, text="↻  LÀM MỚI", width=126, height=36,
            font=("Segoe UI", 12, "bold"), fg_color=PALETTE["BLUE"],
            hover_color=PALETTE["BLUE_HOVER"], command=self.refresh,
        ).grid(row=0, column=1, rowspan=2, sticky="e")

        self.tabs = _HistoryTabview(
            self.top, fg_color=PALETTE["PANEL"], border_width=1,
            command=self._history_tab_changed,
            border_color=PALETTE["BORDER"],
            segmented_button_selected_color=PALETTE["GREEN"],
            segmented_button_selected_hover_color=PALETTE["GREEN_HOVER"],
            segmented_button_unselected_color=PALETTE["SLATE"],
            segmented_button_unselected_hover_color=PALETTE["SLATE_HOVER"],
        )
        self.tabs.grid(row=1, column=0, sticky="nsew", padx=14, pady=(4, 14))
        try:
            self.tabs._segmented_button.configure(font=("Segoe UI", 12, "bold"))
        except AttributeError:
            pass

        style = ttk.Style()
        style.theme_use("clam")
        style.configure(
            "History.Treeview", background=PALETTE["SURFACE"], foreground=PALETTE["TEXT"],
            fieldbackground=PALETTE["SURFACE"], rowheight=50,
            font=FONT_TABLE_VALUE, borderwidth=0,
        )
        style.layout("History.Treeview", [("History.Treeview.treearea", {"sticky": "nswe"})])
        style.configure(
            "History.Treeview.Heading", background=PALETTE["SURFACE_2"],
            foreground=PALETTE["TITLE"], font=FONT_TABLE_HEADING,
            relief="flat", padding=(10, 9),
        )
        style.map(
            "History.Treeview",
            background=[("selected", "#2B6CB0")], foreground=[("selected", "#FFFFFF")],
        )

        self.trees: dict[str, ttk.Treeview] = {}
        self.empty_labels: dict[str, ctk.CTkLabel] = {}
        keys = tuple(item[0] for item in self.COLUMNS)
        for mode in ("REAL", "PAPER"):
            frame = self.tabs.add(f"CKCS {mode}")
            frame.grid_columnconfigure(0, weight=1)
            frame.grid_rowconfigure(0, weight=1)
            tree = ttk.Treeview(
                frame, columns=keys, show="tree headings", selectmode="browse",
                style="History.Treeview",
            )
            tree.heading("#0", text="NGÀY / GIAO DỊCH / SỰ KIỆN", anchor="w")
            tree.column("#0", width=510, minwidth=360, anchor="w", stretch=True)
            for key, title, width_px, anchor in self.COLUMNS:
                tree.heading(key, text=title, anchor=anchor)
                tree.column(key, width=width_px, minwidth=min(width_px, 100), anchor=anchor, stretch=True)
            tree.tag_configure("day", background="#171B20", foreground="#FFFFFF", font=("Segoe UI", 12, "bold"))
            tree.tag_configure("trade_win", background="#173322", foreground="#EAFBF0", font=FONT_TABLE_VALUE)
            tree.tag_configure("trade_loss", background="#382126", foreground="#FFF1F2", font=FONT_TABLE_VALUE)
            tree.tag_configure("trade_open", background="#3C321B", foreground="#FEF3C7", font=FONT_TABLE_VALUE)
            tree.tag_configure("buy", foreground="#65D991")
            tree.tag_configure("sell", foreground="#FF8A8A")
            tree.tag_configure("cancelled", foreground="#F6C35B")
            tree.tag_configure("pending", foreground="#F6C35B")
            tree.grid(row=0, column=0, sticky="nsew", padx=(5, 0), pady=(5, 0))
            yscroll = ttk.Scrollbar(frame, orient="vertical", command=tree.yview)
            xscroll = ttk.Scrollbar(frame, orient="horizontal", command=tree.xview)
            yscroll.grid(row=0, column=1, sticky="ns", pady=(5, 0))
            xscroll.grid(row=1, column=0, sticky="ew", padx=(5, 0))
            tree.configure(yscrollcommand=yscroll.set, xscrollcommand=xscroll.set)
            self.trees[mode] = tree
            self.empty_labels[mode] = ctk.CTkLabel(
                frame, text="CHƯA CÓ GIAO DỊCH", font=("Segoe UI", 12, "bold"),
                text_color=PALETTE["DIM"], fg_color=PALETTE["SURFACE"],
            )
        self._build_signal_tab()
        self.tabs.set("CKCS REAL" if str(initial_mode).upper() == "REAL" else "CKCS PAPER")
        self.refresh()
        self.show()

    SIGNAL_COLUMNS = (
        ("symbol", "MÃ", 95, "center"),
        ("execution_mode", "CHẾ ĐỘ", 95, "center"),
        ("watchlist_priority", "ƯU TIÊN FA", 125, "center"),
        ("slot_usage", "SLOT", 90, "center"),
        ("display_signal", "SỰ KIỆN", 140, "center"),
        ("suggestion", "XỬ LÝ", 140, "center"),
        ("display_price", "GIÁ", 130, "center"),
        ("ema_comparison", "EMA NHANH / CHẬM", 235, "center"),
        ("ema_cross_display", "CẮT EMA", 225, "center"),
        ("rsi_comparison", "RSI HIỆN TẠI / TRƯỚC", 205, "center"),
        ("rsi_previous_date", "PHIÊN RSI TRƯỚC", 170, "center"),
        ("market_state", "THỊ TRƯỜNG", 210, "center"),
        ("reason", "LÝ DO", 320, "w"),
    )

    def _build_signal_tab(self) -> None:
        """Events and periodic samples share one day -> symbol -> time tree."""
        frame = self.tabs.add("TÍN HIỆU")
        frame.grid_columnconfigure(0, weight=1)
        frame.grid_rowconfigure(1, weight=1)
        toolbar = ctk.CTkFrame(frame, fg_color="transparent")
        toolbar.grid(row=0, column=0, columnspan=2, sticky="ew", padx=8, pady=4)
        toolbar.grid_columnconfigure(0, weight=1)
        self.signal_status = ctk.CTkLabel(toolbar, text="", font=("Segoe UI", 12), anchor="w",
                                               text_color=PALETTE["MUTED"])
        self.signal_status.grid(row=0, column=0, columnspan=3, sticky="w", pady=(0, 5))
        self.recording_button = ctk.CTkButton(
            toolbar, text="⚙ GHI TÍN HIỆU", width=145, command=self._open_recording_settings,
            font=("Segoe UI", 12, "bold"), fg_color=PALETTE["BLUE"], hover_color=PALETTE["BLUE_HOVER"])
        self.recording_button.grid(row=1, column=0, sticky="w")
        _HoverHint(self.recording_button, "Bật/tắt ghi định kỳ, giờ bắt đầu/kết thúc và nhịp ghi.\n"
                   "Không đổi rule mua, BOT hoặc Telegram. Sự kiện vẫn ghi ngay khi trạng thái đổi.")
        self.normalization_button = ctk.CTkButton(
            toolbar, text="CHUẨN HOÁ · OFF", width=165, command=self._toggle_normalization,
            font=("Segoe UI", 12, "bold"), fg_color=PALETTE["SLATE"], hover_color=PALETTE["SLATE_HOVER"])
        self.normalization_button.grid(row=1, column=1, padx=8)
        self.normalization_hint = _HoverHint(self.normalization_button,
            "Bảng luôn giữ số và sự kiện bot đã ghi. ON: mở khung đối chiếu riêng dưới bảng.\n"
            "Chọn một dòng giờ để xem số tính lại từ lịch sử DNSE có sẵn, không cần CSV.\n"
            "Một nến 1D/phiên + giá tick tại giờ ghi; EMA chuẩn, RSI Wilder. Không làm tròn giá đầu vào.\n"
            "Thiếu dữ liệu: giữ số gốc. Chu kỳ cũ thiếu: dùng 3/6/14 và ghi rõ trong Chi tiết.\n"
            "Sự kiện, xử lý và cắt EMA vẫn là những gì bot đã ghi; không tính lại lệnh hay gửi Telegram.\n"
            "Không tự suy đoán hệ số cổ tức hoặc cam kết khớp từng số TradingView.", placement="below")
        _HoverHint(self.signal_status, self.normalization_hint.text + "\n"
                   "Lịch sử lấy từ cache DNSE daemon tự tải; LÀM MỚI đọc lại khi cache cập nhật.\n"
                   "Chi tiết giữ số gốc, thời điểm cache và phạm vi lịch sử dùng tính lại.")
        self.signal_export_button = ctk.CTkButton(toolbar, text="XUẤT EXCEL", width=110, command=self._export_signals,
                                                 font=("Segoe UI", 11, "bold"), fg_color=PALETTE["BLUE"])
        self.signal_export_button.grid(row=1, column=2)
        _HoverHint(self.signal_export_button, "Xuất sự kiện và mẫu định kỳ trong danh sách, kể cả nhóm chưa xổ.\n"
                   "Không xuất dòng đã xóa/ẩn. Muốn xuất riêng: chọn dòng/mã/ngày → chuột phải.\n"
                   "Luôn có DNSE GỐC; bật chuẩn hoá thì thêm CHUẨN HOÁ. Có mẫu thì kèm TRACE/SETTING.", placement="below")
        keys = tuple(item[0] for item in self.SIGNAL_COLUMNS)
        tree = self.signal_tree = ttk.Treeview(
            frame, columns=keys, show="tree headings", selectmode="extended", style="Signal.Treeview")
        style = ttk.Style()
        style.configure("Signal.Treeview", background=PALETTE["SURFACE"], foreground=PALETTE["TEXT"],
                        fieldbackground=PALETTE["SURFACE"], rowheight=50, font=FONT_TABLE_VALUE, borderwidth=0)
        style.configure("Signal.Treeview.Heading", background=PALETTE["SURFACE_2"],
                        foreground=PALETTE["TITLE"], font=FONT_TABLE_HEADING, relief="flat", padding=(10, 9))
        style.map("Signal.Treeview", background=[("selected", "#2B6CB0")], foreground=[("selected", "#FFFFFF")])
        tree.heading("#0", text="NGÀY / MÃ / GIỜ GHI NHẬN", anchor="w")
        tree.column("#0", width=390, minwidth=350, anchor="w", stretch=True)
        for key, title, width_px, anchor in self.SIGNAL_COLUMNS:
            tree.heading(key, text=title, anchor=anchor)
            tree.column(key, width=width_px, minwidth=min(width_px, 85), anchor=anchor, stretch=True)
        for tag, color in (("buy", "#65D991"), ("sell", "#FF8A8A"),
                           ("blocked", "#F6C35B"), ("periodic", PALETTE["MUTED"])):
            tree.tag_configure(tag, foreground=color)
        tree.tag_configure("signal_day", background=PALETTE["SURFACE_2"], foreground=PALETTE["TITLE"],
                           font=("Segoe UI", 12, "bold"))
        tree.grid(row=1, column=0, sticky="nsew", padx=(5, 0), pady=(5, 0))
        ys = ttk.Scrollbar(frame, orient="vertical", command=tree.yview)
        xs = ttk.Scrollbar(frame, orient="horizontal", command=tree.xview)
        ys.grid(row=1, column=1, sticky="ns", pady=(5, 0))
        xs.grid(row=2, column=0, sticky="ew", padx=(5, 0))
        tree.configure(yscrollcommand=ys.set, xscrollcommand=xs.set)
        tree.bind("<Button-3>", self._signal_context_menu)
        tree.bind("<Delete>", lambda _event: self._delete_signal_rows())
        tree.bind("<Double-1>", self._signal_double_click, add="+")
        tree.bind("<<TreeviewSelect>>", lambda _event: self._refresh_normalization_preview(), add="+")
        self.signal_menu = tk.Menu(self.top, tearoff=False)
        self.signal_menu.add_command(label="Chi tiết", command=self._show_signal_details)
        self.signal_menu.add_command(label="Sao chép", command=self._copy_signal_rows)
        self.signal_menu.add_command(label="Xuất Excel phần đã chọn", command=lambda: self._export_signals(selected_only=True))
        self.signal_menu.add_separator()
        self.signal_menu.add_command(label="Xóa khỏi lịch sử", command=self._delete_signal_rows)
        self.signal_menu.add_command(label="Khôi phục các dòng đã xóa", command=self._restore_signal_rows)
        self.normalization_preview = ctk.CTkFrame(frame, fg_color=PALETTE["SURFACE_2"],
                                                 border_color=PALETTE["BORDER"], border_width=1)
        self.normalization_preview.grid(row=3, column=0, columnspan=2, sticky="ew", padx=8, pady=(6, 0))
        self.normalization_preview.grid_columnconfigure(0, weight=1)
        self.normalization_title = ctk.CTkLabel(
            self.normalization_preview, text="ĐỐI CHIẾU · DNSE TÍNH LẠI", anchor="w",
            font=("Segoe UI", 11, "bold"), text_color=PALETTE["TITLE"])
        self.normalization_title.grid(row=0, column=0, sticky="ew", padx=12, pady=(6, 0))
        self.normalization_detail = ctk.CTkLabel(
            self.normalization_preview, text="", anchor="w", justify="left", wraplength=850,
            font=("Segoe UI", 12), text_color=PALETTE["TEXT"])
        self.normalization_detail.grid(row=1, column=0, sticky="ew", padx=12, pady=(0, 6))
        self.normalization_preview.bind("<Configure>", lambda event: self.normalization_detail.configure(
            wraplength=max(120, int(self.normalization_preview._reverse_widget_scaling(event.width)) - 30)), add="+")
        self.normalization_preview.grid_remove()
        ctk.CTkLabel(
            frame, text="ENTRY / MẤT ENTRY / EXIT E: sự kiện · ĐỊNH KỲ: mẫu, không gửi Telegram\n"
                        "7 ngày gần nhất · Chuột phải: chi tiết / xóa / khôi phục · Không xóa lệnh giao dịch",
            font=("Segoe UI", 11), text_color=PALETTE["MUTED"], anchor="w", justify="left",
        ).grid(row=4, column=0, columnspan=2, sticky="w", padx=8, pady=(6, 4))
        self.signal_empty = ctk.CTkLabel(frame, text="CHƯA CÓ DỮ LIỆU TÍN HIỆU",
                                       text_color=PALETTE["DIM"], fg_color=PALETTE["SURFACE"])

    def _open_recording_settings(self) -> None:
        from ..config import AppSettings
        try:
            popup = getattr(self, "recording_popup", None)
            if popup is not None and popup.top.winfo_exists():
                popup.top.lift()
                return
            configured = self.trace_settings_provider() if self.trace_settings_provider else AppSettings()
            self.recording_popup = SignalRecordingPopup(
                self.top, configured, self.on_trace_settings, on_saved=self._refresh_signals)
        except (OSError, ValueError, RuntimeError) as exc:
            messagebox.showerror("Ghi tín hiệu", str(exc), parent=self.top)

    def _history_tab_changed(self) -> None:
        self.subtitle.configure(text=("Mỗi ngày → mã → giờ · Sự kiện và mẫu định kỳ"
                                      if self.tabs.get() == "TÍN HIỆU"
                                      else "Mỗi ngày → mỗi giao dịch → các lần đặt/hủy/khớp"))

    def _history_sources(self) -> list[dict[str, Any]]:
        events = [dict(row) for row in (self.signals_provider() if self.signals_provider else [])]
        for row in events:
            row["_history_ids"] = [observation_id(row)]
        samples = [periodic_history_row(row) for row in self.trace_store.read(limit=100000)]
        return self.history_trash.visible([*events, *samples])

    def _normalization_for(self, row: dict[str, Any]) -> dict[str, Any]:
        """Look up a projection by original IDs, never by recalculated numbers."""
        for identity in row.get("_history_ids") or [observation_id(row)]:
            result = self._normalization_rows.get(identity)
            if result is not None:
                return result
        return {"normalization_error": "Đang tính lại" if self._normalization_future else "Chưa tính lại được"}

    def _history_views(self, rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
        """The main grid, clipboard and journal sheet always show saved evidence."""
        events = [row for row in rows if row.get("record_kind") != "PERIODIC"]
        views = [row for day in signal_rows_by_day(events, compact=True) for row in day["rows"]]
        for sample in (row for row in rows if row.get("record_kind") == "PERIODIC"):
            groups = signal_rows_by_day([sample])
            if groups:
                row = groups[0]["rows"][0]
                row["display_signal"] = "ĐỊNH KỲ"
                row["suggestion"] = ("EMA/RSI đạt" if sample.get("entry") is True else
                                     "Chưa đạt" if sample.get("entry") is False else "Thiếu dữ liệu")
                row["reason"] = sample.get("reason", "")
                views.append(row)
        for row in views:
            row["_reason_detail"] = row["reason"]
            row["suggestion"] = {
                "CHỜ GIỜ MUA": "Chờ giờ", "CHỜ BUY": "Chờ xác nhận", "HỦY CHỜ BUY": "Hủy chờ",
                "HẾT GIỜ MUA": "Hết giờ", "BUY BỊ CHẶN": "Chặn", "ĐÃ XẾP BUY": "Đã xếp",
                "ĐÃ XẾP SELL": "Đã xếp", "CHỈ TÍN HIỆU SELL": "Chỉ tín hiệu", "CHƯA XẾP SELL": "Chưa xếp",
                "MẤT ENTRY": "Chỉ tín hiệu",
            }.get(row["suggestion"], row["suggestion"])
            code = str(row.get("blocked_by") or row.get("reason") or "")
            row["reason"] = short_signal_reason(code, row)
            if row.get("_rsi_issue"):
                row["_reason_detail"] += "\n" + row["_rsi_issue"]
                if not row.get("_legacy_record"):
                    row["reason"] += " · " + row["_rsi_issue"]
        return sorted(views, key=history_sort_key, reverse=True)

    def _refresh_signals(self) -> None:
        tree = getattr(self, "signal_tree", None)
        if tree is None or not tree.winfo_exists():
            return
        opened = {iid: tree.item(iid, "open") for day in tree.get_children()
                  for iid in (day, *tree.get_children(day))}
        selected = tree.selection()
        try:
            sources = self._history_sources()
            days = sorted({str(row.get("timestamp", ""))[:10] for row in sources}, reverse=True)[:7]
            sources = [row for row in sources if str(row.get("timestamp", ""))[:10] in days]
            if self.normalization_enabled:
                self._request_normalization(sources)
            views = self._history_views(sources)
            self._visible_history_sources = sources
            status = "DNSE · Số đã ghi · 7 ngày gần nhất"
            if self.normalization_enabled:
                failures = sum(not self._normalization_for(row).get("normalization_ok") for row in views)
                status = ("DNSE · Bảng giữ số đã ghi · Đang tính đối chiếu" if self._normalization_future
                          else f"DNSE · Bảng giữ số đã ghi · Đối chiếu {len(views) - failures}/{len(views)} dòng")
                if failures and not self._normalization_future:
                    status += f" · {failures} dòng chưa tính lại được"
            self.signal_status.configure(text=status)
        except (OSError, ValueError, KeyError, sqlite3.Error) as exc:
            self.signal_status.configure(text=f"Không đọc được tín hiệu: {type(exc).__name__}")
            return
        tree.delete(*tree.get_children())
        self._history_node_rows = {}
        for index, day in enumerate(days):
            details = [row for row in views if str(row.get("timestamp", ""))[:10] == day]
            if not details:
                continue
            try:
                label = datetime.strptime(day, "%Y-%m-%d").strftime("%d/%m/%Y")
            except ValueError:
                label = day
            parent = tree.insert("", "end", iid=f"day:{day}", text=f"  {label} · {history_counts(details)}",
                                 open=opened.get(f"day:{day}", index == 0), tags=("signal_day",))
            self._history_node_rows[parent] = details
            for symbol in sorted({row["symbol"] for row in details}):
                symbol_rows = [row for row in details if row["symbol"] == symbol]
                iid = f"symbol:{day}:{symbol}"
                tree.insert(parent, "end", iid=iid, text=f"  {symbol} · {history_counts(symbol_rows)}",
                            open=opened.get(iid, False), tags=("signal_day",))
                self._history_node_rows[iid] = symbol_rows
                for row in symbol_rows:
                    tag = ("periodic" if row.get("record_kind") == "PERIODIC" else
                           "blocked" if row.get("blocked_by") or row.get("acted") == "WAIT" else
                           "buy" if row.get("signal") == "BUY" else "sell")
                    clock = str(row.get("time", ""))
                    if row.get("legacy_repeat_count", 1) > 1:
                        clock += f" (×{row['legacy_repeat_count']} cũ)"
                    leaf = tree.insert(iid, "end", iid="row:" + row["_history_ids"][0],
                                       text=f"    {clock}", tags=(tag,),
                                       values=tuple(row.get(key, "") for key, *_ in self.SIGNAL_COLUMNS))
                    self._history_node_rows[leaf] = [row]
        tree.selection_set([iid for iid in selected if tree.exists(iid)])
        self._refresh_normalization_preview()
        if tree.get_children():
            self.signal_empty.place_forget()
        else:
            self.signal_empty.place(relx=0.5, rely=0.5, anchor="center")

    def _refresh_normalization_preview(self) -> None:
        if not self.normalization_enabled:
            self.normalization_preview.grid_remove()
            return
        self.normalization_preview.grid()
        rows = self._selected_signal_rows()
        self.normalization_title.configure(text="ĐỐI CHIẾU · DNSE TÍNH LẠI")
        if len(rows) != 1:
            self.normalization_detail.configure(text="Chọn một dòng giờ để đối chiếu. Bảng bên trên luôn giữ số đã ghi.")
            return
        row = rows[0]
        result = self._normalization_for(row)
        self.normalization_title.configure(
            text=f"ĐỐI CHIẾU · {row['symbol']} · {row.get('execution_mode') or '—'} · {row.get('time') or '—'}")
        text = f"Đã ghi: EMA {row['ema_comparison']} · RSI {row['rsi_comparison']}"
        if result.get("normalization_ok"):
            text += ("\nTính lại DNSE: EMA " + number_comparison(result.get("ema_fast"), result.get("ema_slow"), 4)
                     + " · RSI " + rsi_observation_display(result.get("rsi"), result.get("rsi_previous"))[0]
                     + " · Phiên trước " + str(result.get("rsi_previous_date") or "—"))
            text += "\nChỉ đối chiếu · Không đổi sự kiện, cắt EMA hoặc lệnh."
        else:
            text += "\nTính lại DNSE: " + str(result.get("normalization_error") or "Chưa tính lại được")
        self.normalization_detail.configure(text=text)

    def _selected_signal_rows(self) -> list[dict[str, Any]]:
        result, seen = [], set()
        for iid in self.signal_tree.selection():
            for row in self._history_node_rows.get(iid, []):
                identity = tuple(row.get("_history_ids") or [observation_id(row)])
                if identity not in seen:
                    seen.add(identity)
                    result.append(row)
        return result

    def _signal_context_menu(self, event) -> str:
        iid = self.signal_tree.identify_row(event.y)
        if iid:
            if iid not in self.signal_tree.selection():
                self.signal_tree.selection_set(iid)
            self.signal_tree.focus(iid)
        else:
            self.signal_tree.selection_remove(*self.signal_tree.selection())
        enabled = "normal" if self._selected_signal_rows() else "disabled"
        for label in ("Chi tiết", "Sao chép", "Xuất Excel phần đã chọn", "Xóa khỏi lịch sử"):
            self.signal_menu.entryconfigure(label, state=enabled)
        self.signal_menu.entryconfigure("Khôi phục các dòng đã xóa",
                                       state="normal" if self.history_trash.hidden_ids() else "disabled")
        try:
            self.signal_menu.tk_popup(event.x_root, event.y_root)
        finally:
            self.signal_menu.grab_release()
        return "break"

    def _signal_double_click(self, event) -> None:
        iid = self.signal_tree.identify_row(event.y)
        if iid and not self.signal_tree.get_children(iid):
            self.signal_tree.selection_set(iid)
            self._show_signal_details()

    def _copy_signal_rows(self) -> None:
        rows = self._selected_signal_rows()
        if rows:
            self.top.clipboard_clear()
            self.top.clipboard_append("\n".join(
                " · ".join(str(row.get(key, "—")) for key in
                           ("timestamp", "symbol", "execution_mode", "display_signal", "suggestion",
                            "display_price", "ema_comparison", "rsi_comparison", "reason")) for row in rows))

    def _show_signal_details(self) -> None:
        rows = self._selected_signal_rows()
        if not rows:
            return
        row = rows[0]
        text = "\n".join([f"{row.get('timestamp')} · {row.get('symbol')} · {row.get('execution_mode', '—')}",
            f"{row.get('display_signal')} · {row.get('suggestion')}",
            f"Giá {row.get('display_price')} · EMA {row.get('ema_comparison')}",
            f"RSI {row.get('rsi_comparison')} · Phiên trước {row.get('rsi_previous_date') or '—'}",
            f"Cắt EMA: {row.get('ema_cross_display')}",
            f"Lý do: {row.get('_reason_detail')}"])
        previous_ema = number_comparison(row.get("ema_previous_fast"), row.get("ema_previous_slow"), 4)
        text += "\nEMA lần quan sát trước: " + ("Bản cũ chưa lưu" if row.get("_legacy_record") and previous_ema == "—" else previous_ema)
        if row.get("_legacy_record") and (row.get("_rsi_issue") or not row.get("ema_cross_state")):
            text += ("\nBản cũ chưa lưu đủ trường đối chiếu. Không đồng nghĩa bot thiếu dữ liệu lúc chạy."
                     "\nKhông suy ra lần cắt EMA từ một cặp EMA hoặc điền RSI tính lại thành số đã ghi.")
        if self.normalization_enabled:
            result = self._normalization_for(row)
            text += "\n\nĐỐI CHIẾU · DNSE TÍNH LẠI (không phải số đã ghi)"
            if result.get("normalization_ok"):
                text += "\nEMA " + number_comparison(result.get("ema_fast"), result.get("ema_slow"), 4)
                text += " · RSI " + rsi_observation_display(result.get("rsi"), result.get("rsi_previous"))[0]
                text += "\nPhiên RSI trước tính lại: " + str(result.get("rsi_previous_date") or "—")
                periods = result["normalization_periods"]
                text += (f"\nTính lại: EMA {periods[0]}/{periods[1]}, RSI {periods[2]} · "
                         f"{result['normalization_bars']} nến trước phiên"
                         f"\nLịch sử: {result['normalization_history_start']} → {result['normalization_history_end']}"
                         f"\nCache DNSE: {result.get('normalization_cache_at') or 'không có mốc lưu'}"
                         "\nDùng giá tại giờ ghi, không lấy giá đóng cửa ngày đó."
                         "\nLịch sử hiện có có thể đã được cập nhật; đây không phải quyết định mới hay số TradingView.")
                if result.get("normalization_defaults"):
                    text += "\nBản cũ thiếu chu kỳ: dùng mặc định cho " + ", ".join(
                        {"ema_fast_period": "EMA nhanh", "ema_slow_period": "EMA chậm", "rsi_period": "RSI"}.get(key, key)
                        for key in result["normalization_defaults"])
                if row.get("record_kind") == "PERIODIC":
                    text += "\nEMA/RSI tính lại: " + ("Đạt" if result.get("normalization_condition") else "Chưa đạt")
            else:
                text += "\n" + str(result.get("normalization_error") or "Chưa tính lại được")
        if row.get("record_kind") == "PERIODIC":
            text += "\n" + "\n".join(f"{label}: {row.get(key, '—')}" for key, label in (
                ("scheduled_at", "Mốc lấy mẫu"), ("price_source", "Nguồn giá"),
                ("quote_age_seconds", "Tuổi giá (giây)"), ("quote_issue", "Lỗi giá"),
                ("system_issue", "Lỗi hệ thống"), ("exit_e", "EXIT E"),
                ("whipsaw_count", "WHIPSAW đếm"), ("whipsaw_limit", "WHIPSAW ngưỡng"),
                ("slot_usage", "Slot"), ("order_budget", "Vốn/lệnh"),
                ("bot_on", "BOT ON"), ("otp_ok", "OTP OK"), ("queue_summary", "Lệnh đã ghi")))
        messagebox.showinfo("Chi tiết tín hiệu", text, parent=self.top)

    def _delete_signal_rows(self) -> str:
        rows = self._selected_signal_rows()
        if rows and messagebox.askyesno(
            "Xóa lịch sử tín hiệu",
            f"Xóa {len(rows)} dòng đã chọn khỏi danh sách?\nCó thể khôi phục bằng chuột phải. "
            "Không xóa log gốc hay lệnh giao dịch.", parent=self.top):
            try:
                self.history_trash.hide(rows)
                self._refresh_signals()
            except (OSError, ValueError) as exc:
                messagebox.showerror("Xóa tín hiệu", str(exc), parent=self.top)
        return "break"

    def _restore_signal_rows(self) -> None:
        if not self.history_trash.hidden_ids():
            return
        if messagebox.askyesno("Khôi phục tín hiệu", "Khôi phục tất cả dòng đã xóa khỏi danh sách?\n"
                              "Không gửi lại Telegram hay tạo lệnh.", parent=self.top):
            try:
                self.history_trash.restore()
                self._refresh_signals()
            except (OSError, ValueError) as exc:
                messagebox.showerror("Khôi phục tín hiệu", str(exc), parent=self.top)

    def _toggle_normalization(self) -> None:
        self.normalization_enabled = not self.normalization_enabled
        if not self.normalization_enabled:
            self._stop_normalization()
        self.normalization_button.configure(
            text="CHUẨN HOÁ · ON" if self.normalization_enabled else "CHUẨN HOÁ · OFF",
            fg_color=PALETTE["BLUE"] if self.normalization_enabled else PALETTE["SLATE"],
            hover_color=PALETTE["BLUE_HOVER"] if self.normalization_enabled else PALETTE["SLATE_HOVER"])
        self._refresh_signals()

    def _request_normalization(self, sources: list[dict[str, Any]]) -> None:
        """Read/calculate off Tk; a worker never calls Tk, the broker or Telegram."""
        if self._normalization_future is not None:
            return
        key = (self.indicator_normalizer.signature(), tuple(
            (observation_id(row), *(str(row.get(field, "")) for field in (
                "ema_fast_period", "ema_slow_period", "rsi_period", "price_vnd"))) for row in sources))
        if key == self._normalization_request_key:
            return
        self._normalization_request_key = key
        self._normalization_rows = {}
        if not sources:
            return
        if self._normalization_executor is None:
            self._normalization_executor = ThreadPoolExecutor(max_workers=1, thread_name_prefix="dnse-display-normalize")
        inputs = [dict(row) for row in sources]
        self._normalization_future = self._normalization_executor.submit(self.indicator_normalizer.normalize, inputs)
        generation = self._normalization_generation
        self._normalization_after = self.top.after(30, lambda: self._poll_normalization(generation, inputs))

    def _poll_normalization(self, generation: int, inputs: list[dict[str, Any]]) -> None:
        if generation != self._normalization_generation or not self.normalization_enabled or not self.top.winfo_exists():
            return
        self._normalization_after = None
        future = self._normalization_future
        if future is None:
            return
        if not future.done():
            self._normalization_after = self.top.after(30, lambda: self._poll_normalization(generation, inputs))
            return
        try:
            results = future.result()
        except Exception as exc:
            results = [{**row, "normalization_error": f"Chưa tính lại được ({type(exc).__name__}); giữ số gốc"} for row in inputs]
        self._normalization_rows = {observation_id(raw): calculated for raw, calculated in zip(inputs, results)}
        self._normalization_future = None
        self._refresh_signals()

    def _stop_normalization(self) -> None:
        self._normalization_generation += 1
        if self._normalization_after is not None:
            try:
                self.top.after_cancel(self._normalization_after)
            except tk.TclError:
                pass
            self._normalization_after = None
        if self._normalization_future is not None:
            self._normalization_future.cancel()
            self._normalization_future = None
        self._normalization_request_key = None
        self._normalization_rows = {}

    def _export_signals(self, *, selected_only: bool = False) -> None:
        if self.normalization_enabled and self._normalization_future is not None:
            messagebox.showinfo("Xuất Excel", "Đang chuẩn hoá. Chờ tính xong rồi xuất Excel.", parent=self.top)
            return
        original = list(getattr(self, "_visible_history_sources", []))
        if selected_only:
            ids = {identity for row in self._selected_signal_rows() for identity in row.get("_history_ids", [])}
            original = [row for row in original if observation_id(row) in ids]
        if not original:
            messagebox.showinfo("Xuất Excel", "Không có dòng để xuất.", parent=self.top)
            return
        path = filedialog.asksaveasfilename(parent=self.top, title="Xuất tín hiệu", defaultextension=".xlsx",
                                          filetypes=[("Excel", "*.xlsx")], initialfile="TIN_HIEU.xlsx")
        if not path:
            return
        book = None
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font
            captures = [row for row in original if row.get("record_kind") == "PERIODIC"]
            book = trace_workbook(captures) if captures else Workbook()
            if not captures:
                book.remove(book.active)
            views = self._history_views(original)
            for index, normalized in enumerate((False, True) if self.normalization_enabled else (False,)):
                sheet = book.create_sheet("CHUẨN HOÁ" if normalized else "DNSE GỐC", index)
                metadata = ["NGUỒN", "CHI TIẾT", "LỖI CHUẨN HOÁ", "CACHE DNSE", "SỐ NẾN", "CHU KỲ MẶC ĐỊNH",
                            "LỊCH SỬ TỪ", "LỊCH SỬ ĐẾN", "CHU KỲ TÍNH", "EMA/RSI TÍNH LẠI"]
                columns = (["MÃ", "CHẾ ĐỘ", "SỰ KIỆN ĐÃ GHI", "XỬ LÝ ĐÃ GHI", "GIÁ ĐÃ GHI",
                            "EMA ĐÃ GHI", "RSI ĐÃ GHI", "PHIÊN RSI TRƯỚC ĐÃ GHI", "CẮT EMA ĐÃ GHI",
                            "EMA TÍNH LẠI", "RSI TÍNH LẠI", "PHIÊN RSI TRƯỚC TÍNH LẠI"] if normalized
                           else [title for _, title, *_ in self.SIGNAL_COLUMNS])
                sheet.append(["GIỜ GHI NHẬN", *columns, *metadata])
                for cell in sheet[1]:
                    cell.font = Font(bold=True)
                for row in views:
                    result = self._normalization_for(row) if normalized else {}
                    ok = bool(result.get("normalization_ok"))
                    if normalized:
                        values = [row.get(key, "") for key in (
                            "symbol", "execution_mode", "display_signal", "suggestion", "display_price",
                            "ema_comparison", "rsi_comparison", "rsi_previous_date", "ema_cross_display")]
                        values.extend([
                            number_comparison(result.get("ema_fast"), result.get("ema_slow"), 4) if ok else "Chưa tính lại",
                            rsi_observation_display(result.get("rsi"), result.get("rsi_previous"))[0] if ok else "Chưa tính lại",
                            result.get("rsi_previous_date", "") if ok else ""])
                    else:
                        values = [row.get(key, "") for key, *_ in self.SIGNAL_COLUMNS]
                    detail = ("Chỉ đối chiếu; không đổi sự kiện/xử lý/cắt EMA đã ghi. Lịch sử DNSE hiện có, không phải số TradingView."
                              if normalized else row.get("_reason_detail", ""))
                    sheet.append([row["timestamp"], *values,
                                  "DNSE · Tính lại" if normalized else "DNSE · Đã ghi",
                                  detail, result.get("normalization_error", ""),
                                  result.get("normalization_cache_at", ""), result.get("normalization_bars", ""),
                                  ", ".join({"ema_fast_period": "EMA nhanh", "ema_slow_period": "EMA chậm", "rsi_period": "RSI"}.get(key, key)
                                            for key in result.get("normalization_defaults") or []),
                                  result.get("normalization_history_start", ""), result.get("normalization_history_end", ""),
                                  "/".join(str(value) for value in result.get("normalization_periods") or []),
                                  result.get("normalization_condition", "") if ok else ""])
                sheet.freeze_panes = "A2"
                sheet.auto_filter.ref = sheet.dimensions
                for sheet_row in sheet.iter_rows():
                    for cell in sheet_row:
                        if isinstance(cell.value, str) and cell.value.startswith(("=", "+", "-", "@")):
                            cell.data_type = "s"
                for column in sheet.columns:
                    sheet.column_dimensions[column[0].column_letter].width = min(
                        65, max(16, max(len(str(cell.value or "")) for cell in column) + 2))
            book.save(path)
        except (OSError, ValueError, ImportError) as exc:
            messagebox.showerror("Xuất Excel", str(exc), parent=self.top)
        finally:
            if book is not None:
                book.close()
    def show(self) -> None:
        self.top.deiconify()
        self.top.lift()
        self.top.focus_force()
        if self.on_visibility_changed:
            self.on_visibility_changed(True)

    def hide(self) -> None:
        if self.top.winfo_exists():
            self.top.withdraw()

        if self.on_visibility_changed:
            self.on_visibility_changed(False)

    def close(self) -> None:
        self._stop_normalization()
        if self._normalization_executor is not None:
            self._normalization_executor.shutdown(wait=False, cancel_futures=True)
            self._normalization_executor = None
        if self.on_visibility_changed:
            self.on_visibility_changed(False)
        if self.top.winfo_exists():
            self.top.destroy()

    def refresh(self) -> None:
        self._refresh_signals()
        for mode, tree in self.trees.items():
            tree.delete(*tree.get_children())
            groups = self.rows_provider(mode)
            for day in groups:
                day_id = tree.insert(
                    "", "end", text=str(day.get("label", "")),
                    values=tuple(day.get("values", ())), tags=("day",), open=True,
                )
                for trade in day.get("trades", []):
                    trade_id = tree.insert(
                        day_id, "end", text=str(trade.get("label", "")),
                        values=tuple(trade.get("values", ())),
                        tags=(str(trade.get("tag", "trade_open")),), open=True,
                    )
                    for event in trade.get("events", []):
                        tree.insert(
                            trade_id, "end", text=str(event.get("label", "")),
                            values=tuple(event.get("values", ())),
                            tags=(str(event.get("tag", "pending")),),
                        )
            empty = self.empty_labels[mode]
            if groups:
                empty.place_forget()
            else:
                empty.place(relx=0.5, rely=0.5, anchor="center")
