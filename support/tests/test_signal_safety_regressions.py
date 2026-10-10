"""Regression coverage for trading signals, delivery and durable audit writes.

Run through support/tools/run_offline.py to disable credentials and networking.
"""
from copy import deepcopy
from datetime import datetime, timedelta
from types import SimpleNamespace
from unittest.mock import Mock

import pytest

from viking_v2.config import AppSettings
from viking_v2.dashboard.actions import DashboardActionsMixin
from viking_v2.rules.business import StaticRule, StaticRuleParameters, indicator_snapshot
from viking_v2.rules.business import ema, rsi, crossover_count
from viking_v2.rules.entry_filters import apply_buy_filters
from viking_v2.rules.state import RuleStateStore
from viking_v2.services.daemon import indicator_snapshot_at_close
from viking_v2.services.indicator_comparison import DNSEIndicatorNormalizer
from viking_v2.storage import SignalLog
from viking_v2.trading.market import VN_TZ


def test_reference_ema_and_wilder_rsi_arithmetic():
    prices = [10, 11, 10, 12, 11]
    assert ema(prices, 3) == pytest.approx([10, 10.5, 10.25, 11.125, 11.0625])
    values = rsi(prices, 3)
    assert values[:3] == [None, None, None]
    assert values[3:] == pytest.approx([75, 600 / 11])


def test_whipsaw_daily_counter_and_threshold_control():
    bars = [dict(close=price, closed=True) for price in [10, 9, 11, 8, 12]]
    assert crossover_count(bars, 1, 2, 4) == 4
    source = dict(symbol='MSN', bars=bars, signal_mode='REALTIME',
                  confirmed_market_state='UPTREND')
    p = dict(available_capital=100_000_000)
    for enabled, action in [(True, 'WAIT'), (False, 'BUY')]:
        rule = StaticRule(StaticRuleParameters(buy_ema_fast=1, buy_ema_slow=2,
            buy_signal_session_cross_enabled=False,
            buy_signal_use_rsi=False, whipsaw_enabled=enabled, whipsaw_n=3, whipsaw_x=4))
        result = rule.evaluate(source, p)
        assert result.action == action
        if enabled:
            assert result.reason == 'WHIPSAW_LOCK'


def test_normalization_reference_units_dates_and_original_evidence(tmp_path):
    import json
    cache = tmp_path / 'history.json'
    bars = [dict(time=int(datetime(2026, 10, day, 9, tzinfo=VN_TZ).timestamp()),
                 close=price * 1000) for day, price in zip(range(5, 9), [10, 11, 10, 12])]
    bars.append(dict(time=int(at('14:45:00').timestamp()), close=999_000))
    cache.write_text(json.dumps(dict(symbols={'MSN': list(reversed(bars))},
        price_unit='VND', resolution='1D')), encoding='utf-8')
    original = dict(timestamp='2026-10-09 14:00:00', symbol='MSN', price=11,
        ema_fast=44, ema_slow=55, rsi=66, rsi_previous=77,
        ema_fast_period=2, ema_slow_period=3, rsi_period=3, signal='BUY', acted='WAIT')
    before, bytes_before = deepcopy(original), cache.read_bytes()
    result = DNSEIndicatorNormalizer(cache).normalize([original])[0]
    assert result['normalization_ok']
    assert result['ema_fast'] == pytest.approx(11.135802469135802)
    assert result['ema_slow'] == pytest.approx(11.0625)
    assert result['rsi'] == pytest.approx(600 / 11)
    assert result['rsi_previous'] == pytest.approx(75)
    assert result['rsi_previous_date'] == '2026-10-08'
    assert original == before and cache.read_bytes() == bytes_before


def at(clock):
    return datetime.fromisoformat('2026-10-09T' + clock + '+07:00')


def marks(fast=101.0, slow=100.0, rsi=60.0):
    return dict(buy_ema_fast=fast, buy_ema_slow=slow, sell_ema_fast=fast,
                sell_ema_slow=slow, rsi=rsi, rsi_previous=55.0,
                sample_count=30, buy_ema_fast_period=3, buy_ema_slow_period=6,
                sell_ema_fast_period=3, sell_ema_slow_period=6, rsi_period=14)


def context(current=None, previous=None, clock='14:00:00'):
    return dict(symbol='MSN', bars=[dict(close=100, closed=True)] * 30,
                signal_mode='REALTIME', exchange='HOSE',
                confirmed_market_state='UPTREND', observation_time=at(clock).isoformat(),
                indicator_snapshot=current or marks(),
                previous_indicators=previous or marks(fast=99))


def portfolio():
    return dict(available_capital=100_000_000, scale_in_allowed=True,
                entry_orders_available=True, entry_orders_max=2,
                position=dict(quantity=100, avg_price=100, current_price=100,
                              managed_by_bot=True, sl_enabled=True))


@pytest.mark.parametrize('filter_kind', ['window', 'confirmation'])
@pytest.mark.parametrize('exit_event', ['STOP_LOSS', 'TAKE_PROFIT', 'INDICATOR_EXIT'])
def test_buy_filters_must_preserve_exits_during_scale_in_wait(filter_kind, exit_event):
    params = StaticRuleParameters(whipsaw_enabled=False,
        buy_signal_session_cross_enabled=False,
        indicator_exit_policy='AUTO',
        buy_window_enabled=filter_kind == 'window',
        buy_confirmation_enabled=filter_kind == 'confirmation',
        buy_confirmation_minutes=5)
    rule, p = StaticRule(params), portfolio()
    first_clock = '13:59:00' if filter_kind == 'window' else '14:00:00'
    source = context(clock=first_clock)
    raw = rule.evaluate(source, p)
    assert raw.action == 'BUY'
    state, waiting = apply_buy_filters(rule, raw, source, p, {},
        observed_at=at(first_clock), exchange='HOSE')
    assert waiting.action == 'WAIT' and state
    later = at(first_clock) + timedelta(seconds=1)
    p = deepcopy(p)
    if exit_event == 'STOP_LOSS':
        p['position']['current_price'] = 96
    elif exit_event == 'TAKE_PROFIT':
        p['position'].update(current_price=110, tp_mode='PRICE', tp_value=108)
    else:
        p['position']['em_modes'] = ['IND_EXIT']
        source = context(current=marks(fast=99, rsi=50), previous=marks(fast=101))
    raw_exit = rule.evaluate(source, p)
    assert raw_exit.action == 'SELL' and raw_exit.event == exit_event
    _, filtered = apply_buy_filters(rule, raw_exit, source, p, state,
        observed_at=later, exchange='HOSE')
    assert (filtered.action, filtered.event) == ('SELL', exit_event), filtered.to_dict()


@pytest.mark.parametrize('filter_kind', ['window', 'confirmation'])
def test_valid_scale_in_candidate_must_release_without_a_second_cross(filter_kind):
    rule = StaticRule(StaticRuleParameters(whipsaw_enabled=False,
        buy_signal_session_cross_enabled=False,
        buy_window_enabled=filter_kind == 'window',
        buy_confirmation_enabled=filter_kind == 'confirmation',
        buy_confirmation_minutes=5))
    first_clock = '13:59:00' if filter_kind == 'window' else '14:00:00'
    end_clock = '14:00:00' if filter_kind == 'window' else '14:05:00'
    p, source = portfolio(), context(clock=first_clock)
    state, initial = apply_buy_filters(rule, rule.evaluate(source, p), source, p, {},
        observed_at=at(first_clock), exchange='HOSE')
    assert initial.reason in {'BUY_WINDOW_WAIT', 'BUY_CONFIRMATION_WAIT'}
    source = context(previous=marks(), clock=end_clock)
    raw = rule.evaluate(source, p)
    assert raw.reason == 'HOLD_POSITION' and raw.signal == ''
    _, final = apply_buy_filters(rule, raw, source, p, state,
        observed_at=at(end_clock), exchange='HOSE')
    assert final.action == 'BUY', final.to_dict()


@pytest.mark.parametrize('mode', ['REAL', 'PAPER'])
@pytest.mark.parametrize('interval', ['1M', '2M', '5M'])
@pytest.mark.parametrize('use_ema', [True, False])
def test_minute_cold_start_must_not_buy_on_yesterdays_bullish_snapshot(tmp_path, mode, interval, use_ema):
    history = [100] * 15 + [99, 98, 97, 98, 100]
    closed = [dict(close=value, closed=True) for value in history]
    live = [*closed, dict(close=90, closed=False)]
    params = StaticRuleParameters(whipsaw_enabled=False,
        buy_signal_require_ema_cross=False, buy_signal_use_ema=use_ema, buy_window_enabled=False)
    baseline = indicator_snapshot_at_close(closed, 100, params)
    store = RuleStateStore(tmp_path / 'bucket.json')
    result = store.observe_indicator_bucket('MSN', mode, '2026-10-09', interval,
        1000, 90, baseline, lambda close: indicator_snapshot_at_close(live, close, params))
    rule = StaticRule(params)
    fresh_tick_decision = rule.evaluate(dict(symbol='MSN', bars=live,
        signal_mode='REALTIME', indicator_snapshot=indicator_snapshot(live),
        previous_indicators={}, confirmed_market_state='UPTREND'),
        dict(available_capital=100_000_000))
    assert fresh_tick_decision.action == 'WAIT'
    actual = rule.evaluate(dict(symbol='MSN', bars=live, signal_mode='REALTIME',
        indicator_snapshot=result['current'], previous_indicators=result['previous'],
        confirmed_market_state='UPTREND'), dict(available_capital=100_000_000))
    assert actual.action == 'WAIT', actual.to_dict()


def test_minute_feed_recovery_must_not_buy_on_an_old_pending_close(tmp_path):
    closed = [dict(close=value, closed=True) for value in
              [100] * 20 + [99, 98, 97]]
    params = StaticRuleParameters(whipsaw_enabled=False, buy_window_enabled=False)
    baseline = indicator_snapshot_at_close(closed, 97, params)
    store = RuleStateStore(tmp_path / 'feed-gap.json')
    # 13:59: last observed price = 101; feed then stops until 14:10.
    before_gap = [*closed, dict(close=101, closed=False)]
    store.observe_indicator_bucket('MSN', 'REAL', '2026-10-09', '1M',
        1000, 101, baseline,
        lambda price: indicator_snapshot_at_close(before_gap, price, params))
    # 14:10: the fresh quote is bearish at 90. The pending bucket is 11 minutes old.
    after_gap = [*closed, dict(close=90, closed=False)]
    resumed = store.observe_indicator_bucket('MSN', 'REAL', '2026-10-09', '1M',
        1011, 90, baseline,
        lambda price: indicator_snapshot_at_close(after_gap, price, params))
    rule = StaticRule(params)
    actual = rule.evaluate(dict(symbol='MSN', bars=after_gap, signal_mode='REALTIME',
        indicator_snapshot=resumed['current'], previous_indicators=resumed['previous'],
        confirmed_market_state='UPTREND', observation_time=at('14:10:00').isoformat()),
        dict(available_capital=100_000_000))
    assert actual.action == 'WAIT', actual.to_dict()


def test_failed_technical_telegram_delivery_can_retry_same_candidate(tmp_path, monkeypatch):
    clock = [1000.0]
    monkeypatch.setattr('viking_v2.rules.state.time.time', lambda: clock[0])
    class ImmediateThread:
        def __init__(self, *, target, kwargs, daemon):
            self.target, self.kwargs = target, kwargs
        def start(self):
            self.target(**self.kwargs)
    monkeypatch.setattr('viking_v2.dashboard.actions.threading.Thread', ImmediateThread)
    app = DashboardActionsMixin()
    app.settings = AppSettings(telegram_enabled=True).normalize()
    app.settings.telegram_cooldown_minutes['blocked_buy'] = 0
    app.rule_state = RuleStateStore(tmp_path / 'telegram.json')
    service = SimpleNamespace(notify_technical_buy=Mock(side_effect=[False, True]))
    source = context()
    decision = StaticRule(StaticRuleParameters(whipsaw_enabled=False)).evaluate(source,
        dict(available_capital=100_000_000))
    decision.details['signal_cycle'] = 'same-candidate'
    for moment in (1000.0, 1060.0, 1061.0):
        clock[0] = moment  # A failed delivery waits 60 seconds before retry.
        app._notify_signal_only(service, 'MSN', 'BUY', decision, 100, 'REAL')
    assert service.notify_technical_buy.call_count == 2


def test_minute_cross_timestamp_must_not_move_on_unchanged_bucket(tmp_path):
    store = RuleStateStore(tmp_path / 'cross.json')
    base = marks(fast=99)
    store.observe_indicator_bucket('MSN', 'REAL', '2026-10-09', '1M', 1000, 101,
        base, lambda price: marks(fast=price))
    first = store.observe_indicator_bucket('MSN', 'REAL', '2026-10-09', '1M', 1001,
        102, base, lambda price: marks(fast=price))
    repeated = store.observe_indicator_bucket('MSN', 'REAL', '2026-10-09', '1M', 1001,
        103, base, lambda price: marks(fast=price))
    assert first['advanced'] and not repeated['advanced']
    rule = StaticRule(StaticRuleParameters(whipsaw_enabled=False, buy_signal_session_cross_enabled=False))
    values = []
    for result, clock in [(first, '14:01:00'), (repeated, '14:01:55')]:
        decision = rule.evaluate(context(result['current'], result['previous'], clock),
            dict(available_capital=100_000_000))
        values.append(decision.details['ema_cross']['cross_at'])
    assert values[0] == values[1], values


def test_signal_csv_must_not_duplicate_after_crash_between_append_and_dedup_state(tmp_path, monkeypatch):
    path = tmp_path / 'signals.csv'
    log = SignalLog(path)
    row = dict(timestamp='2026-10-09 14:00:00', execution_mode='REAL', symbol='MSN',
        signal='BUY', signal_cycle='C1', acted='BUY', record_kind='SIGNAL_EVENT')
    monkeypatch.setattr(log.state, 'write', Mock(side_effect=OSError('offline disk failure')))
    with pytest.raises(OSError):
        log.observe(row, entry_condition=True, exit_condition=False)
    assert len(log.read_all()) == 1
    restarted = SignalLog(path)
    restarted.observe(row, entry_condition=True, exit_condition=False)
    assert len(restarted.read_all()) == 1


@pytest.mark.parametrize('event', ['INDICATOR_EXIT', 'INDICATOR_EXIT_ALERT', 'NORMAL_ARMED', 'PROTECT_ALERT'])
def test_confirmed_scale_in_preserves_position_management(event):
    params = StaticRuleParameters(whipsaw_enabled=False, buy_window_enabled=False,
        buy_confirmation_enabled=True, indicator_exit_policy='ALERT' if event.endswith('ALERT') else 'AUTO',
        normal_policy='ALERT')
    rule, p, source = StaticRule(params), portfolio(), context()
    if event.startswith('INDICATOR_EXIT'):
        p['position']['em_modes'] = ['IND_EXIT']
        source = context(current=marks(fast=99, rsi=50), previous=marks(fast=101))
    else:
        p['position'].update(em_modes=['NORMAL'], current_price=103.5, peak_profit_pct=7.2,
                             normal_armed=event != 'NORMAL_ARMED')
    raw = rule.evaluate(source, p)
    confirmed = rule.evaluate({**source, 'confirmed_buy': True}, p)
    assert raw.event == confirmed.event == event
    _, filtered = apply_buy_filters(rule, confirmed, source, p,
        {'confirmation': {'active': True}}, observed_at=at('14:01:00'), exchange='HOSE')
    assert filtered.to_dict() == confirmed.to_dict()


@pytest.mark.parametrize('updates,reason', [
    ({'entry_orders_available': False}, 'MAX_SYMBOL_ORDERS'),
    ({'loss_streak': 3}, 'LOCKED_AFTER_LOSSES'),
    ({'entry_slot_available': False}, 'MAX_POSITIONS'),
    ({'available_capital': 0}, 'NO_AVAILABLE_CAPITAL'),
    ({'corporate_action_blocked': True}, 'CORPORATE_ACTION_BLOCK'),
])
def test_scale_in_release_rechecks_guards(updates, reason):
    rule = StaticRule(StaticRuleParameters(whipsaw_enabled=False, buy_window_enabled=True,
                                         buy_signal_session_cross_enabled=False))
    p, source = portfolio(), context(clock='13:59:00')
    state, waiting = apply_buy_filters(rule, rule.evaluate(source, p), source, p, {},
        observed_at=at('13:59:00'), exchange='HOSE')
    assert waiting.reason == 'BUY_WINDOW_WAIT'
    p.update(updates)
    source = context(previous=marks(), clock='14:00:00')
    _, released = apply_buy_filters(rule, rule.evaluate(source, p), source, p, state,
        observed_at=at('14:00:00'), exchange='HOSE')
    assert (released.action, released.reason) == ('WAIT', reason)


@pytest.mark.parametrize('gap_kind', ['elapsed', 'skipped_bucket', 'rejected_quote'])
def test_minute_recovery_waits_for_a_new_close_then_resumes(tmp_path, monkeypatch, gap_kind):
    clock = [1000.0]
    monkeypatch.setattr('viking_v2.rules.state.time.time', lambda: clock[0])
    store = RuleStateStore(tmp_path / 'restart.json')
    baseline = marks(fast=99)
    build = lambda price: marks(fast=price)
    store.observe_indicator_bucket('MSN', 'REAL', '2026-10-09', '5M', 100, 101, baseline, build)
    ready = store.observe_indicator_bucket('MSN', 'REAL', '2026-10-09', '5M', 101, 102, baseline, build)
    assert ready['current']['signal_ready'] is True
    if gap_kind == 'elapsed':
        clock[0] += 120
    elif gap_kind == 'rejected_quote':
        store.reset_indicator_bucket('MSN', 'REAL')
    bucket = 105 if gap_kind == 'skipped_bucket' else 101
    store = RuleStateStore(store.store.path)
    resumed = store.observe_indicator_bucket('MSN', 'REAL', '2026-10-09', '5M', bucket, 90, baseline, build)
    assert resumed['current']['signal_ready'] is False
    clock[0] += 1
    completed = store.observe_indicator_bucket('MSN', 'REAL', '2026-10-09', '5M', bucket + 1, 89, baseline, build)
    assert completed['current']['signal_ready'] is True
    assert completed['current']['indicator_price'] == 90
    assert completed['current']['buy_ema_fast'] == 90


def test_unready_bucket_clears_buy_confirmation_without_blocking_stop_loss():
    rule = StaticRule(StaticRuleParameters(whipsaw_enabled=False, buy_confirmation_enabled=True))
    p, source = portfolio(), context(current={**marks(), 'signal_ready': False})
    decision = rule.evaluate(source, p)
    state, filtered = apply_buy_filters(rule, decision, source, p,
        {'confirmation': {'active': True}}, observed_at=at('14:01:00'), exchange='HOSE')
    assert not state and filtered.action != 'BUY'
    p['position']['current_price'] = 96
    stop = rule.evaluate(source, p)
    assert stop.action == 'SELL' and stop.event == 'STOP_LOSS'


def test_confirmed_scale_in_still_obeys_whipsaw():
    rule = StaticRule(StaticRuleParameters(whipsaw_enabled=True, whipsaw_n=3,
        whipsaw_x=4, buy_ema_fast=1, buy_ema_slow=2, buy_window_enabled=False))
    source = context(previous=marks())
    source['bars'] = [dict(close=value, closed=True) for value in [10, 9, 11, 8, 12]]
    result = rule.evaluate({**source, 'confirmed_buy': True}, portfolio())
    assert (result.action, result.reason) == ('WAIT', 'WHIPSAW_LOCK')


@pytest.mark.parametrize('field,price', [('indicator_price', 11), ('indicator_price_vnd', 11000)])
def test_normalization_uses_frozen_indicator_price_and_preserves_raw_quote(tmp_path, field, price):
    import json
    cache = tmp_path / 'history.json'
    bars = [dict(time=at('14:45:00').replace(day=day).timestamp(), close=value)
            for day, value in zip(range(5, 9), [10, 11, 10, 12])]
    cache.write_text(json.dumps(dict(symbols={'MSN': bars}, resolution='1D',
                                    price_unit='THOUSAND_VND')), encoding='utf-8')
    original = dict(timestamp='2026-10-09 14:00:00', symbol='MSN', price=90,
                    ema_fast_period=2, ema_slow_period=3, rsi_period=3,
                    signal_ready=True, **{field: price})
    before = deepcopy(original)
    normalizer = DNSEIndicatorNormalizer(cache)
    result = normalizer.normalize([original])[0]
    assert result['normalization_ok'] and result['rsi'] == pytest.approx(600 / 11)
    assert result['ema_fast'] == pytest.approx(11.135802469135802)
    assert result['price'] == 90 and original == before
    unready = normalizer.normalize([{**original, 'signal_ready': 'False', 'ema_fast': 77}])[0]
    assert not unready.get('normalization_ok') and unready['ema_fast'] == 77


@pytest.mark.parametrize('failure', ['csv_append', 'state_write', 'pending_clear'])
def test_signal_log_recovers_exact_original_row_after_write_failure(tmp_path, monkeypatch, failure):
    path = tmp_path / 'audit.csv'
    log = SignalLog(path)
    row = dict(timestamp='2026-10-09 14:00:00', execution_mode='REAL', symbol='MSN',
               signal='BUY', signal_cycle='C1', acted='BUY', record_kind='SIGNAL_EVENT',
               price=100.0, indicator_price=100.0)
    with monkeypatch.context() as patch:
        if failure == 'csv_append':
            from pathlib import Path
            original_open = Path.open
            def fail_append(target, *args, **kwargs):
                if target == path and args and args[0] == 'a':
                    raise OSError('offline append failure')
                return original_open(target, *args, **kwargs)
            patch.setattr(Path, 'open', fail_append)
        elif failure == 'state_write':
            patch.setattr(log.state, 'write', Mock(side_effect=OSError('offline state failure')))
        else:
            original_write = log.pending.write
            def fail_clear(value):
                if value == {}:
                    raise OSError('offline clear failure')
                original_write(value)
            patch.setattr(log.pending, 'write', fail_clear)
        with pytest.raises(OSError):
            log.record(row)
    restarted = SignalLog(path)
    restarted.record({**row, 'timestamp': '2026-10-09 14:01:00'})
    records = restarted.read_all()
    assert len(records) == 1 and records[0]['timestamp'] == row['timestamp']
    assert records[0]['record_id'] and restarted.pending.read() == {}
    from openpyxl import load_workbook
    book = load_workbook(restarted.excel_archive.path_for(row['timestamp']), read_only=True)
    try:
        assert book.active.max_row == 2  # Header plus the same event, once.
    finally:
        book.close()
    # A genuine new signal is a separate row, even when its price is unchanged.
    assert restarted.record({**row, 'signal': 'SELL', 'signal_cycle': 'C2'})
    assert len(restarted.read_all()) == 2
    assert len({record['record_id'] for record in restarted.read_all()}) == 2
