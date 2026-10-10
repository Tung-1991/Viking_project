from copy import deepcopy
from datetime import datetime, timedelta
import json
import tkinter as tk

import customtkinter as ctk
import pytest

from viking_v2.rules.business import indicator_snapshot
from viking_v2.services.indicator_comparison import DNSEIndicatorNormalizer, number_comparison, rsi_observation_display
from viking_v2.trading.market import VN_TZ


def daily_history(count=260, final_price=99.0):
    final = datetime(2026, 10, 9, 9, tzinfo=VN_TZ)
    dates, current = [], final
    while len(dates) < count:
        if current.weekday() < 5:
            dates.append(current)
        current -= timedelta(days=1)
    return [{"time": int(day.timestamp()), "close": (final_price if day == final else
             21.53846156 + ((index % 11) - 5) * .038461538)} for index, day in enumerate(reversed(dates))]


def cache(path, bars=None, **fields):
    # AtomicJSONStore is also the production daemon's cache writer.
    from viking_v2.storage import AtomicJSONStore
    AtomicJSONStore(path).write({"symbols": {"HDB": bars if bars is not None else daily_history()},
                                 "updated_at": 1791540000, **fields})
    return DNSEIndicatorNormalizer(path)


def event(**changes):
    return {"timestamp": "2026-10-09 14:00:01", "symbol": "HDB", "price": 22.2,
            "ema_fast": 22.3163, "ema_slow": 22.3459, "rsi": 52.29,
            "ema_fast_period": 3, "ema_slow_period": 6, "rsi_period": 14,
            "signal": "SELL", "acted": "WAIT", **changes}


def settle(popup, root):
    finished = tk.BooleanVar(master=root, value=False)
    def check():
        if popup._normalization_future is None:
            finished.set(True)
        else:
            root.after(10, check)
    timeout = root.after(3000, lambda: finished.set(True))
    check()
    if not finished.get():
        root.wait_variable(finished)
    root.after_cancel(timeout)
    assert popup._normalization_future is None, "Normalization worker did not finish"


def popup_for(root, path, rows):
    from viking_v2.dashboard.windows import HistoryPopup
    return HistoryPopup(root, lambda _mode: [], signals_provider=lambda: rows,
                        indicator_normalizer=DNSEIndicatorNormalizer(path))


def test_uses_daily_history_and_each_recorded_tick_never_final_or_future_close(tmp_path):
    path = tmp_path / "market_bars.json"
    bars = daily_history()
    future = {"time": int(datetime(2026, 10, 12, 9, tzinfo=VN_TZ).timestamp()), "close": 1000.0}
    normalizer = cache(path, [future, *reversed(bars)])
    raw = [event(), event(price=22.5, timestamp="2026-10-09 14:02:01")]
    original, before = deepcopy(raw), path.read_bytes()
    result = normalizer.normalize(raw)
    for source, calculated in zip(raw, result):
        expected = indicator_snapshot([*bars[:-1], {"close": source["price"]}])
        for key in ("ema_fast", "ema_slow", "rsi", "rsi_previous"):
            assert calculated[key] == pytest.approx(expected[key])
        assert calculated["rsi_previous_date"] == "2026-10-08"
        assert calculated["dnse_indicators"]["rsi"] == 52.29
        assert calculated["normalization_bars"] == 259
        assert calculated["normalization_cache_at"] == datetime.fromtimestamp(1791540000, VN_TZ).isoformat(timespec="seconds")
        assert calculated["signal"] == "SELL" and calculated["acted"] == "WAIT"
        assert calculated["price"] == source["price"] and calculated["normalization_source"] == "DNSE"
        assert "comparison_entry" not in calculated
    assert raw == original and path.read_bytes() == before


def test_no_csv_or_api_needed_and_normalization_never_writes(tmp_path):
    path = tmp_path / "missing" / "market_bars.json"
    raw = event(rsi_previous=56.79)
    result = DNSEIndicatorNormalizer(path).normalize([raw])[0]
    assert result["rsi"] == 52.29 and result["rsi_previous"] == 56.79
    assert "lịch sử DNSE" in result["normalization_error"] and "CSV" not in result["normalization_error"]
    assert not path.parent.exists()


@pytest.mark.parametrize("payload", [[], None, {"symbols": []}, "not json"])
def test_bad_cache_keeps_original_numbers(tmp_path, payload):
    from viking_v2.storage import AtomicJSONStore
    path = tmp_path / "market_bars.json"
    AtomicJSONStore(path).write(payload)
    result = DNSEIndicatorNormalizer(path).normalize([event()])[0]
    assert result["normalization_error"] and result["rsi"] == 52.29


def test_explicit_units_and_capture_vnd_produce_same_result_without_rounding(tmp_path):
    bars = daily_history()
    thousands = cache(tmp_path / "thousands.json", bars)
    vnd = cache(tmp_path / "vnd.json", [{**bar, "close": bar["close"] * 1000} for bar in bars], price_unit="VND")
    a = thousands.normalize([event()])[0]
    b = vnd.normalize([event(price_vnd=22200)])[0]
    for key in ("ema_fast", "ema_slow", "rsi", "rsi_previous"):
        assert a[key] == pytest.approx(b[key])
    rounded = cache(tmp_path / "rounded.json", [{**bar, "close": round(bar["close"], 2)} for bar in bars])
    assert a["rsi_previous"] != rounded.normalize([event()])[0]["rsi_previous"]


@pytest.mark.parametrize("fields", [{"price_unit": "UNKNOWN"}, {"resolution": "1"}])
def test_does_not_guess_units_or_use_minute_bars_for_daily_indicators(tmp_path, fields):
    result = cache(tmp_path / "market_bars.json", **fields).normalize([event()])[0]
    assert result["normalization_error"] and result["ema_fast"] == 22.3163


def test_missing_symbol_never_borrows_another_history_and_valid_rows_still_work(tmp_path):
    normalizer = cache(tmp_path / "market_bars.json")
    valid, missing = normalizer.normalize([event(), event(symbol="MSN")])
    assert valid["normalization_ok"] and missing["normalization_error"]
    assert missing["rsi"] == 52.29


@pytest.mark.parametrize("changes", [
    {"price": 0}, {"price": float("nan")}, {"price": None}, {"price_vnd": 0},
    {"timestamp": "bad"}, {"ema_fast_period": 0}, {"rsi_period": "NaN"},
    {"ema_slow_period": 2.5}, {"symbol": "../orders"},
])
def test_invalid_observation_keeps_raw_indicators_without_using_latest_close(tmp_path, changes):
    result = cache(tmp_path / "market_bars.json").normalize([event(**changes)])[0]
    assert result["normalization_error"] and result["rsi"] == 52.29
    assert not result.get("normalization_ok")


@pytest.mark.parametrize("changes", [{"time": None}, {"close": float("inf")}, {"close": -1}])
def test_invalid_daily_history_is_not_silently_dropped(tmp_path, changes):
    bars = daily_history()
    bars[0].update(changes)
    result = cache(tmp_path / "market_bars.json", bars).normalize([event()])[0]
    assert result["normalization_error"] and result["rsi"] == 52.29


def test_conflicting_intraday_rows_are_not_passed_off_as_daily_history(tmp_path):
    bars = daily_history()
    bars.append({**bars[-2], "time": bars[-2]["time"] + 60, "close": 10})
    result = cache(tmp_path / "market_bars.json", bars).normalize([event()])[0]
    assert result["normalization_error"]


def test_periods_use_saved_values_and_legacy_defaults_are_explicit(tmp_path):
    normalizer = cache(tmp_path / "market_bars.json")
    legacy = event()
    for key in ("ema_fast_period", "ema_slow_period", "rsi_period"):
        del legacy[key]
    a, b = normalizer.normalize([legacy, event(ema_fast_period=4, ema_slow_period=9, rsi_period=5)])
    assert a["normalization_periods"] == [3, 6, 14]
    assert len(a["normalization_defaults"]) == 3 and b["normalization_defaults"] == []
    expected = indicator_snapshot([*daily_history()[:-1], {"close": 22.2}], 4, 9, 5)
    assert b["rsi"] == pytest.approx(expected["rsi"])


@pytest.mark.parametrize("count,ok", [(15, False), (16, True), (20, True)])
def test_requires_enough_history_for_previous_rsi_not_arbitrary_100_csv_bars(tmp_path, count, ok):
    result = cache(tmp_path / "market_bars.json", daily_history(count=count)).normalize([event()])[0]
    assert bool(result.get("normalization_ok")) is ok


def test_session_date_is_vietnamese_and_previous_rsi_is_previous_session_not_calendar_day(tmp_path):
    bars = daily_history()
    normalizer = cache(tmp_path / "market_bars.json", bars)
    # UTC Sunday night is already Monday in Vietnam; previous session is Friday.
    result = normalizer.normalize([event(timestamp="2026-10-11T23:30:00Z")])[0]
    assert result["rsi_previous_date"] == "2026-10-09"


@pytest.mark.parametrize("a,b,expected", [
    (56.79, 56.86, "56.79 < 56.86"), (56.79, 56.79, "56.79 = 56.79"),
    (56.791, 56.789, "56.791 > 56.789"), (52.29, "", "—"), (float("nan"), 50, "—"),
])
def test_numeric_comparison_never_changes_raw_sign_due_to_rounding(a, b, expected):
    assert number_comparison(a, b) == expected


@pytest.mark.parametrize("current,previous,text,issue", [
    (52.29, "", "52.29 · Trước: —", "Thiếu RSI trước"),
    ("52.29", None, "52.29 · Trước: —", "Thiếu RSI trước"),
    (None, 56.79, "— · Trước: 56.79", "Thiếu RSI hiện tại"),
    (0, None, "0.00 · Trước: —", "Thiếu RSI trước"),
    (100, float("nan"), "100.00 · Trước: —", "Thiếu RSI trước"),
    (float("inf"), -1, "—", "Thiếu RSI"),
    (52.29, 56.79, "52.29 < 56.79", ""),
    (56.791, 56.789, "56.791 > 56.789", ""),
    (56.79, 56.79, "56.79 = 56.79", ""),
])
def test_rsi_display_keeps_valid_side_without_guessing_missing_baseline(current, previous, text, issue):
    assert rsi_observation_display(current, previous) == (text, issue)


def test_single_normalize_button_preserves_events_selection_and_raw_values(ui_root, tmp_path):
    path = tmp_path / "market_bars.json"
    cache(path)
    raw = [event(rsi_previous=56.79, blocked_by="BUY_WINDOW_WAIT", signal="BUY"), event(symbol="MSN")]
    before, cache_before = deepcopy(raw), path.read_bytes()
    popup = popup_for(ui_root, path, raw)
    try:
        assert not popup.normalization_enabled
        assert not hasattr(popup, "signal_basis_button") and not hasattr(popup, "signal_import_button")
        tree, hdb = popup.signal_tree, "symbol:2026-10-09:HDB"
        leaf = tree.get_children(hdb)[0]
        tree.item(hdb, open=True)
        tree.selection_set(leaf)
        assert tree.set(leaf, "rsi_comparison") == "52.29 < 56.79"
        popup.normalization_button.invoke()
        settle(popup, ui_root)
        expected = indicator_snapshot([*daily_history()[:-1], {"close": 22.2}])
        assert tree.set(leaf, "rsi_comparison") == number_comparison(expected["rsi"], expected["rsi_previous"])
        assert tree.set(leaf, "display_signal") == "ENTRY" and tree.set(leaf, "suggestion") == "Chờ giờ"
        assert tree.item(hdb, "open") and tree.selection() == (leaf,)
        assert tree.heading("suggestion", "text") == "XỬ LÝ ĐÃ GHI"
        msn = tree.get_children("symbol:2026-10-09:MSN")[0]
        assert tree.set(msn, "rsi_comparison") == "52.29 · Trước: —"
        assert "1 dòng giữ số gốc" in popup.signal_status.cget("text")
        popup.normalization_button.invoke()
        assert tree.set(leaf, "rsi_comparison") == "52.29 < 56.79"
        assert raw == before and path.read_bytes() == cache_before
    finally:
        popup.close()


def test_periodic_processing_is_recorded_not_a_new_decision(ui_root, tmp_path, monkeypatch):
    from types import SimpleNamespace
    from viking_v2.dashboard.windows import HistoryPopup
    raw = {**event(signal="PERIODIC"), "record_kind": "PERIODIC", "price_vnd": 22200,
           "entry": False, "reason": "WHIPSAW_LOCK", "ema_cross_state": "WAIT_DOWN",
           "ema_cross_required": True, "scheduled_at": "2026-10-09T14:00:00+07:00"}
    path = tmp_path / "market_bars.json"
    cache(path)
    popup = HistoryPopup(ui_root, lambda _: [], trace_store=SimpleNamespace(path=tmp_path / "trace.sqlite3", read=lambda **_: [raw]),
                         indicator_normalizer=DNSEIndicatorNormalizer(path))
    try:
        popup._toggle_normalization()
        settle(popup, ui_root)
        tree = popup.signal_tree
        leaf = tree.get_children("symbol:2026-10-09:HDB")[0]
        assert tree.set(leaf, "suggestion") == "Chưa đạt"
        assert tree.set(leaf, "display_signal") == "ĐỊNH KỲ"
        assert "WHIPSAW" in tree.set(leaf, "reason") or "EMA nhiễu" in tree.set(leaf, "reason")
        assert tree.set(leaf, "ema_cross_display") != "—"
        shown = []
        monkeypatch.setattr("viking_v2.dashboard.windows.messagebox.showinfo", lambda *args, **_: shown.append(args[1]))
        tree.selection_set(leaf)
        popup._show_signal_details()
        assert "DNSE đã ghi:" in shown[0] and "không phải quyết định mới hay số TradingView" in shown[0]
    finally:
        popup.close()


def test_legacy_rsi_survives_copy_excel_and_restore(ui_root, tmp_path, monkeypatch):
    from openpyxl import load_workbook
    original = [event()]
    before = deepcopy(original)
    popup = popup_for(ui_root, tmp_path / "market_bars.json", original)
    try:
        tree, hdb = popup.signal_tree, "symbol:2026-10-09:HDB"
        leaf = tree.get_children(hdb)[0]
        assert tree.set(leaf, "rsi_comparison") == "52.29 · Trước: —"
        tree.selection_set(leaf)
        copied = []
        monkeypatch.setattr(popup.top, "clipboard_clear", lambda: None)
        monkeypatch.setattr(popup.top, "clipboard_append", copied.append)
        popup._copy_signal_rows()
        assert "52.29 · Trước: —" in copied[0]
        target = tmp_path / "legacy.xlsx"
        monkeypatch.setattr("viking_v2.dashboard.windows.filedialog.asksaveasfilename", lambda **_: str(target))
        popup._export_signals(selected_only=True)
        book = load_workbook(target)
        try:
            assert book.sheetnames == ["DNSE GỐC"]
            assert "52.29 · Trước: —" in [cell.value for cell in book["DNSE GỐC"][2]]
        finally:
            book.close()
        monkeypatch.setattr("viking_v2.dashboard.windows.messagebox.askyesno", lambda *_, **__: True)
        popup._delete_signal_rows()
        assert not tree.get_children()
        popup._restore_signal_rows()
        assert tree.set(tree.get_children(hdb)[0], "rsi_comparison") == "52.29 · Trước: —"
        assert original == before and "rsi_previous" not in original[0]
    finally:
        popup.close()


def test_distinct_signal_cycles_remain_even_in_same_second():
    from viking_v2.dashboard.windows import signal_rows_by_day
    rows = [event(signal_cycle="first"), event(signal_cycle="second")]
    assert len(signal_rows_by_day(rows)[0]["rows"]) == 2
    assert len(signal_rows_by_day([rows[0], dict(rows[0])])[0]["rows"]) == 1


def test_refresh_recalculates_from_updated_cache_without_modifying_old_events(ui_root, tmp_path):
    path = tmp_path / "market_bars.json"
    cache(path)
    raw = [event()]
    popup = popup_for(ui_root, path, raw)
    try:
        popup._toggle_normalization()
        settle(popup, ui_root)
        leaf = popup.signal_tree.get_children("symbol:2026-10-09:HDB")[0]
        before = popup.signal_tree.set(leaf, "rsi_comparison")
        cache(path, [{**bar, "close": bar["close"] * 0.5} for bar in daily_history()])
        popup._refresh_signals()
        settle(popup, ui_root)
        assert popup.signal_tree.set(leaf, "rsi_comparison") != before
        popup._toggle_normalization()
        assert popup.signal_tree.set(leaf, "rsi_comparison") == "52.29 · Trước: —"
    finally:
        popup.close()


def test_excel_exports_raw_and_normalized_only_when_enabled_and_blocks_formula_injection(ui_root, tmp_path, monkeypatch):
    from openpyxl import load_workbook
    path = tmp_path / "market_bars.json"
    cache(path)
    rows = [event(), event(symbol="=DANGEROUS()", rsi_previous=56.79)]
    before = deepcopy(rows)
    target = tmp_path / "signals.xlsx"
    monkeypatch.setattr("viking_v2.dashboard.windows.filedialog.asksaveasfilename", lambda **_: str(target))
    monkeypatch.setattr("viking_v2.dashboard.windows.messagebox.showerror", lambda *_, **__: pytest.fail("Export failed"))
    popup = popup_for(ui_root, path, rows)
    try:
        popup._toggle_normalization()
        settle(popup, ui_root)
        popup._export_signals()
        book = load_workbook(target)
        try:
            assert book.sheetnames == ["DNSE GỐC", "CHUẨN HOÁ"]
            for sheet in book:
                assert sheet.max_row == 3
                unsafe = next(row for row in sheet.iter_rows() if row[1].value == "=DANGEROUS()")
                assert unsafe[1].data_type == "s"
            assert rows == before
            assert any("52.29 · Trước: —" == cell.value for row in book["DNSE GỐC"].iter_rows() for cell in row)
            assert any("DNSE · Tính lại" == cell.value for row in book["CHUẨN HOÁ"].iter_rows() for cell in row)
            normalized = book["CHUẨN HOÁ"]
            headers = [cell.value for cell in normalized[1]]
            hdb = next(row for row in normalized.iter_rows(min_row=2) if row[1].value == "HDB")
            assert hdb[headers.index("LỊCH SỬ ĐẾN")].value == "2026-10-08"
            assert hdb[headers.index("CHU KỲ TÍNH")].value == "3/6/14"
        finally:
            book.close()
    finally:
        popup.close()


def test_off_and_close_during_background_calculation_ignore_late_results(ui_root, tmp_path):
    from concurrent.futures import Future
    popup = popup_for(ui_root, tmp_path / "missing.json", [event()])
    pending = Future()
    class Executor:
        def submit(self, *_, **__):
            return pending
        def shutdown(self, **_):
            pass
    popup._normalization_executor = Executor()
    try:
        popup._toggle_normalization()
        generation = popup._normalization_generation
        assert popup._normalization_after is not None
        popup._toggle_normalization()
        assert popup._normalization_after is None and not popup._normalization_rows
        popup._poll_normalization(generation, [event()])
        assert not popup._normalization_rows
        assert popup.signal_tree.set(popup.signal_tree.get_children("symbol:2026-10-09:HDB")[0], "rsi_comparison") == "52.29 · Trước: —"
    finally:
        popup.close()


@pytest.mark.parametrize("scaling", [1.0, 1.5])
def test_normalized_toolbar_fits_and_has_only_one_normalization_control(ui_root, tmp_path, scaling):
    from support.tests.test_signal_history import _capture_ui
    ctk.set_widget_scaling(scaling)
    path = tmp_path / "market_bars.json"
    cache(path)
    popup = popup_for(ui_root, path, [event(rsi_previous=56.79, signal="BUY", blocked_by="BUY_WINDOW_WAIT")])
    try:
        popup.top.geometry("980x560")
        popup.tabs.set("TÍN HIỆU")
        popup._history_tab_changed()
        popup._toggle_normalization()
        settle(popup, ui_root)
        ui_root.update()
        toolbar = popup.normalization_button.master
        controls = [child for child in toolbar.winfo_children() if child.grid_info().get("row") == 1]
        assert len(controls) == 3
        assert all(child.winfo_x() + child.winfo_width() <= toolbar.winfo_width() for child in controls)
        assert popup.normalization_button.cget("text") == "CHUẨN HOÁ · ON"
        assert "không cần CSV" in popup.normalization_hint.text
        assert "Không tự suy đoán hệ số" in popup.normalization_hint.text
        popup.top.geometry("1600x820+0+0")
        popup.signal_tree.item("symbol:2026-10-09:HDB", open=True)
        _capture_ui(ui_root, popup.top, f"signal-normalized-{scaling}")
    finally:
        popup.close()
        ctk.set_widget_scaling(1.0)
