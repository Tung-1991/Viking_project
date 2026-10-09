from datetime import date, datetime, timedelta

from viking_v2.backtest.entry_audit import audit_entry_day
from viking_v2.config import AppSettings
from viking_v2.trading.market import VN_TZ


def test_daily_seed_plus_one_running_candle_not_minute_ema_or_today_final_close(tmp_path):
    start = datetime(2026, 10, 9, tzinfo=VN_TZ)
    daily = [{"time": int((start - timedelta(days=150-i)).timestamp()),
              "open": 20 + i/100, "high": 20.2+i/100, "low": 19.9+i/100,
              "close": 20+i/20+(0.1 if i%2 == 0 else -0.1), "volume": 1000} for i in range(150)]
    daily.append({"time": int(start.timestamp()), "close": 10000})  # Never use today's completed candle.
    minutes = [{"time": int(start.replace(hour=14, minute=i).timestamp()), "open": 28,
                "high": 28.1, "low": 27.9, "close": 28, "volume": 100} for i in range(3)]
    settings = AppSettings(rule_parameters={"buy_signal_require_ema_cross": True}).normalize()
    rows = audit_entry_day("IDC", date(2026, 10, 9), daily, minutes, settings, exchange="HNX", baseline_source="TEST")
    assert len(rows) == 3 and rows[0]["timestamp"].endswith("14:00:59+07:00")
    assert all(r["entry"] for r in rows)
    assert not any(r["fresh_cross_entry"] for r in rows)
    assert all(r["reason"] == "EMA_CROSS_REQUIRED" for r in rows)
    assert rows[0]["ema_fast"] == rows[-1]["ema_fast"]  # Same daily provisional close, not a growing minute EMA.
    assert rows[0]["rsi_previous_date"] == "2026-10-08"
    assert all(r["execution_mode"] == "REPLAY" for r in rows)
    settings.rule_parameters["buy_signal_require_ema_cross"] = False
    ready = audit_entry_day("IDC", date(2026,10,9), daily, minutes, settings, exchange="HNX", baseline_source="TEST")
    assert all(r["reason"] == "TECHNICAL_READY" for r in ready)
    assert all(r["rule_action"] != "BUY" for r in ready)  # No invented executable order.
    from viking_v2.services.signal_trace import export_trace
    from openpyxl import load_workbook
    target = tmp_path / "entry.xlsx"
    export_trace(rows, target)
    book = load_workbook(target)
    try:
        assert book.sheetnames == ["ENTRY REPLAY", "SETTING"]
        sheet = book.worksheets[0]
        assert sheet.max_column == 16
        assert "Chưa có lần EMA vừa cắt lên" in [cell.value for cell in sheet[2]]
        assert "OTP còn hạn" not in [cell.value for cell in sheet[1]]
    finally:
        book.close()
