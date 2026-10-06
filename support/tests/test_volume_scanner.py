from __future__ import annotations

from datetime import datetime
from zoneinfo import ZoneInfo

import pytest
from openpyxl import load_workbook

from viking_v2.services.volume_scanner import (
    VN100_SYMBOLS_BY_MARKET_CAP,
    VolumeScanOptions,
    VolumeScanner,
    evaluate_volume_change,
    export_volume_scan,
    export_watchlist,
    import_watchlist,
)


def _bars(volumes: list[float], *, open_last: bool = False) -> list[dict[str, object]]:
    return [
        {
            "time": index + 1,
            "volume": volume,
            "closed": not (open_last and index == len(volumes) - 1),
        }
        for index, volume in enumerate(volumes)
    ]


def test_volume_change_uses_two_adjacent_five_session_blocks() -> None:
    row, reason = evaluate_volume_change("fpt", _bars([100] * 5 + [120] * 5), 5)

    assert reason == ""
    assert row is not None
    assert row.symbol == "FPT"
    assert row.previous_average == pytest.approx(100)
    assert row.recent_average == pytest.approx(120)
    assert row.change_pct == pytest.approx(20)
    assert row.status == "TĂNG"


def test_volume_change_supports_ten_sessions_and_ignores_open_bar() -> None:
    volumes = [1_000] * 10 + [800] * 10 + [9_999_999]
    row, reason = evaluate_volume_change("MBB", _bars(volumes, open_last=True), 10)

    assert reason == ""
    assert row is not None
    assert row.previous_average == pytest.approx(1_000)
    assert row.recent_average == pytest.approx(800)
    assert row.change_pct == pytest.approx(-20)
    assert row.status == "GIẢM"


def test_volume_change_rejects_missing_data_and_zero_base() -> None:
    assert evaluate_volume_change("FPT", _bars([100] * 9), 5) == (None, "INSUFFICIENT_DATA")
    assert evaluate_volume_change("FPT", _bars([0] * 5 + [100] * 5), 5) == (None, "ZERO_BASE")


def test_scanner_applies_inclusive_threshold_sorting_and_limits() -> None:
    symbols = VN100_SYMBOLS_BY_MARKET_CAP[:4]
    changes = {
        symbols[0]: [100] * 5 + [130] * 5,
        symbols[1]: [100] * 5 + [50] * 5,
        symbols[2]: [100] * 5 + [140] * 5,
        symbols[3]: [100] * 5 + [120] * 5,
    }
    calls: list[str] = []

    def load(symbol: str, **_kwargs: object) -> list[dict[str, object]]:
        calls.append(symbol)
        return _bars(changes[symbol])

    result = VolumeScanner(load).scan(VolumeScanOptions(
        scan_count=4,
        result_count=3,
        sessions=5,
        threshold_pct=20,
        direction="CẢ HAI",
    ))

    assert calls == list(symbols)
    assert [row.symbol for row in result.rows] == [symbols[1], symbols[2], symbols[0]]
    assert result.summary.matched == 4
    assert result.summary.returned == 3


@pytest.mark.parametrize(
    ("direction", "expected"),
    [
        ("TĂNG", [30.0, 20.0]),
        ("GIẢM", [-30.0, -20.0]),
    ],
)
def test_scanner_direction_includes_exact_twenty_percent(
    direction: str, expected: list[float],
) -> None:
    symbols = VN100_SYMBOLS_BY_MARKET_CAP[:4]
    recent = (130, 120, 80, 70)

    def load(symbol: str, **_kwargs: object) -> list[dict[str, object]]:
        value = recent[symbols.index(symbol)]
        return _bars([100] * 5 + [value] * 5)

    result = VolumeScanner(load).scan(VolumeScanOptions(
        scan_count=4,
        result_count=4,
        sessions=5,
        threshold_pct=20,
        direction=direction,
    ))

    assert [row.change_pct for row in result.rows] == pytest.approx(expected)


def test_scanner_counts_api_missing_and_zero_base_without_stopping() -> None:
    symbols = VN100_SYMBOLS_BY_MARKET_CAP[:4]
    last_status = {"value": 200}

    def load(symbol: str, **_kwargs: object) -> list[dict[str, object]]:
        if symbol == symbols[0]:
            raise RuntimeError("DNSE unavailable")
        if symbol == symbols[1]:
            last_status["value"] = 200
            return _bars([100] * 9)
        if symbol == symbols[2]:
            last_status["value"] = 200
            return _bars([0] * 5 + [100] * 5)
        last_status["value"] = 200
        return _bars([100] * 5 + [130] * 5)

    scanner = VolumeScanner(load, api_health=lambda: {"last_status": last_status["value"]})
    result = scanner.scan(VolumeScanOptions(scan_count=4, result_count=4))

    assert [row.symbol for row in result.rows] == [symbols[3]]
    assert result.summary.api_errors == 1
    assert result.summary.insufficient == 1
    assert result.summary.zero_base == 1


def test_vn100_universe_has_one_hundred_unique_symbols() -> None:
    assert len(VN100_SYMBOLS_BY_MARKET_CAP) == 100
    assert len(set(VN100_SYMBOLS_BY_MARKET_CAP)) == 100
    assert VN100_SYMBOLS_BY_MARKET_CAP[:3] == ("VIC", "VHM", "VCB")


def test_export_writes_numeric_excel_to_exports_directory(tmp_path) -> None:
    row, _reason = evaluate_volume_change("FPT", _bars([100] * 5 + [125] * 5), 5)
    assert row is not None

    path = export_volume_scan(
        [row],
        output_dir=tmp_path / "exports",
        now=datetime(2026, 9, 22, 15, 1, 2, tzinfo=ZoneInfo("Asia/Ho_Chi_Minh")),
    )

    assert path.name == "vn100_volume_20260922_150102.xlsx"
    book = load_workbook(path, data_only=True)
    try:
        sheet = book["VOLUME"]
        assert [cell.value for cell in sheet[1]] == [
            "MÃ CK", "TB KỲ TRƯỚC", "TB KỲ GẦN NHẤT", "% THAY ĐỔI", "TRẠNG THÁI",
        ]
        assert [sheet.cell(2, column).value for column in range(1, 6)] == [
            "FPT", 100, 125, 0.25, "TĂNG",
        ]
        assert sheet["D2"].number_format == "0.00%;[Red]-0.00%"
    finally:
        book.close()


def test_export_rejects_empty_results_without_creating_directory(tmp_path) -> None:
    output = tmp_path / "exports"
    with pytest.raises(ValueError, match="Không có kết quả"):
        export_volume_scan([], output_dir=output)
    assert not output.exists()


def test_watchlist_excel_round_trip_preserves_order_exchange_and_priority(tmp_path) -> None:
    path = export_watchlist(
        ["FPT", "MBB", "SSI"],
        priority_symbols=["MBB"],
        symbol_exchanges={"FPT": "HOSE", "MBB": "HSX"},
        output_dir=tmp_path,
        now=datetime(2026, 9, 23, 10, 11, 12, tzinfo=ZoneInfo("Asia/Ho_Chi_Minh")),
    )

    assert path.name == "watchlist_20260923_101112.xlsx"
    imported = import_watchlist(path)
    assert imported.symbols == ("FPT", "MBB", "SSI")
    assert imported.priority_symbols == ("MBB",)
    assert imported.symbol_exchanges == {"FPT": "HOSE", "MBB": "HOSE"}


def test_watchlist_import_rejects_invalid_symbol_without_partial_result(tmp_path) -> None:
    from openpyxl import Workbook

    path = tmp_path / "invalid.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.append(("MÃ CK",))
    sheet.append(("FPT!",))
    book.save(path)
    book.close()

    with pytest.raises(ValueError, match="không hợp lệ"):
        import_watchlist(path)
