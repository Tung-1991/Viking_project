from copy import deepcopy
from datetime import datetime, timedelta
import json
from types import SimpleNamespace

import pytest

from viking_v2.config import AppSettings
from viking_v2.services.signal_trace import SignalTraceStore, export_trace, compare_entry
from viking_v2.trading.market import VN_TZ


NOW = datetime(2026, 10, 9, 14, 0, 2, tzinfo=VN_TZ)


def settings(**kwargs):
    return AppSettings(watchlist=["IDC", "HDB"], priority_symbols=["IDC"], **kwargs).normalize()


def status(now=NOW, *, mode="REAL", signal="", reason="NO_NEW_BUY_SIGNAL", rsi=70):
    return {
        "bot_enabled": True, "active_symbols": ["IDC", "HDB"], "market_status": "OPEN",
        "ticks": {s: {"symbol": s, "price": 35.2, "source": "WS", "timestamp": now.timestamp()} for s in ("IDC", "HDB")},
        "decisions_by_mode": {mode: {"IDC": {
            "symbol": "IDC", "signal": signal, "reason": reason, "action": "WAIT",
            "details": {"updated_at": now.isoformat(), "execution_mode": mode,
                        "indicators": {"ema_fast": 35.001, "ema_slow": 34.644, "rsi": rsi,
                                       "rsi_previous": 66.5, "rsi_previous_time": (now - timedelta(days=1)).timestamp()},
                        "entry_checks": {"order_budget": 7000000, "available_cash": 50000000,
                                         "whipsaw_crossovers": 0, "whipsaw_limit": 3, "whipsaw_window": 7},
                        "buy_window": {"state": "OPEN"}}}}}}


@pytest.mark.parametrize("time,expected", [("13:59:59", False), ("14:00:00", True),
    ("14:01:59", True), ("14:30:01", True), ("14:30:59", True), ("14:31:00", False)])
def test_window_boundaries_and_final_sample(time, expected):
    now = datetime.fromisoformat("2026-10-09T" + time).replace(tzinfo=VN_TZ)
    assert (SignalTraceStore.schedule(settings(), now) is not None) is expected


def test_disabled_or_outside_window_does_not_create_files(tmp_path):
    store = SignalTraceStore(tmp_path / "trace.sqlite3")
    assert store.capture(status(), settings(signal_trace_enabled=False), active_mode="REAL", now=NOW) == 0
    assert store.capture(status(), settings(), active_mode="REAL", now=NOW.replace(hour=18)) == 0
    assert store.read() == [] and store.days() == [] and not store.path.exists()


def test_all_symbols_no_signal_still_recorded_with_numeric_reason_and_restart_dedup(tmp_path):
    store = SignalTraceStore(tmp_path / "trace.sqlite3")
    source, configured = status(), settings(telegram_chat_id="DO_NOT_SAVE_THIS")
    source_before = deepcopy(source)
    assert store.capture(source, configured, active_mode="REAL", otp_ok=True, now=NOW) == 2
    assert store.capture(source, configured, active_mode="REAL", now=NOW + timedelta(seconds=1)) == 0
    reopened = SignalTraceStore(store.path)
    assert reopened.capture(source, configured, active_mode="REAL", now=NOW + timedelta(seconds=2)) == 0
    rows = store.read()
    idc = next(r for r in rows if r["symbol"] == "IDC")
    assert idc["entry"] is True and idc["rule_action"] == "WAIT" and idc["ema_cross_required"] is True
    assert idc["price_vnd"] == 35200 and idc["rsi_comparison"] == "70.00 > 66.50"
    assert idc["rsi_previous_date"] == "2026-10-08" and idc["otp_ok"] is True
    assert "DO_NOT_SAVE_THIS" not in json.dumps(rows)
    assert next(r for r in rows if r["symbol"] == "HDB")["entry"] is None
    assert source == source_before
    later = NOW + timedelta(minutes=2)
    assert reopened.capture(status(later), configured, active_mode="REAL", now=later) == 2
    assert len(reopened.read()) == 4


def test_paper_real_separate_and_missing_real_does_not_use_paper(tmp_path):
    store = SignalTraceStore(tmp_path / "trace.sqlite3")
    value = status(mode="PAPER")
    store.capture(value, settings(), active_mode="REAL", now=NOW)
    assert len(store.read(mode="REAL")) == len(store.read(mode="PAPER")) == 2
    assert store.read(mode="REAL", symbol="IDC")[0]["entry"] is None
    assert store.read(mode="PAPER", symbol="IDC")[0]["entry"] is True
    assert store.read(mode="PAPER")[0]["otp_ok"] is None


@pytest.mark.parametrize("issue", ["stale", "frozen", "old_decision"])
def test_bad_quote_or_decision_not_reported_as_entry(tmp_path, issue):
    value = status()
    if issue == "old_decision":
        value["decisions_by_mode"]["REAL"]["IDC"]["details"]["updated_at"] = (NOW-timedelta(minutes=1)).isoformat()
    else:
        value["ticks"]["IDC"][issue] = True
    store = SignalTraceStore(tmp_path / "trace.sqlite3")
    store.capture(value, settings(), active_mode="REAL", now=NOW)
    assert store.read(symbol="IDC")[0]["entry"] is None


def test_retention_is_only_trace_and_does_not_touch_orders(tmp_path):
    store = SignalTraceStore(tmp_path / "trace.sqlite3")
    old = NOW - timedelta(days=35)
    store.capture(status(old), settings(), active_mode="REAL", now=old)
    store.capture(status(), settings(), active_mode="REAL", now=NOW)
    assert store.days() == ["2026-10-09"]


def test_nonfinite_indicator_is_unknown_and_does_not_drop_other_symbols(tmp_path):
    value = status(rsi=float("nan"))
    store = SignalTraceStore(tmp_path / "trace.sqlite3")
    assert store.capture(value, settings(), active_mode="REAL", now=NOW) == 2
    row = store.read(symbol="IDC")[0]
    assert row["entry"] is None and row["rsi"] is None and row["rsi_comparison"] == "—"


def test_excel_keeps_original_numbers_settings_and_no_formulas(tmp_path):
    from openpyxl import load_workbook
    store = SignalTraceStore(tmp_path / "trace.sqlite3")
    store.capture(status(), settings(), active_mode="REAL", now=NOW)
    rows = store.read()
    rows[0]["queue_summary"] = "=BAD()"
    path = tmp_path / "trace.xlsx"
    export_trace(rows, path)
    book = load_workbook(path)
    try:
        assert book.sheetnames == ["TRACE", "SETTING"]
        cells = list(book["TRACE"][2])
        cell = next(c for c in cells if c.value == "=BAD()")
        assert cell.data_type == "s"
        assert any(c.value == 66.5 for r in book["TRACE"].iter_rows() for c in r)
        assert book["SETTING"].max_row == 2
    finally:
        book.close()


def test_capture_is_in_the_same_symbol_tree_as_events_and_exportable(ui_root, tmp_path, monkeypatch):
    from viking_v2.dashboard.windows import HistoryPopup
    store = SignalTraceStore(tmp_path / "trace.sqlite3")
    store.capture(status(), settings(), active_mode="REAL", now=NOW)
    events = [{"timestamp": "2026-10-09 13:40:53", "symbol": "HDB", "execution_mode": "REAL",
               "signal": "BUY", "acted": "WAIT", "blocked_by": "BUY_WINDOW_WAIT"}]
    popup = HistoryPopup(ui_root, lambda _mode: [], trace_store=store, signals_provider=lambda: events)
    try:
        assert set(popup.tabs._tab_dict) == {"CKCS REAL", "CKCS PAPER", "TÍN HIỆU"}
        assert not hasattr(popup, "signal_view_button") and not hasattr(popup, "trace_tree")
        tree = popup.signal_tree
        hdb = tree.get_children("symbol:2026-10-09:HDB")
        assert [tree.set(child, "display_signal") for child in hdb] == ["ĐỊNH KỲ", "ENTRY"]
        idc = tree.get_children("symbol:2026-10-09:IDC")[0]
        assert tree.set(idc, "display_signal") == "ĐỊNH KỲ"
        assert tree.set(idc, "suggestion") == "EMA/RSI đạt"
        assert tree.set(idc, "reason") == "Chờ cắt EMA"
        monkeypatch.setattr("viking_v2.dashboard.windows.filedialog.asksaveasfilename", lambda **_kw: str(tmp_path / "out.xlsx"))
        popup._export_signals()
        assert (tmp_path / "out.xlsx").exists()
        from openpyxl import load_workbook
        book = load_workbook(tmp_path / "out.xlsx")
        try:
            assert book.sheetnames == ["DNSE GỐC", "TRACE", "SETTING"]
            assert book["DNSE GỐC"].max_row == 4 and book["TRACE"].max_row == 3
        finally:
            book.close()
    finally:
        popup.close()


def test_trace_rule_fields_save_without_modifying_indicator_interval(ui_root, tmp_path, monkeypatch):
    from viking_v2 import config
    from viking_v2.rules.window import RuleSettingsPopup
    monkeypatch.setattr(config, "ACCOUNTS_ROOT", tmp_path)
    initial = settings(signal_trace_enabled=False, realtime_indicator_interval="TICK")
    popup = RuleSettingsPopup(ui_root, initial, "TRACE_TEST", lambda: None)
    try:
        popup.signal_trace_enabled.set(True)
        popup.signal_trace_interval.delete(0, "end")
        popup.signal_trace_interval.insert(0, "3")
        popup.save()
        saved = config.load_settings("TRACE_TEST")
        assert saved.signal_trace_enabled and saved.signal_trace_interval_minutes == 3
        assert saved.signal_trace_start == "14:00" and saved.signal_trace_end == "14:30"
        assert saved.realtime_indicator_interval == "TICK"
    finally:
        popup._close()


def test_gui_trace_uses_final_planning_block_and_never_changes_daemon_source(tmp_path, monkeypatch):
    from viking_v2.dashboard.actions import DashboardActionsMixin
    import viking_v2.dashboard.actions as actions
    class Clock(datetime):
        @classmethod
        def now(cls, tz=None):
            return NOW
    monkeypatch.setattr(actions, "datetime", Clock)
    source = status()
    original = deepcopy(source)
    final = deepcopy(source["decisions_by_mode"]["REAL"]["IDC"])
    final["reason"] = "NO_AVAILABLE_CAPITAL"
    view = SimpleNamespace(settings=settings(),
        bridge=SimpleNamespace(root=tmp_path, read_config=lambda: SimpleNamespace(paper_mode=False)),
        queue=SimpleNamespace(list_all=lambda: []), real=SimpleNamespace(has_trading_token=lambda: True),
        rule_state=SimpleNamespace(entry_pause=lambda _mode: {}), logger=SimpleNamespace(info=lambda *_a: None),
        _trace_final_decisions={("REAL", "IDC"): final})
    DashboardActionsMixin._capture_signal_trace(view, source, "RUNNING")
    assert SignalTraceStore(tmp_path / "signal_trace.sqlite3").read(symbol="IDC")[0]["reason"] == "NO_AVAILABLE_CAPITAL"
    assert source == original


def test_trace_disk_failure_is_coalesced_detailed_and_does_not_stop_trading(tmp_path, monkeypatch):
    from viking_v2.dashboard.actions import DashboardActionsMixin
    import viking_v2.dashboard.actions as actions
    class Clock(datetime):
        @classmethod
        def now(cls, tz=None):
            return NOW
    monkeypatch.setattr(actions, "datetime", Clock)
    warnings = []
    def fail(*_args, **_kwargs):
        raise OSError("Disk is full")
    view = SimpleNamespace(settings=settings(),
        _signal_trace_store=SimpleNamespace(capture=fail),
        bridge=SimpleNamespace(root=tmp_path, read_config=lambda: SimpleNamespace(paper_mode=False)),
        queue=SimpleNamespace(list_all=lambda: []), real=SimpleNamespace(has_trading_token=lambda: True),
        rule_state=SimpleNamespace(entry_pause=lambda _mode: {}),
        logger=SimpleNamespace(warning=lambda *args: warnings.append(args)))
    DashboardActionsMixin._capture_signal_trace(view, status(), "RUNNING")
    DashboardActionsMixin._capture_signal_trace(view, status(), "RUNNING")
    assert len(warnings) == 1 and "Disk is full" in str(warnings[0])


def test_entire_capture_window_never_dispatches_telegram_or_orders_even_with_raw_buy_sell(tmp_path, monkeypatch):
    from viking_v2.dashboard.actions import DashboardActionsMixin
    import viking_v2.dashboard.actions as actions
    clock = [NOW]
    class Clock(datetime):
        @classmethod
        def now(cls, tz=None):
            return clock[0]
    monkeypatch.setattr(actions, "datetime", Clock)
    attempted, warnings = [], []
    class ReadOnlyGuard:
        def __init__(self, **allowed):
            self.allowed = allowed
        def __getattr__(self, name):
            if name in self.allowed:
                return self.allowed[name]
            attempted.append(name)
            raise AssertionError("Capture attempted external mutation: " + name)
    view = SimpleNamespace(settings=settings(),
        bridge=SimpleNamespace(root=tmp_path, read_config=lambda: SimpleNamespace(paper_mode=False)),
        queue=ReadOnlyGuard(list_all=lambda: []), real=ReadOnlyGuard(has_trading_token=lambda: True),
        telegram=ReadOnlyGuard(), rule_state=ReadOnlyGuard(entry_pause=lambda _mode: {}),
        logger=SimpleNamespace(info=lambda *_a: None, warning=lambda *args: warnings.append(args)))
    for index in range(16):
        clock[0] = NOW + timedelta(minutes=index * 2)
        source = status(clock[0], signal="BUY" if index % 2 else "SELL")
        source["decisions_by_mode"]["REAL"]["IDC"]["action"] = "BUY"
        before = deepcopy(source)
        DashboardActionsMixin._capture_signal_trace(view, source, "RUNNING")
        assert source == before
    assert attempted == [] and warnings == []
    assert len(SignalTraceStore(tmp_path / "signal_trace.sqlite3").read()) == 32
