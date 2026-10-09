from copy import deepcopy
from types import SimpleNamespace

import customtkinter as ctk
import pytest

from viking_v2.config import AppSettings
from viking_v2.dashboard.windows import HistoryPopup, SignalRecordingPopup
from viking_v2.services.indicator_comparison import IndicatorComparisonStore
from viking_v2.services.signal_history import (
    SignalHistoryTrash, observation_id, periodic_history_row, recording_options,
)
from viking_v2.storage import SignalLog


def event(clock="13:00:00", **changes):
    return dict(timestamp="2026-10-09 " + clock, symbol="HDB", execution_mode="REAL",
                signal="BUY", acted="WAIT", blocked_by="BUY_WINDOW_WAIT",
                record_kind="SIGNAL_EVENT", signal_event="ENTRY", signal_cycle="C1",
                ema_fast=22.45, ema_slow=22.40, rsi=58, rsi_previous=56.8,
                price=22.5, slot_usage="0/4", watchlist_priority=3, **changes)


def sample(clock="14:00:01", symbol="HDB"):
    return dict(timestamp="2026-10-09T" + clock + "+07:00", scheduled_at="2026-10-09T" + clock[:5],
                symbol=symbol, execution_mode="REAL", priority=3, price_vnd=22500,
                ema_fast=22.45, ema_slow=22.40, rsi=58, rsi_previous=56.8,
                rsi_previous_date="2026-10-08", entry=True, exit_e=False, slot_usage="0/4",
                ema_cross_required=True, ema_cross_state="WAIT_DOWN", reason="NO_NEW_BUY_SIGNAL")


def popup_for(ui_root, tmp_path, rows=None, samples=None):
    trace = SimpleNamespace(path=tmp_path / "trace.sqlite3", read=lambda **_kw: deepcopy(samples or []))
    return HistoryPopup(ui_root, lambda _mode: [], signals_provider=lambda: rows or [], trace_store=trace,
                        indicator_comparison=IndicatorComparisonStore(tmp_path / "comparison"))


def test_trash_hides_only_ids_persists_and_can_restore_without_modifying_records(tmp_path):
    log = SignalLog(tmp_path / "signals.csv")
    log.observe(event(), entry_condition=True, exit_condition=False)
    before = {path: path.read_bytes() for path in tmp_path.iterdir() if path.is_file()}
    row = log.read_all()[0]
    trash = SignalHistoryTrash(tmp_path / "signal_history_trash.json")
    assert trash.hide([row]) == 1
    assert SignalHistoryTrash(trash.store.path).visible([row]) == []
    assert all(path.read_bytes() == data for path, data in before.items())
    assert "HDB" not in trash.store.path.read_text(encoding="utf-8")
    # Hiding an observation does not rearm the backend event/Telegram dedup.
    assert not log.observe(event("13:01:00"), entry_condition=True, exit_condition=False)
    assert trash.restore() == 1 and trash.visible([row]) == [row]


def test_periodic_identity_is_book_symbol_and_schedule_not_display_numbers():
    first = periodic_history_row(sample())
    assert observation_id(first) == observation_id({**first, "ema_fast": 99, "rsi": 99})
    assert observation_id(first) != observation_id(periodic_history_row(sample("14:02:01")))
    assert observation_id(first) != observation_id({**first, "execution_mode": "PAPER"})
    assert first["signal"] == "PERIODIC" and first["acted"] == ""


def test_malformed_trash_metadata_does_not_break_history(tmp_path):
    trash = SignalHistoryTrash(tmp_path / "trash.json")
    trash.store.write({"hidden": None})
    assert trash.visible([event()]) == [event()]


@pytest.mark.parametrize("interval,start,end", [
    (0, "14:00", "14:30"), (31, "14:00", "14:30"), ("2.5", "14:00", "14:30"),
    (2, "25:00", "14:30"), (2, "14:00", "14:60"), (2, "14:30", "14:00"),
])
def test_recording_options_reject_invalid_values(interval, start, end):
    with pytest.raises(ValueError):
        recording_options(True, interval, start, end)


def test_unified_tree_keeps_every_sample_orders_mixed_timestamps_and_short_reasons(ui_root, tmp_path):
    rows = [event("14:01:00"), event("13:30:00")]
    popup = popup_for(ui_root, tmp_path, rows, [sample(), sample("14:02:01"), sample(symbol="IDC")])
    try:
        tree = popup.signal_tree
        hdb = tree.get_children("symbol:2026-10-09:HDB")
        assert [tree.item(iid, "text").strip() for iid in hdb] == ["14:02:01", "14:01:00", "14:00:01", "13:30:00"]
        assert [tree.set(iid, "display_signal") for iid in hdb] == ["ĐỊNH KỲ", "ENTRY", "ĐỊNH KỲ", "ENTRY"]
        assert all(tree.set(iid, "rsi_comparison") == "58.00 > 56.80" for iid in hdb)
        assert "2 sự kiện · 2 mẫu" in tree.item("symbol:2026-10-09:HDB", "text")
        assert tree.set("day:2026-10-09", "reason") == ""
        assert tree.bind("<Button-3>") and tree.bind("<Delete>")
        assert str(tree.cget("selectmode")) == "extended"
        assert not hasattr(popup, "trace_tree") and not hasattr(popup, "signal_view_button")
        # Selection survives refresh, not just the expanded day/symbol state.
        tree.selection_set(hdb[0])
        popup._refresh_signals()
        assert tree.selection() == (hdb[0],)
    finally:
        popup.close()


@pytest.mark.parametrize("scope", ["row", "symbol", "day"])
def test_context_delete_is_confirmed_scoped_recoverable_and_export_omits_trash(ui_root, tmp_path, monkeypatch, scope):
    from openpyxl import load_workbook
    rows = [event(), event("13:05:00")]
    raw = deepcopy(rows)
    popup = popup_for(ui_root, tmp_path, rows, [sample(), sample(symbol="IDC")])
    try:
        tree = popup.signal_tree
        hdb = "symbol:2026-10-09:HDB"
        selected = tree.get_children(hdb)[0] if scope == "row" else hdb if scope == "symbol" else "day:2026-10-09"
        tree.selection_set(selected)
        count = len(popup._selected_signal_rows())
        monkeypatch.setattr("viking_v2.dashboard.windows.messagebox.askyesno", lambda *_a, **_kw: False)
        popup._delete_signal_rows()
        assert popup.history_trash.hidden_ids() == set()
        monkeypatch.setattr("viking_v2.dashboard.windows.messagebox.askyesno", lambda *_a, **_kw: True)
        popup._delete_signal_rows()
        assert len(popup._visible_history_sources) == 4 - count
        assert rows == raw
        if scope != "day":
            path = tmp_path / "out.xlsx"
            monkeypatch.setattr("viking_v2.dashboard.windows.filedialog.asksaveasfilename", lambda **_kw: str(path))
            popup._export_signals()
            book = load_workbook(path)
            try:
                assert book["DNSE GỐC"].max_row == 1 + 4 - count
                assert book["TRACE"].max_row == 2
            finally:
                book.close()
        popup._restore_signal_rows()
        assert len(popup._visible_history_sources) == 4
        assert popup.history_trash.hidden_ids() == set()
    finally:
        popup.close()


def test_delete_legacy_aggregate_hides_all_repeats_but_not_new_cycle(ui_root, tmp_path, monkeypatch):
    rows = [{**event(), "signal": "SELL", "record_kind": "", "signal_event": "", "signal_cycle": ""},
            {**event("13:01:00"), "signal": "SELL", "record_kind": "", "signal_event": "", "signal_cycle": ""}]
    popup = popup_for(ui_root, tmp_path, rows)
    try:
        tree = popup.signal_tree
        leaf = tree.get_children("symbol:2026-10-09:HDB")[0]
        assert "×2 cũ" in tree.item(leaf, "text")
        tree.selection_set(leaf)
        monkeypatch.setattr("viking_v2.dashboard.windows.messagebox.askyesno", lambda *_a, **_kw: True)
        popup._delete_signal_rows()
        assert not tree.get_children() and len(popup.history_trash.hidden_ids()) == 2
        rows.append(event("14:10:00"))
        popup._refresh_signals()
        assert len(tree.get_children("symbol:2026-10-09:HDB")) == 1
    finally:
        popup.close()


def test_delete_identity_remains_original_when_switching_to_tradingview(ui_root, tmp_path, monkeypatch):
    popup = popup_for(ui_root, tmp_path, [event()], [sample()])
    try:
        popup._change_indicator_basis("TRADINGVIEW")
        popup.signal_tree.selection_set("symbol:2026-10-09:HDB")
        monkeypatch.setattr("viking_v2.dashboard.windows.messagebox.askyesno", lambda *_a, **_kw: True)
        popup._delete_signal_rows()
        popup._change_indicator_basis("DNSE")
        assert not popup.signal_tree.get_children()
        popup._restore_signal_rows()
        assert len(popup.signal_tree.get_children("symbol:2026-10-09:HDB")) == 2
    finally:
        popup.close()


def test_copy_and_details_use_selected_observation_and_do_not_include_secrets(ui_root, tmp_path, monkeypatch):
    popup = popup_for(ui_root, tmp_path, samples=[sample()])
    try:
        tree = popup.signal_tree
        tree.selection_set(tree.get_children("symbol:2026-10-09:HDB")[0])
        copied = []
        monkeypatch.setattr(popup.top, "clipboard_clear", lambda: None)
        monkeypatch.setattr(popup.top, "clipboard_append", copied.append)
        popup._copy_signal_rows()
        assert "ĐỊNH KỲ" in copied[0] and "58.00 > 56.80" in copied[0]
        shown = []
        monkeypatch.setattr("viking_v2.dashboard.windows.messagebox.showinfo", lambda *_a, **_kw: shown.append(_a))
        popup._show_signal_details()
        assert "RSI 58.00 > 56.80" in shown[0][1] and "WHIPSAW" in shown[0][1]
    finally:
        popup.close()


@pytest.mark.parametrize("scaling", [1.0, 1.5])
def test_single_settings_button_layout_and_small_form_fit_at_remote_scaling(ui_root, tmp_path, scaling):
    ctk.set_widget_scaling(scaling)
    configured = AppSettings()
    before = configured.to_dict()
    def save(options):
        for key, value in options.items():
            setattr(configured, key, value)
    popup = popup_for(ui_root, tmp_path, [event()], [sample(), sample("14:02:01"), sample(symbol="IDC")])
    popup.on_trace_settings = save
    popup.trace_settings_provider = lambda: configured
    try:
        popup.top.geometry("980x560")
        popup.tabs.set("TÍN HIỆU")
        popup._open_recording_settings()
        form = popup.recording_popup
        ui_root.update_idletasks()
        assert form.save_button.winfo_y() + form.save_button.winfo_height() <= form.save_button.master.winfo_height()
        assert form.save_button.winfo_rooty() + form.save_button.winfo_height() <= form.top.winfo_rooty() + form.top.winfo_height()
        toolbar = popup.recording_button.master
        buttons = [child for child in toolbar.winfo_children() if child.grid_info().get("row") == 1]
        assert all(child.winfo_x() + child.winfo_width() <= toolbar.winfo_width() for child in buttons)
        assert len(form.entries) == 3
        _capture_ui(ui_root, form.top, f"signal-recording-{scaling}")
        form.enabled.set(False)
        form.entries["interval"].delete(0, "end")
        form.entries["interval"].insert(0, "3")
        form.save()
        assert configured.to_dict() == {**before, "signal_trace_enabled": False, "signal_trace_interval_minutes": 3}
        popup.top.geometry("1600x820+0+0")
        popup.signal_tree.item("symbol:2026-10-09:HDB", open=True)
        _capture_ui(ui_root, popup.top, f"signal-history-{scaling}")
    finally:
        if hasattr(popup, "recording_popup") and popup.recording_popup.top.winfo_exists():
            popup.recording_popup.top.destroy()
        popup.close()
        ctk.set_widget_scaling(1.0)


def test_quick_tab_changes_do_not_let_stale_ctk_callback_hide_signals(ui_root, tmp_path):
    import tkinter as tk
    popup = popup_for(ui_root, tmp_path, [event()])
    try:
        for name in ("TÍN HIỆU", "CKCS REAL", "CKCS PAPER", "TÍN HIỆU"):
            popup.tabs.set(name)
        popup._history_tab_changed()
        settled = tk.BooleanVar(master=ui_root, value=False)
        ui_root.after(150, lambda: settled.set(True))
        ui_root.wait_variable(settled)
        assert popup.tabs.tab("TÍN HIỆU").winfo_ismapped()
        assert popup.signal_tree.winfo_ismapped()
        assert "mã → giờ" in popup.subtitle.cget("text")
    finally:
        popup.close()


def _capture_ui(root, top, name):
    """Optional review artifacts from fixtures only, never the live account UI."""
    import os
    if os.getenv("VIKING_CAPTURE_UI") != "1":
        return
    import ctypes
    from pathlib import Path
    import tkinter as tk
    from PIL import ImageGrab
    root.update_idletasks()
    top.lift()
    settled = tk.BooleanVar(master=root, value=False)
    root.after(180, lambda: settled.set(True))
    root.wait_variable(settled)
    directory = Path(__file__).resolve().parents[2] / ".artifacts" / "ui-review"
    directory.mkdir(parents=True, exist_ok=True)
    ImageGrab.grab(window=ctypes.windll.user32.GetParent(top.winfo_id())).save(directory / f"{name}.png")


def test_recording_callback_updates_only_four_settings_and_ignores_unrelated_payload(tmp_path, monkeypatch):
    from viking_v2.dashboard.actions import DashboardActionsMixin
    import viking_v2.dashboard.actions as actions
    configured = AppSettings()
    before = configured.to_dict()
    saved = []
    monkeypatch.setattr(actions.config, "load_settings", lambda _account: deepcopy(configured))
    monkeypatch.setattr(actions, "save_settings", lambda value, account: saved.append((value, account)))
    view = SimpleNamespace(account_id="OFFLINE_TEST", settings=configured)
    values = {**recording_options(False, 3, "14:05", "14:25"),
              "rule_parameters": {}, "telegram_enabled": False, "bot_enabled": False}
    DashboardActionsMixin._save_signal_recording_settings(view, values)
    assert view.settings.to_dict() == {**before, **recording_options(False, 3, "14:05", "14:25")}
    assert len(saved) == 1 and saved[0][1] == "OFFLINE_TEST"
    with pytest.raises(ValueError):
        DashboardActionsMixin._save_signal_recording_settings(view, {**values, "signal_trace_interval_minutes": 0})
    assert len(saved) == 1
