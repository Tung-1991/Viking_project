from __future__ import annotations

import json
import os
import pytest

from viking_v2 import config
from viking_v2.config import AppSettings
from viking_v2.rules.business import StaticRuleParameters
from viking_v2.connections.window import ConnectionPopup
from viking_v2.dashboard.windows import parse_symbols
from viking_v2.config import update_env


def test_v2_settings_only_keep_two_bot_execution_modes():
    assert AppSettings(bot_order_mode="lo_local").normalize().bot_order_mode == "LO_LOCAL"
    assert AppSettings(bot_order_mode="resting-lo").normalize().bot_order_mode == "MARKET"


def test_execution_defaults_are_explicit_and_minimal():
    settings = AppSettings().normalize()
    assert settings.daily_stats_mode == "DAILY"
    assert settings.daily_stats_reset_time == "00:00"
    assert settings.skip_order_popups is True
    assert settings.bot_sl_enabled is True
    assert settings.bot_em_modes == ["NORMAL", "IND_EXIT"]
    assert settings.sell_wait_policy == "RECHECK"
    assert settings.signal_mode == "REALTIME"
    assert AppSettings(sell_wait_policy="keep").normalize().sell_wait_policy == "KEEP"
    assert AppSettings(sell_wait_policy="unknown").normalize().sell_wait_policy == "RECHECK"
    assert StaticRuleParameters().whipsaw_n == 3
    assert StaticRuleParameters().whipsaw_x == 7
    assert StaticRuleParameters().loss_lock_hours == 24
    assert settings.manual_sell_pause_minutes == 15
    assert StaticRuleParameters().initial_sl_pct == -3.5
    assert StaticRuleParameters().reentry_sl_pct == -2.5
    assert settings.rule_parameters == config.DEFAULT_RULE_PARAMETERS
    assert settings.rule_parameters is not config.DEFAULT_RULE_PARAMETERS


@pytest.mark.parametrize("legacy", [True, False])
def test_old_confirmation_setting_uses_new_skip_default_without_changing_other_settings(legacy):
    raw = AppSettings(telegram_enabled=True, telegram_chat_id="123", watchlist=["MSN"]).to_dict()
    raw.pop("skip_order_popups")
    raw["confirm_real_orders"] = legacy
    settings = AppSettings.from_dict(raw)
    assert settings.skip_order_popups is True
    assert settings.telegram_enabled is True and settings.telegram_chat_id == "123"
    assert settings.watchlist == ["MSN"]
    assert "confirm_real_orders" not in settings.to_dict()


@pytest.mark.parametrize("skip", [True, False])
def test_confirmation_preference_survives_save_and_reload(monkeypatch, tmp_path, skip):
    monkeypatch.setattr(config, "ACCOUNTS_ROOT", tmp_path / "accounts")
    settings = AppSettings(skip_order_popups=skip, watchlist=["MSN"])
    config.save_settings(settings, "CONFIRM_TEST")
    assert config.load_settings("CONFIRM_TEST").skip_order_popups is skip
    assert AppSettings.from_dict({"skip_order_popups": skip, "confirm_real_orders": True}).skip_order_popups is skip


@pytest.mark.parametrize("skip", [True, False])
def test_confirmation_only_preference_migrates_to_full_order_popup_preference(skip):
    settings = AppSettings.from_dict({"skip_real_order_confirmation": skip})
    assert settings.skip_order_popups is skip
    assert "skip_real_order_confirmation" not in settings.to_dict()
    assert AppSettings.from_dict({"skip_real_order_confirmation": not skip, "skip_order_popups": skip}).skip_order_popups is skip


@pytest.mark.parametrize("malformed", ["false", "true", 1, 0, None, [], {}])
def test_malformed_confirmation_flag_keeps_confirmation_required(malformed):
    assert AppSettings.from_dict({"skip_order_popups": malformed}).skip_order_popups is False


@pytest.mark.parametrize("skip", [True, False])
def test_rule_ui_confirmation_switch_has_clear_label_and_saves(ui_root, monkeypatch, tmp_path, skip):
    import customtkinter as ctk
    from viking_v2.rules.window import RuleSettingsPopup

    monkeypatch.setattr(config, "ACCOUNTS_ROOT", tmp_path / "accounts")
    popup = RuleSettingsPopup(
        ui_root, AppSettings(skip_order_popups=skip), "CONFIRM_UI", lambda: None,
    )
    try:
        variable = popup.skip_order_popups
        assert variable.get() is skip
        switches = [child for child in variable._viking_row.winfo_children() if isinstance(child, ctk.CTkSwitch)]
        assert len(switches) == 1
        assert switches[0].cget("text") == "BỎ POPUP ĐẶT LỆNH"
        variable.set(not skip)
        popup.save()
        assert "ĐÃ LƯU" in popup.status.cget("text"), popup.status.cget("text")
        assert config.load_settings("CONFIRM_UI").skip_order_popups is (not skip)
    finally:
        popup._close()


def test_dashboard_stats_mode_is_normalized():
    assert AppSettings(daily_stats_mode="since_reset").normalize().daily_stats_mode == "SINCE_RESET"
    assert AppSettings(daily_stats_mode="bad").normalize().daily_stats_mode == "DAILY"
    assert AppSettings(daily_stats_reset_time="15:30").normalize().daily_stats_reset_time == "15:30"
    assert AppSettings(daily_stats_reset_time="25:00").normalize().daily_stats_reset_time == "00:00"


def test_static_rule_fallback_matches_the_single_canonical_default_source():
    assert StaticRuleParameters().to_dict() == config.DEFAULT_RULE_PARAMETERS
    assert StaticRuleParameters.from_dict({}).to_dict() == config.DEFAULT_RULE_PARAMETERS


def test_sparse_account_rules_merge_over_complete_operating_defaults():
    settings = AppSettings(rule_parameters={"loss_lock_hours": 12}).normalize()
    assert settings.rule_parameters["loss_lock_hours"] == 12
    assert settings.rule_parameters["initial_sl_pct"] == -3.5
    assert settings.rule_parameters["reentry_sl_pct"] == -2.5
    assert settings.rule_parameters["normal_atr_activation_multiplier"] == 0.55
    assert settings.rule_parameters["normal_dynamic_enabled"] is True
    assert settings.manual_sell_pause_minutes == 15


def test_phase1_override_settings_are_normalized():
    settings = AppSettings(
        market_phase_override_enabled=1,
        market_phase_override="bad",
        market_phase_override_exposure_pct=150,
    ).normalize()
    assert settings.market_phase_override_enabled is True
    assert settings.market_phase_override == "ACCUMULATION"
    assert settings.market_phase_override_exposure_pct == 100


def test_default_watchlist_keeps_the_40_legacy_ckcs_symbols(monkeypatch):
    monkeypatch.delenv("DNSE_CKCS_WATCHLIST", raising=False)
    symbols = config._watchlist_from_env()
    assert len(symbols) == 40
    assert symbols == list(config.DEFAULT_CKCS_WATCHLIST)
    assert "FPT" not in symbols


def test_priority_symbols_are_normalized_as_watchlist_subset():
    settings = AppSettings(
        watchlist=["fpt", "mbb"],
        priority_symbols=["MBB", "vcb", "mbb"],
    ).normalize()

    assert settings.watchlist == ["FPT", "MBB"]
    assert settings.priority_symbols == ["MBB"]


def test_telegram_settings_only_keep_connection_values():
    settings = AppSettings(telegram_enabled=True, telegram_chat_id=" 123 ").normalize()
    assert settings.telegram_enabled is True
    assert settings.telegram_chat_id == "123"
    assert settings.telegram_token_env == "TELE_BOT_KEY"
    assert settings.telegram_buy_batch_minutes == 30
    assert settings.telegram_notifications["protect"] is False
    assert settings.telegram_notifications["indicator_exit"] is True
    assert settings.telegram_notifications["blocked_buy"] is False
    assert settings.telegram_notifications["system"] is True
    assert settings.telegram_cooldown_minutes["indicator_exit"] == 30
    assert settings.telegram_cooldown_minutes["system"] == 15
    assert not hasattr(settings, "telegram_system_alerts")


def test_legacy_telegram_alert_switch_migrates_to_explicit_categories():
    settings = AppSettings.from_dict({"telegram_signal_alerts": True})

    assert settings.telegram_notifications["protect"] is True
    assert settings.telegram_notifications["indicator_exit"] is True
    assert settings.telegram_notifications["blocked_buy"] is True
    assert settings.telegram_notifications["system"] is True
    assert not hasattr(settings, "telegram_signal_alerts")


def test_official_2026_exchange_holidays_are_always_applied():
    settings = AppSettings(custom_holidays=["2026-08-13", "2026-09-02"]).normalize()

    assert "2026-01-02" in settings.trading_holidays
    assert "2026-02-16" in settings.trading_holidays
    assert "2026-09-02" in settings.trading_holidays
    assert "2026-08-13" in settings.trading_holidays
    assert settings.custom_holidays == ["2026-08-13"]


def test_rule_parameter_overrides_are_single_flat_source():
    params = StaticRuleParameters.from_dict({"ema_fast": 4, "whipsaw_n": 5, "unknown": 999})
    assert params.ema_fast == 4
    assert params.buy_ema_fast == 4
    assert params.sell_ema_fast == 4
    assert params.buy_ema_slow == 6
    assert params.sell_ema_slow == 6
    assert params.whipsaw_n == 5
    assert not hasattr(params, "unknown")


def test_rule_parameter_supports_independent_buy_and_sell_ema_pairs():
    params = StaticRuleParameters.from_dict(
        {
            "buy_ema_fast": 3,
            "buy_ema_slow": 6,
            "sell_ema_fast": 5,
            "sell_ema_slow": 10,
        }
    )

    assert params.buy_ema_fast == 3
    assert params.buy_ema_slow == 6
    assert params.sell_ema_fast == 5
    assert params.sell_ema_slow == 10
    saved = params.to_dict()
    assert "ema_fast" not in saved
    assert saved["buy_ema_fast"] == 3
    assert saved["sell_ema_slow"] == 10


def test_v2_account_workspace_is_self_contained_and_safe(tmp_path, monkeypatch):
    accounts = tmp_path / "viking_v2" / "runtime" / "accounts"
    monkeypatch.setattr(config, "ACCOUNTS_ROOT", accounts)

    path = config.account_root("../ACC 01")

    assert path.parent == accounts
    assert path.name == "_ACC_01"
    assert ".." not in path.parts


def test_missing_account_settings_are_materialized_as_json(tmp_path, monkeypatch):
    accounts = tmp_path / "runtime" / "accounts"
    monkeypatch.setattr(config, "ACCOUNTS_ROOT", accounts)

    settings = config.load_settings("ACC01")
    saved = accounts / "ACC01" / "settings.json"

    assert saved.is_file()
    assert json.loads(saved.read_text(encoding="utf-8"))["watchlist"] == settings.watchlist


def test_saved_json_rule_overrides_win_over_config_defaults(tmp_path, monkeypatch):
    accounts = tmp_path / "runtime" / "accounts"
    monkeypatch.setattr(config, "ACCOUNTS_ROOT", accounts)

    settings = AppSettings().normalize()
    settings.rule_parameters["initial_sl_pct"] = -4.25
    settings.rule_parameters["normal_atr_activation_multiplier"] = 0.72
    saved = config.save_settings(settings, "ACC01")

    raw = json.loads(saved.read_text(encoding="utf-8"))
    loaded = config.load_settings("ACC01")

    assert raw["rule_parameters"]["initial_sl_pct"] == -4.25
    assert raw["rule_parameters"]["normal_atr_activation_multiplier"] == 0.72
    assert loaded.rule_parameters["initial_sl_pct"] == -4.25
    assert loaded.rule_parameters["normal_atr_activation_multiplier"] == 0.72
    assert loaded.rule_parameters["normal_atr_multiplier"] == 0.8
    assert config.DEFAULT_RULE_PARAMETERS["initial_sl_pct"] == -3.5


def test_connection_popup_normalizes_watchlist_and_dnse_account_payload():
    symbols, invalid = parse_symbols("fpt, SSI  fpt;FUEVFVND bad!")
    rows, custody, owner = ConnectionPopup._account_payload(
        {"data": {"accounts": [{"id": "001", "dealAccount": True}], "custodyCode": "C01", "name": "K"}}
    )

    assert symbols == ["FPT", "SSI", "FUEVFVND"]
    assert invalid == ["BAD!"]
    assert rows == [{"id": "001", "dealAccount": True}]
    assert (custody, owner) == ("C01", "K")
    assert ConnectionPopup._dnse_error_message("Authorization field missing, malformed or invalid") == (
        "XÁC THỰC API KHÔNG HỢP LỆ · KIỂM TRA API KEY VÀ API SECRET"
    )


def test_connection_popup_parses_paper_balance_in_common_formats():
    assert ConnectionPopup._parse_paper_balance("100,000,000") == 100_000_000
    assert ConnectionPopup._parse_paper_balance("100.000.000") == 100_000_000
    assert ConnectionPopup._parse_paper_balance("100000000") == 100_000_000


def test_env_store_can_remove_only_selected_secret(tmp_path, monkeypatch):
    path = tmp_path / ".env"
    path.write_text("DNSE_API_KEY=keep\nDNSE_TRADING_TOKEN=remove\nDNSE_TRADING_TOKEN_EXPIRES_AT=123\n", encoding="utf-8")
    monkeypatch.setenv("DNSE_TRADING_TOKEN", "remove")
    monkeypatch.setenv("DNSE_TRADING_TOKEN_EXPIRES_AT", "123")

    update_env(
        {
            "DNSE_TRADING_TOKEN": None,
            "DNSE_TRADING_TOKEN_EXPIRES_AT": None,
        },
        path,
    )

    assert path.read_text(encoding="utf-8") == "DNSE_API_KEY=keep\n"
    assert "DNSE_TRADING_TOKEN" not in os.environ
    assert "DNSE_TRADING_TOKEN_EXPIRES_AT" not in os.environ


def test_no_two_cards_ever_share_a_grid_cell(ui_root):
    """A card gridded onto an occupied cell is drawn over and never seen.

    This has bitten twice already, so the layout is checked instead of eyeballed.
    """
    import customtkinter as ctk
    from viking_v2.config import load_settings
    from viking_v2.rules.window import RuleSettingsPopup
    from viking_v2.backtest.window import BacktestPopup

    root = ui_root
    try:
        clashes = []

        def scan(label, widget):
            # CustomTkinter stacks a canvas and a raw label inside its own
            # widgets on purpose, so only real containers are checked.
            if not isinstance(widget, ctk.CTkFrame):
                return
            taken: dict[tuple[int, int], str] = {}
            for child in widget.winfo_children():
                info = child.grid_info()
                if not info or not isinstance(child, ctk.CTkBaseClass | ctk.CTkFrame):
                    continue
                row, column = int(info.get("row", 0)), int(info.get("column", 0))
                span = int(info.get("columnspan", 1))
                for offset in range(span):
                    cell = (row, column + offset)
                    if cell in taken:
                        clashes.append(f"{label} ô {cell}")
                    taken[cell] = str(child)

        rules = RuleSettingsPopup(root, load_settings("PAPER"), "PAPER", lambda: None)
        for name in ("NGHIỆP VỤ", "E/M", "THỰC THI"):
            rules.tabs.set(name)
            root.update_idletasks()
            for frame in rules.tabs.tab(name).winfo_children():
                scan(f"RULE/{name}", frame)

        back = BacktestPopup(root, load_settings("PAPER"), None)
        for name in ("MODE 1", "MODE 2", "THAM SỐ"):
            back.tabs.set(name)
            root.update_idletasks()
            for frame in back.tabs.tab(name).winfo_children():
                scan(f"BACKTEST/{name}", frame)

        assert not clashes, "thẻ chồng lên nhau: " + ", ".join(clashes)
    finally:
        if 'back' in locals() and back.top.winfo_exists():
            back.close()
        if 'rules' in locals() and rules.top.winfo_exists():
            rules._close()


def test_exit_sell_controls_are_in_em_card_and_save_separately(ui_root, monkeypatch):
    from viking_v2.config import load_settings
    from viking_v2.rules import window as rule_window

    settings = load_settings("PAPER")
    saved = []
    monkeypatch.setattr(
        rule_window, "save_settings",
        lambda current, account_id: saved.append((current, account_id)),
    )
    rules = rule_window.RuleSettingsPopup(ui_root, settings, "PAPER", lambda: None)
    try:
        assert rules.exit_card.master is rules.exit_left_column
        assert rules.exit_card.master.master is rules.protect_card
        assert int(rules.exit_card.grid_info()["column"]) == 2
        assert rules.sell_ema_fast.master.master is rules.exit_card
        assert rules.sell_ema_slow.master.master is rules.exit_card
        assert rules.buy_rsi_period.get() == rules.rsi_period.get()
        assert rules.dynamic_settings_card.master is rules.exit_card.master
        assert rules.normal_arm.master.master is rules.protect_trail_card
        assert rules.normal_giveback.master.master is rules.protect_trail_card
        assert rules.normal_sell.master.master is rules.protect_trail_card
        for entry in (
            rules.normal_atr_activation_multiplier,
            rules.normal_atr_multiplier,
        ):
            assert entry.master.master is rules.dynamic_atr_card
        for entry in (
            rules.normal_retention_pct,
            rules.normal_retention_until_pct,
        ):
            assert entry.master.master is rules.dynamic_retention_card
        assert rules.dynamic_atr_card.master.master is rules.dynamic_settings_card
        assert rules.dynamic_retention_card.master.master is rules.dynamic_settings_card
        for key in (
            "normal_atr_activation_enabled", "normal_atr_trail_enabled",
            "normal_retention_enabled", "normal_retention_until_enabled",
        ):
            getattr(rules, key).set(False)

        original_buy_ema = (rules.params.buy_ema_fast, rules.params.buy_ema_slow)
        rules.sell_ema_fast.delete(0, "end")
        rules.sell_ema_fast.insert(0, "4")
        rules.sell_ema_slow.delete(0, "end")
        rules.sell_ema_slow.insert(0, "8")
        rules.sell_signal_rsi.set(False)
        rules.shared_rsi_period.set("15")
        assert rules.buy_rsi_period.get() == rules.rsi_period.get() == "15"

        rules.save()
        assert saved and saved[0][1] == "PAPER"
        values = saved[0][0].rule_parameters
        assert (values["buy_ema_fast"], values["buy_ema_slow"]) == original_buy_ema
        assert (values["sell_ema_fast"], values["sell_ema_slow"]) == (4, 8)
        assert values["sell_signal_use_rsi"] is False
        assert values["rsi_period"] == 15
        for key in (
            "normal_atr_activation_enabled", "normal_atr_trail_enabled",
            "normal_retention_enabled", "normal_retention_until_enabled",
        ):
            assert values[key] is False
    finally:
        if rules.top.winfo_exists():
            rules._close()


def test_backtest_exit_sell_controls_are_separate_and_round_trip(ui_root, monkeypatch):
    from viking_v2.backtest.window import BacktestPopup
    from viking_v2.config import load_settings

    back = BacktestPopup(ui_root, load_settings("PAPER"), None)
    try:
        monkeypatch.setattr(back.data, "save_settings", lambda _values: None)
        assert "sell_ema_fast" in back._rule_entries
        assert "sell_ema_slow" in back._rule_entries
        assert back.exit_card.winfo_exists()

        for variable in back.dynamic_subrules.values():
            variable.set(False)
        result = back._collect()
        assert len(back.dynamic_subrules) == 4
        assert all(result.rule_parameters[key] is False for key in back.dynamic_subrules)
    finally:
        if back.top.winfo_exists():
            back.close()


def test_no_tab_is_wider_than_the_window_it_lives_in(ui_root):
    """A tab that needs more width than it has gets its right edge cut off.

    Both popups sized themselves past the window twice already, hiding a whole
    switch, so the widths are measured instead of eyeballed.
    """
    import customtkinter as ctk
    from viking_v2.config import load_settings
    from viking_v2.rules.window import RuleSettingsPopup
    from viking_v2.backtest.window import BacktestPopup

    root = ui_root
    try:
        too_wide = []
        settings = load_settings("PAPER")
        pages = (
            ("RULE", RuleSettingsPopup(root, settings, "PAPER", lambda: None),
             ("NGHIỆP VỤ", "E/M", "THỰC THI")),
            ("BACKTEST", BacktestPopup(root, settings, None),
             ("MODE 1", "MODE 2", "THAM SỐ")),
        )
        for label, popup, names in pages:
            for name in names:
                popup.tabs.set(name)
                root.update_idletasks()
                popup.top.update()
                tab = popup.tabs.tab(name)
                # MODE 2 holds a table that scrolls sideways on purpose.
                if name == "MODE 2":
                    continue
                if tab.winfo_reqwidth() > tab.winfo_width():
                    too_wide.append(
                        f"{label}/{name} cần {tab.winfo_reqwidth()} có {tab.winfo_width()}"
                    )
        assert not too_wide, "tab tràn khỏi cửa sổ: " + ", ".join(too_wide)
    finally:
        for _label, popup, _names in locals().get("pages", ()):
            closer = getattr(popup, "close", None) or getattr(popup, "_close", None)
            if callable(closer) and popup.top.winfo_exists():
                closer()


def test_signal_log_records_a_change_not_every_loop(tmp_path):
    """The scanner evaluates every symbol every second; only changes are worth keeping."""
    from viking_v2.storage import SignalLog

    log = SignalLog(tmp_path / "signal_log.csv")
    base = dict(price=21.5, ema_fast=21.4, ema_slow=21.2, rsi=58.3,
                market_state="UPTREND", acted="WAIT", blocked_by="MAX_POSITIONS")

    assert log.record({**base, "timestamp": "10:00", "symbol": "HSG", "signal": "BUY"})
    assert not log.record({**base, "timestamp": "10:01", "symbol": "HSG", "signal": "BUY"})
    assert not log.record({**base, "timestamp": "10:02", "symbol": "HSG", "signal": "BUY"})
    # Losing the signal is a change, but an empty signal is not worth a row.
    assert not log.record({**base, "timestamp": "14:00", "symbol": "HSG", "signal": ""})
    assert log.record({**base, "timestamp": "14:30", "symbol": "HSG", "signal": "BUY"})
    assert log.record({**base, "timestamp": "10:00", "symbol": "FPT", "signal": "SELL"})

    rows = log.read_all()
    assert [(r["symbol"], r["signal"], r["timestamp"]) for r in rows] == [
        ("HSG", "BUY", "10:00"), ("HSG", "BUY", "14:30"), ("FPT", "SELL", "10:00"),
    ]
    # The reason the bot stood still is the whole point of keeping this.
    assert rows[0]["blocked_by"] == "MAX_POSITIONS"


def test_signal_log_records_the_final_arbitration_outcome(tmp_path):
    from viking_v2.storage import SignalLog

    log = SignalLog(tmp_path / "signal_log.csv")
    common = {
        "timestamp": "2026-09-09 10:00:00", "execution_mode": "PAPER",
        "symbol": "FPT", "signal": "BUY", "candle_key": "D1",
        "signal_cycle": "D1|2026-09-09T09:45:00+07:00",
        "signal_time": "2026-09-09T09:45:00+07:00", "watchlist_priority": 2,
        "slot_usage": "5/5",
    }
    assert log.record({**common, "acted": "WAIT", "blocked_by": "MAX_POSITIONS"})
    assert log.record({**common, "acted": "BUY", "blocked_by": "", "slot_usage": "4/5"})
    rows = log.read_all()
    assert [(row["acted"], row["blocked_by"]) for row in rows] == [
        ("WAIT", "MAX_POSITIONS"), ("BUY", ""),
    ]
    assert rows[-1]["execution_mode"] == "PAPER"
    assert rows[-1]["candle_key"] == "D1"
    assert rows[-1]["signal_cycle"] == "D1|2026-09-09T09:45:00+07:00"
    assert rows[-1]["watchlist_priority"] == "2"


def test_signal_log_dedupe_keeps_paper_and_real_independent(tmp_path):
    from viking_v2.storage import SignalLog

    log = SignalLog(tmp_path / "signal_log.csv")
    common = {
        "timestamp": "2026-09-09 10:00:00", "symbol": "FPT",
        "signal": "BUY", "candle_key": "D1", "acted": "WAIT",
        "blocked_by": "BOT_OFF",
    }
    assert log.record({**common, "execution_mode": "PAPER"})
    assert log.record({**common, "execution_mode": "REAL"})
    assert [row["execution_mode"] for row in log.read_all()] == ["PAPER", "REAL"]


def test_signal_history_groups_detailed_rows_by_day_and_hides_restart_duplicates():
    from viking_v2.dashboard.windows import signal_rows_by_day

    def row(timestamp, symbol, *, acted="BUY", blocked=""):
        return {
            "timestamp": timestamp,
            "symbol": symbol,
            "signal": "BUY",
            "price": "10.5",
            "ema_fast": "10.2",
            "ema_slow": "10.1",
            "rsi": "55",
            "market_state": "ACCUMULATION",
            "acted": acted,
            "blocked_by": blocked,
        }

    rows = [
        row("2026-08-21 19:29:52", "VIX"),
        row("2026-08-21 19:29:52", "QCG", acted="WAIT", blocked="WHIPSAW_LOCK"),
        row("2026-08-21 20:29:52", "VIX"),
        row("2026-08-21 20:29:52", "QCG", acted="WAIT", blocked="WHIPSAW_LOCK"),
        row("2026-08-22 10:51:45", "VIX"),
        row("2026-08-22 10:51:45", "QCG", acted="WAIT", blocked="WHIPSAW_LOCK"),
    ]

    days = signal_rows_by_day(rows)
    assert [group["date"] for group in days] == ["2026-08-22", "2026-08-21"]
    assert len(days[0]["rows"]) == 2
    assert len(days[1]["rows"]) == 2
    assert days[1]["allowed_count"] == 1
    assert days[1]["blocked_count"] == 1
    vix = next(row for row in days[1]["rows"] if row["symbol"] == "VIX")
    qcg = next(row for row in days[1]["rows"] if row["symbol"] == "QCG")
    assert vix["suggestion"] == "CÓ THỂ MUA"
    assert vix["reason"] == "Đủ điều kiện tín hiệu BUY đang bật"
    assert vix["repeat_count"] == 2
    assert qcg["suggestion"] == "KHÔNG MUA"
    assert qcg["reason"] == "EMA nhiễu, khóa mua"


def test_signal_history_does_not_count_observed_sell_as_blocked_buy():
    from viking_v2.dashboard.windows import signal_rows_by_day

    days = signal_rows_by_day([{
        "timestamp": "2026-09-22 14:20:00",
        "symbol": "ANV",
        "signal": "SELL",
        "acted": "WAIT",
        "blocked_by": "NO_NEW_BUY_SIGNAL",
    }])

    assert days[0]["sell_count"] == 1
    assert days[0]["allowed_count"] == 0
    assert days[0]["blocked_count"] == 0


def test_signal_log_keeps_candle_dedupe_across_daemon_restart(tmp_path):
    from viking_v2.storage import SignalLog

    path = tmp_path / "signal_log.csv"
    row = {
        "timestamp": "2026-08-21 10:00:00", "symbol": "HSG", "signal": "BUY",
        "price": 10.5, "ema_fast": 10.2, "ema_slow": 10.1, "rsi": 55,
        "market_state": "ACCUMULATION", "acted": "BUY", "blocked_by": "",
        "candle_key": "2026-08-21",
    }
    assert SignalLog(path).record(row)
    assert not SignalLog(path).record({**row, "timestamp": "2026-08-21 10:05:00"})
    assert SignalLog(path).record({
        **row, "timestamp": "2026-08-24 10:00:00", "candle_key": "2026-08-24",
    })
    assert len(SignalLog(path).read_all()) == 2


def test_main_click_minimizes_popup_without_withdrawing_or_destroying_it():
    from types import SimpleNamespace

    from viking_v2.dashboard.actions import DashboardActionsMixin

    visibility = []

    class FakeTop:
        def __init__(self):
            self.window_state = "normal"
            self.iconify_calls = 0

        def winfo_exists(self):
            return True

        def state(self):
            return self.window_state

        def iconify(self):
            self.iconify_calls += 1
            self.window_state = "iconic"

    popup = SimpleNamespace(
        top=FakeTop(),
        on_visibility_changed=lambda visible: visibility.append(visible),
    )
    dashboard = SimpleNamespace(
        _rule_settings_popup=popup,
        _advanced_popup=None,
        _backtest_popup=None,
        _info_popup=None,
        _history_popup=None,
        _data_popups={},
    )
    event = SimpleNamespace(widget=SimpleNamespace(master=None))

    DashboardActionsMixin._minimize_popups_from_main_click(dashboard, event)
    DashboardActionsMixin._minimize_popups_from_main_click(dashboard, event)

    assert popup.top.window_state == "iconic"
    assert popup.top.iconify_calls == 1
    assert visibility == [False]


def test_history_is_archived_to_monthly_excel_without_manual_export(tmp_path):
    from openpyxl import load_workbook

    from viking_v2.dashboard.windows import HistoryPopup
    from viking_v2.storage import CSVOrderJournal, SignalLog

    signal = SignalLog(tmp_path / "signal_log.csv")
    assert signal.record({
        "timestamp": "2026-08-22 11:00:00", "symbol": "CTS", "signal": "BUY",
        "price": 32.5, "ema_fast": 32.1, "ema_slow": 31.9, "rsi": 55,
        "market_state": "ACCUMULATION", "acted": "BUY", "blocked_by": "",
        "candle_key": "2026-08-22",
    })
    signal_book_path = tmp_path / "excel_archive" / "signal_log_2026-08.xlsx"
    assert signal_book_path.exists()
    signal_book = load_workbook(signal_book_path, read_only=True)
    assert signal_book["TÍN HIỆU"].max_row == 2
    signal_book.close()

    orders = CSVOrderJournal(tmp_path / "order_history.csv")
    orders.append_event({
        "ts": "2026-08-22 11:01:00",
        "intent": {
            "id": "cache-1", "symbol": "CTS", "side": "BUY", "action": "BUY",
            "execution_mode": "PAPER", "order_type": "MARKET", "quantity": 100,
        },
        "result": {"order_id": "order-1", "status": "FILLED"},
    })
    order_book_path = tmp_path / "excel_archive" / "order_history_2026-08.xlsx"
    assert order_book_path.exists()
    order_book = load_workbook(order_book_path, read_only=True)
    assert order_book["LỆNH"].max_row == 2
    order_book.close()

    assert not hasattr(HistoryPopup, "_export_signals")


def test_signal_history_limit_reads_only_latest_rows(tmp_path):
    import csv

    from viking_v2.storage import SignalLog

    path = tmp_path / "signal_log.csv"
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=SignalLog.FIELDS)
        writer.writeheader()
        for index in range(20):
            writer.writerow({
                "timestamp": f"2026-08-{index + 1:02d} 10:00:00",
                "symbol": f"S{index:02d}", "signal": "BUY",
            })

    rows = SignalLog(path).read_all(limit=3)

    assert [row["symbol"] for row in rows] == ["S17", "S18", "S19"]


def test_legacy_signal_csv_is_archived_before_recent_file_is_compacted(tmp_path, monkeypatch):
    import csv

    from openpyxl import load_workbook

    from viking_v2.storage import SignalLog

    monkeypatch.setattr(SignalLog, "RECENT_CSV_ROWS", 3)
    path = tmp_path / "signal_log.csv"
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=SignalLog.FIELDS)
        writer.writeheader()
        for index in range(5):
            writer.writerow({
                "timestamp": f"2026-08-{index + 1:02d} 10:00:00",
                "symbol": f"OLD{index}", "signal": "BUY",
            })

    log = SignalLog(path)
    assert log.record({
        "timestamp": "2026-08-06 10:00:00", "symbol": "NEW", "signal": "SELL",
        "candle_key": "2026-08-06",
    })

    assert [row["symbol"] for row in log.read_all()] == ["OLD3", "OLD4", "NEW"]
    book = load_workbook(
        tmp_path / "excel_archive" / "signal_log_2026-08.xlsx", read_only=True,
    )
    assert book["TÍN HIỆU"].max_row == 7
    book.close()
