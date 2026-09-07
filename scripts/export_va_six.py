"""Combine the verified four-symbol and CTS/VIX batches into one audit workbook."""
from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill


ROOT = Path(__file__).resolve().parents[1]
RUNS = ROOT / "viking_v2/runtime/backtest/runs"
BATCHES = [
    RUNS / "exports/VA-4MA-E-ONLY-20260905-101932/manifest.json",
    RUNS / "exports/VA-CTS-VIX-E-ONLY-20260905-103220/manifest.json",
]
OUTPUT = RUNS / "exports/VA-CTS-VIX-E-ONLY-20260905-103220/TỔNG HỢP 6 MÃ · GỐC VS MUA TỪ 14H.xlsx"


def main() -> None:
    records = []
    for manifest_path in BATCHES:
        payload = json.loads(manifest_path.read_text(encoding="utf-8-sig"))
        records.extend(payload["results"])

    book = Workbook()
    summary = book.active
    summary.title = "SO SÁNH 6 MÃ"
    summary.append([
        "MÃ", "KỊCH BẢN", "MUA", "THẮNG", "THUA", "ĐANG GIỮ",
        "LÃI/LỖ", "% TRÊN 600 TRIỆU", "% TRÊN NAV 1 TỶ", "RUN ID", "FILE CHI TIẾT",
    ])
    details = book.create_sheet("CHI TIẾT TẤT CẢ LỆNH")
    details.append([
        "MÃ", "KỊCH BẢN", "LƯỢT", "VÀO", "RA", "GIÁ VÀO", "GIÁ RA",
        "KHỐI LƯỢNG", "THOÁT BỞI", "KẾT QUẢ", "%", "LÃI/LỖ",
    ])
    for record in records:
        variant = "MUA TỪ 14:00" if record["variant"] in {"14H", "MUA TỪ 14:00"} else "GỐC"
        summary.append([
            record["symbol"], variant, record["buys"], record["wins"], record["losses"],
            record["open_positions"], record["net_pnl"], record["return_600m_pct"] / 100,
            record["return_nav_pct"] / 100, record["run_id"], Path(record["workbook"]).name,
        ])
        run = json.loads((RUNS / f"{record['run_id']}.json").read_text(encoding="utf-8-sig"))
        for index, trade in enumerate(run["trades"], 1):
            entry_value = float(trade.get("entry_value", 0) or 0)
            pct = float(trade.get("net_pnl", 0) or 0) / entry_value if entry_value else 0
            details.append([
                record["symbol"], variant, index, trade["opened_date"],
                trade.get("closed_date") or "ĐANG GIỮ", trade["avg_entry_price"],
                trade.get("avg_exit_price") or None, trade["entry_quantity"],
                " + ".join(trade.get("exit_events") or []) or "ĐANG GIỮ",
                trade["outcome"], pct, trade["net_pnl"],
            ])

    dark = PatternFill("solid", fgColor="17365D")
    for sheet in (summary, details):
        for cell in sheet[1]:
            cell.fill = dark
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(horizontal="center")
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for column in sheet.columns:
            width = max(len(str(cell.value or "")) for cell in column) + 2
            sheet.column_dimensions[column[0].column_letter].width = min(72, max(12, width))
    for row in range(2, summary.max_row + 1):
        summary.cell(row, 7).number_format = '#,##0;[Red]-#,##0'
        summary.cell(row, 8).number_format = '0.00%;[Red]-0.00%'
        summary.cell(row, 9).number_format = '0.00%;[Red]-0.00%'
    for row in range(2, details.max_row + 1):
        details.cell(row, 11).number_format = '0.00%;[Red]-0.00%'
        details.cell(row, 12).number_format = '#,##0;[Red]-#,##0'

    info = book.create_sheet("CẤU HÌNH")
    for row in [
        ("Rule gốc", "EMA3/6 cắt lên + RSI14 tăng; WHIPSAW bật"),
        ("Kịch bản mới", "Chỉ BUY từ 14:00; SELL/SL xử lý ngay"),
        ("Xác nhận X phút", "OFF"),
        ("Mode", "MODE 2 · E ONLY · ACCUMULATION · vốn 1 tỷ · dùng 60% · không compound"),
        ("Nguồn", "CTS: giao dịch 1m, warm-up intraday TradingView; VIX và 4 mã còn lại: 2m"),
        ("Lưu ý", "Các mã là backtest độc lập; VIX cuối kỳ còn vị thế mở."),
    ]:
        info.append(row)
    info.column_dimensions["A"].width = 24
    info.column_dimensions["B"].width = 105
    book.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
