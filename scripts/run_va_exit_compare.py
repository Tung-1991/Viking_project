"""Create six VA workbooks comparing NORMAL AUTO with NORMAL ALERT.

The export keeps the established trade-table layout. The only new trade
metric is MFE; timestamps and settlement delay reuse existing columns.
"""
from __future__ import annotations

from datetime import datetime, timedelta
from pathlib import Path
import sys

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from viking_v2.backtest.data import HistoricalDataStore
from viking_v2.backtest.engine import BacktestEngine, exit_comparison_variants
from viking_v2.backtest.models import BacktestEvent, BacktestResult, BacktestScenario, BacktestTrade
from viking_v2.backtest.replay import ReplayDataStore


BASE = ROOT / "viking_v2/runtime/backtest"
EXPORTS = BASE / "runs/exports"
SPECS = {
    "SHS": ("2026-03-02", "2026-09-03", "HNX"),
    "FTS": ("2026-03-02", "2026-09-03", "HOSE"),
    "SSI": ("2026-03-02", "2026-09-03", "HOSE"),
    "VND": ("2026-03-02", "2026-09-03", "HOSE"),
    "CTS": ("2026-03-14", "2026-08-20", "HOSE"),
    "VIX": ("2026-03-02", "2026-09-03", "HOSE"),
}

HEADER_FILL = PatternFill("solid", fgColor="111820")
WIN_FILL = PatternFill("solid", fgColor="C6EFCE")
LOSS_FILL = PatternFill("solid", fgColor="FFC7CE")


def _stamp(value: str, fallback: str = "—") -> str:
    if value:
        try:
            return datetime.fromisoformat(value).strftime("%d/%m/%Y %H:%M")
        except ValueError:
            pass
    if fallback and fallback != "—":
        try:
            return datetime.fromisoformat(fallback).strftime("%d/%m/%Y")
        except ValueError:
            return fallback
    return "—"


def _events(result: BacktestResult) -> tuple[dict[str, BacktestEvent], dict[str, BacktestEvent]]:
    entries: dict[str, BacktestEvent] = {}
    exits: dict[str, BacktestEvent] = {}
    for event in result.events:
        if event.event == "ENTRY_BUY":
            entries.setdefault(event.trade_id, event)
        elif event.side == "SELL":
            exits[event.trade_id] = event
    return entries, exits


def _exit_name(trade: BacktestTrade) -> str:
    names = {
        "INDICATOR_EXIT": "E",
        "NORMAL_PROTECTION": "NORMAL",
        "STOP_LOSS": "SL",
        "NORMAL_ALERT_EXIT": "NORMAL ALERT",
    }
    return names.get(trade.exit_mode, trade.exit_mode or "CÒN MỞ")


def _settlement_note(event: BacktestEvent | None) -> str:
    if not event or not event.decision_time or not event.fill_time:
        return ""
    try:
        decision = datetime.fromisoformat(event.decision_time)
        fill = datetime.fromisoformat(event.fill_time)
    except ValueError:
        return ""
    try:
        bucket = max(1, int(event.source_resolution or 2))
    except ValueError:
        bucket = 2
    return " · CHỜ T+2" if fill - decision > timedelta(minutes=bucket + 1) else ""


def _exit_detail(trade: BacktestTrade, event: BacktestEvent | None) -> str:
    if not trade.exit_fills:
        return "CÒN MỞ"
    detail = " + ".join(
        f"{_exit_name(trade)} {int(fill.get('quantity', 0)):,}@{float(fill.get('price', 0)):.2f}"
        for fill in trade.exit_fills
    )
    if trade.outcome == "OPEN" and trade.remaining_quantity > 0:
        detail += f" · CÒN GIỮ {trade.remaining_quantity:,}"
    return detail + _settlement_note(event)


def _write_symbol(path: Path, results: list[BacktestResult]) -> None:
    """Write exactly two detailed policy sheets."""
    book = Workbook()
    book.remove(book.active)

    for result in results:
        sheet = book.create_sheet()
        policy = str(result.config.rule_parameters.get("normal_policy", "")).upper()
        sheet.title = f"E + NORMAL {policy}"
        params = result.config.rule_parameters
        fast_in = int(params.get("buy_ema_fast", 3))
        slow_in = int(params.get("buy_ema_slow", 6))
        fast_out = int(params.get("sell_ema_fast", 3))
        slow_out = int(params.get("sell_ema_slow", 6))
        headers = (
            "LƯỢT", "MÃ", "VÀO", "GIÁ VÀO", "KHỐI LƯỢNG", "VỐN", "CẮT LỖ",
            f"EMA{fast_in} / EMA{slow_in} VÀO", "RSI VÀO", "RA", "PHIÊN", "THOÁT BỞI",
            f"EMA{fast_out} / EMA{slow_out} RA", "RSI RA", "MFE %", "PHÍ+THUẾ",
            "LÃI/LỖ", "%", "KẾT QUẢ",
        )
        sheet.append(headers)
        for cell in sheet[1]:
            cell.fill = HEADER_FILL
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center")

        entry_events, exit_events = _events(result)
        for trade in sorted(result.trades, key=lambda row: (row.opened_date, row.cycle_id)):
            entry = entry_events.get(trade.trade_id)
            exit_event = exit_events.get(trade.trade_id)
            sheet.append((
                trade.cycle_id or trade.trade_id,
                trade.symbol,
                _stamp(entry.fill_time if entry else "", trade.opened_date),
                round(trade.avg_entry_price, 2),
                trade.entry_quantity,
                round(trade.entry_value),
                round(trade.stop_price, 2),
                f"{trade.entry_ema_fast:.2f}/{trade.entry_ema_slow:.2f}",
                round(trade.entry_rsi, 1),
                _stamp(exit_event.fill_time if exit_event else "", trade.closed_date),
                trade.sessions_held,
                _exit_detail(trade, exit_event),
                (f"{trade.exit_ema_fast:.2f}/{trade.exit_ema_slow:.2f}"
                 if trade.exit_fills else "—"),
                round(trade.exit_rsi, 1) if trade.exit_fills else "—",
                round(trade.peak_profit_pct, 2),
                round(trade.fees + trade.tax),
                round(trade.net_pnl),
                round(trade.pnl_pct, 2),
                {"WIN": "THẮNG", "LOSS": "LỖ"}.get(trade.outcome, "CÒN MỞ"),
            ))

        for row in range(2, sheet.max_row + 1):
            for column in (4, 7, 9, 14, 15, 18):
                sheet.cell(row, column).number_format = "0.00;[Red]-0.00"
            for column in (5, 6, 16, 17):
                sheet.cell(row, column).number_format = "#,##0;[Red]-#,##0"
            outcome = sheet.cell(row, 19)
            outcome.fill = WIN_FILL if outcome.value == "THẮNG" else LOSS_FILL

        widths = (13, 8, 19, 11, 14, 16, 11, 18, 11, 19, 9, 31, 18, 10, 10, 14, 16, 9, 11)
        for column, width in enumerate(widths, 1):
            sheet.column_dimensions[get_column_letter(column)].width = width
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = f"A1:S{sheet.max_row}"
        sheet.sheet_view.showGridLines = False
    book.save(path)


def main() -> None:
    data = HistoricalDataStore(root=BASE)
    reference = data.load_settings()
    base_params = dict(reference["rule_parameters"])
    base_params.update(
        # Each workbook is an independent one-symbol NAV, matching the VA
        # comparison runs rather than splitting its capital across five slots.
        max_positions=1,
        buy_window_enabled=True,
        buy_window_start="14:00",
        normal_arm_pct=7.0,
        normal_giveback_pct=2.0,
        normal_sell_pct=100.0,
    )
    EXPORTS.mkdir(parents=True, exist_ok=True)
    replay = ReplayDataStore(BASE / "replay")
    engine = BacktestEngine(data, replay)

    for symbol, (start, end, exchange) in SPECS.items():
        scenario = BacktestScenario(
            symbol, [symbol], start, end,
            reference.get("fixed_market_phase", "ACCUMULATION"),
            exposure_pct=float(reference.get("fixed_exposure_pct", 60.0)),
            max_positions=int(base_params.get("max_positions", 1) or 1),
            em_modes=["NORMAL", "IND_EXIT"],
            whipsaw_enabled=bool(reference.get("whipsaw_enabled", True)),
        )
        results: list[BacktestResult] = []
        for variant, params in exit_comparison_variants(scenario, base_params):
            result = engine.run_scenario(
                variant,
                initial_capital=float(reference.get("initial_capital", 1_000_000_000)),
                rule_parameters=params,
                loss_lock_enabled=bool(reference.get("loss_lock_enabled", True)),
                loss_lock_hours=int(reference.get("loss_lock_hours", 24)),
                sell_wait_policy=str(reference.get("sell_wait_policy", "RECHECK")),
                fill_session=str(reference.get("fill_session", "CONTINUOUS")),
                buy_fee_rate=float(reference.get("buy_fee_rate", 0.00045)),
                sell_fee_rate=float(reference.get("sell_fee_rate", 0.00045)),
                sell_tax_rate=float(reference.get("sell_tax_rate", 0.001)),
                simulation_mode="REPLAY",
                execution_resolution="2",
                symbol_exchanges={symbol: exchange},
                save=False,
            )
            result.data_quality["comparison"] = "INDEPENDENT_NORMAL_POLICY"
            results.append(result)
        _write_symbol(EXPORTS / f"{symbol} · E + NORMAL AUTO vs ALERT.xlsx", results)

    print("EXPORTS=" + str(EXPORTS.resolve()), flush=True)


if __name__ == "__main__":
    main()
