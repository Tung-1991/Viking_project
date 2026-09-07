"""Export exactly one two-scenario workbook for each of the six VA symbols."""
from __future__ import annotations

import json
from pathlib import Path
import sys

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from openpyxl import load_workbook

from viking_v2.backtest.models import BacktestResult
from viking_v2.backtest.report import export_run_excel
from viking_v2.storage import AtomicJSONStore


RUNS = ROOT / "viking_v2/runtime/backtest/runs"
OUTPUT = RUNS / "exports/VA-6-MA-RIENG-20260905"
PAIRS = {
    "SHS": ["BT-20260905-101934-2E713A", "BT-20260905-101937-B0B3D0"],
    "FTS": ["BT-20260905-101939-F676A4", "BT-20260905-101941-D8AC7B"],
    "SSI": ["BT-20260905-101943-88C163", "BT-20260905-101946-03B65D"],
    "VND": ["BT-20260905-101948-A8025B", "BT-20260905-101951-7CC16B"],
    "CTS": ["BT-20260905-103223-7CF61D", "BT-20260905-103226-EA71D7"],
    "VIX": ["BT-20260905-103228-4E2697", "BT-20260905-103231-0639C4"],
}


def main() -> None:
    OUTPUT.mkdir(parents=True, exist_ok=True)
    manifest = {"files": []}
    for symbol, ids in PAIRS.items():
        results = [BacktestResult.from_dict(json.loads(
            (RUNS / f"{run_id}.json").read_text(encoding="utf-8-sig")
        )) for run_id in ids]
        for result in results:
            enabled = bool(result.config.rule_parameters.get("buy_window_enabled"))
            # Excel forbids ':' in sheet names, so use the unambiguous 14H label.
            result.config.run_name = "MUA TỪ 14H" if enabled else "GỐC"
        path = export_run_excel(
            results, OUTPUT, mode="MODE 2 - E ONLY", stamp="VA-20260905",
        )
        book = load_workbook(path, read_only=True, data_only=True)
        try:
            assert {"TÍN HIỆU", "KHỚP LỆNH", "THÔNG TIN"} <= set(book.sheetnames)
            scenario_sheets = [name for name in book.sheetnames if name not in {
                "TÍN HIỆU", "KHỚP LỆNH", "THÔNG TIN",
            }]
            assert len(scenario_sheets) == 2
        finally:
            book.close()
        manifest["files"].append({"symbol": symbol, "path": str(path.resolve()), "run_ids": ids})
        print(path.resolve())
    AtomicJSONStore(OUTPUT / "manifest.json").write(manifest)


if __name__ == "__main__":
    main()
