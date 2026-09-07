"""Reproduce VA's four-symbol E-only batch through the normal Mode 2 engine."""
from __future__ import annotations

import argparse
from copy import deepcopy
from datetime import datetime, timedelta
import hashlib
import json
from pathlib import Path
import shutil
import sys

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from viking_v2.backtest.data import HistoricalDataStore, VN_TZ, bar_date
from viking_v2.backtest.engine import BacktestEngine
from viking_v2.backtest.models import BacktestConfig
from viking_v2.backtest.replay import ReplayDataStore
from viking_v2.backtest.report import export_run_excel
from viking_v2.connections.dnse.client import DNSEClient
from viking_v2.storage import AtomicJSONStore

FILES = {
    "SHS": "HNX_DLY_SHS, 2.csv", "FTS": "HOSE_DLY_FTS, 2.csv",
    "SSI": "HOSE_DLY_SSI, 2.csv", "VND": "HOSE_DLY_VND, 2.csv",
}
BASE = ROOT / "viking_v2/runtime/backtest"
REFERENCE = BASE / "runs/BT-20260904-102513-2EFF47.json"


def prepare(start: str, end: str) -> list[dict]:
    replay = ReplayDataStore(BASE / "replay")
    client = DNSEClient()
    data = HistoricalDataStore(client, root=BASE)
    settlement_end = (datetime.fromisoformat(end) + timedelta(days=14)).date().isoformat()
    audits = []
    try:
        for symbol, filename in FILES.items():
            path = Path("C:/Users/tungn/Downloads") / filename
            preview = replay.import_file(path)
            daily = data.load_daily(symbol, start, settlement_end, warmup_sessions=250,
                                    progress=lambda msg: print(msg, flush=True))
            by_day = {bar_date(row).isoformat(): row for row in daily}
            missing, differences, sessions = [], [], []
            for day, row in sorted(by_day.items()):
                if not start <= day <= end:
                    continue
                bars, resolution, quality = replay.load_day(symbol, day)
                if not bars or quality != "FULL" or resolution != "2":
                    missing.append(day)
                    continue
                sessions.append(day)
                differences.append({"day": day, "daily": row["close"], "tv": bars[-1]["close"],
                                    "difference_pct": (bars[-1]["close"] / row["close"] - 1) * 100})
            warmup = sum(day < start for day in by_day)
            audit = {"symbol": symbol, "source": str(path), "sha256": preview.source_hash,
                     "exchange": preview.exchange, "resolution": "2", "source_bars": preview.bar_count,
                     "source_from": preview.coverage_start, "source_to": preview.coverage_end,
                     "partial_dates": preview.partial_dates, "warmup_daily_bars": warmup,
                     "sessions": len(sessions), "missing": missing,
                     "first_session": sessions[0] if sessions else "",
                     "last_session": sessions[-1] if sessions else "",
                     "max_close_difference_pct": max((abs(x["difference_pct"]) for x in differences), default=0),
                     "largest_close_differences": sorted(differences, key=lambda x: abs(x["difference_pct"]), reverse=True)[:5]}
            audits.append(audit)
            print(json.dumps(audit, ensure_ascii=False), flush=True)
            if missing or warmup < 220:
                raise RuntimeError(f"{symbol}: insufficient daily warmup or FULL 2-minute days: {missing}")
    finally:
        client.close()
    return audits


def run_batch(args, audits):
    stamp = datetime.now(VN_TZ).strftime("%Y%m%d-%H%M%S")
    batch = BASE / "runs/exports" / f"VA-4MA-E-ONLY-{stamp}"
    batch.mkdir(parents=True, exist_ok=False)
    source_dir = batch / "inputs"
    source_dir.mkdir()
    for audit in audits:
        original = Path(audit["source"])
        shutil.copy2(original, source_dir / original.name)
        shutil.copy2(BASE / "cache" / f"{audit['symbol']}_1D.json", source_dir / f"{audit['symbol']}_1D.json")
        audit["daily_sha256"] = hashlib.sha256((source_dir / f"{audit['symbol']}_1D.json").read_bytes()).hexdigest()
    base_config = json.loads(REFERENCE.read_text(encoding="utf-8-sig"))["config"]
    scenarios = [False, True] if args.variant == "both" else [args.variant == "afternoon"]
    data = HistoricalDataStore(root=BASE)  # Use the inspected snapshots; no more network requests.
    engine = BacktestEngine(data, ReplayDataStore(BASE / "replay"))
    manifest = {"reference_run": REFERENCE.name, "start": args.start, "end": args.end,
                "variant": args.variant, "inputs": audits, "results": []}
    ledger = AtomicJSONStore(batch / "manifest.json")
    ledger.write(manifest)
    for symbol in FILES:
        for window in scenarios:
            raw = deepcopy(base_config)
            raw.update(symbols=[symbol], start_date=args.start, end_date=args.end,
                       run_name=f"{symbol} · E ONLY · {'14H' if window else 'GOC'}")
            raw["rule_parameters"].update(buy_confirmation_enabled=False, buy_window_enabled=window,
                                           buy_window_start="14:00")
            config = BacktestConfig.from_dict(raw)
            last_bucket = [-1]
            def progress(value, text):
                bucket = int(value * 10)
                if bucket != last_bucket[0]:
                    last_bucket[0] = bucket
                    print(f"{symbol} {'14H' if window else 'GOC'} {value:.0%} {text}", flush=True)
            result = engine.run(config, progress=progress, save=True)
            workbook = export_run_excel(result, batch, mode="MODE 2 - E ONLY", stamp=result.run_id)
            verify_workbook(workbook, result, window)
            record = {"symbol": symbol, "variant": "14H" if window else "GOC", "run_id": result.run_id,
                      "workbook": str(workbook), "buys": result.buy_count, "sells": result.sell_count,
                      "wins": result.win_count, "losses": result.loss_count,
                      "open_positions": sum(trade.outcome == "OPEN" for trade in result.trades),
                      "net_pnl": result.net_pnl, "return_nav_pct": result.return_pct,
                      "return_600m_pct": result.net_pnl / 600_000_000 * 100,
                      "fees": result.total_fees, "tax": result.total_tax,
                      "max_drawdown_nav_pct": result.max_drawdown_pct,
                      "worst_closed_trade_pct": min((trade.pnl_pct for trade in result.trades if trade.outcome != "OPEN"), default=0),
                      "warnings": result.warnings}
            manifest["results"].append(record)
            ledger.write(manifest)
            print(json.dumps(record, ensure_ascii=False), flush=True)
    summary = export_summary(batch, manifest)
    manifest["summary_workbook"] = str(summary)
    ledger.write(manifest)
    print("BATCH=" + str(batch), flush=True)


def export_summary(batch: Path, manifest: dict) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    path = batch / "TỔNG HỢP · 4 MÃ · GỐC VS MUA TỪ 14H.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "SO SÁNH"
    headers = [
        "MÃ", "KỊCH BẢN", "LỆNH MUA", "THẮNG", "THUA", "ĐANG GIỮ",
        "LÃI/LỖ (VNĐ)", "% TRÊN 600 TRIỆU", "% TRÊN NAV 1 TỶ",
        "MAX DRAWDOWN NAV", "LỆNH ĐÓNG TỆ NHẤT", "RUN ID", "FILE CHI TIẾT",
    ]
    sheet.append(headers)
    for record in manifest["results"]:
        sheet.append([
            record["symbol"], "MUA TỪ 14:00" if record["variant"] == "14H" else "GỐC",
            record["buys"], record["wins"], record["losses"], record["open_positions"],
            record["net_pnl"], record["return_600m_pct"] / 100,
            record["return_nav_pct"] / 100, record["max_drawdown_nav_pct"] / 100,
            record["worst_closed_trade_pct"] / 100, record["run_id"],
            Path(record["workbook"]).name,
        ])
    header_fill = PatternFill("solid", fgColor="17365D")
    for cell in sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in range(2, sheet.max_row + 1):
        sheet.cell(row, 7).number_format = '#,##0;[Red]-#,##0'
        for column in range(8, 12):
            sheet.cell(row, column).number_format = '0.00%;[Red]-0.00%'
    widths = [10, 18, 12, 10, 10, 12, 18, 20, 18, 20, 20, 28, 78]
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[chr(64 + index)].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions

    info = book.create_sheet("CẤU HÌNH")
    info_rows = [
        ("Khoảng test", f"{manifest['start']} đến {manifest['end']}"),
        ("Dữ liệu tín hiệu/khớp", "TradingView 2 phút · FULL 128 phiên/mã"),
        ("Chỉ báo", "EMA 3/6 + RSI14"),
        ("Exit", "E ONLY · EXIT SELL"),
        ("Vốn", "1 tỷ NAV · dùng tối đa 60% = 600 triệu · không compound"),
        ("Gốc", "Không giới hạn giờ BUY"),
        ("14H", "Chỉ BUY từ 14:00; SELL/SL không chờ"),
        ("Xác nhận BUY X phút", "OFF trong cả hai kịch bản"),
        ("Ghi chú", "Kịch bản 14H là phép thử, không phải kết luận tối ưu."),
    ]
    for row in info_rows:
        info.append(row)
    info.column_dimensions["A"].width = 28
    info.column_dimensions["B"].width = 95
    book.save(path)
    return path


def verify_workbook(path, result, window):
    from openpyxl import load_workbook
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        assert {"TÍN HIỆU", "KHỚP LỆNH", "THÔNG TIN"} <= set(book.sheetnames)
        assert book["KHỚP LỆNH"].max_row - 1 == len(result.events)
        assert result.config.em_modes == ["IND_EXIT"]
        assert result.data_quality["fallback_count"] == 0
        assert all(value["source_resolution"] == "2" for days in result.data_quality["source_coverage"].values() for value in days.values())
        assert abs(result.final_equity - result.initial_capital - result.net_pnl) < 0.01
        if window:
            for event in result.events:
                if event.side == "BUY":
                    local = datetime.fromisoformat(event.fill_time)
                    assert local.hour >= 14
    finally:
        book.close()


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--start", default="2026-03-02")
    parser.add_argument("--end", default="2026-09-03")
    parser.add_argument("--variant", choices=["baseline", "afternoon", "both"], default="baseline")
    parser.add_argument("--prepare-only", action="store_true")
    args = parser.parse_args()
    audit_rows = prepare(args.start, args.end)
    if not args.prepare_only:
        run_batch(args, audit_rows)
