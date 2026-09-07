"""Paired E-only baseline/14:00 replay for CTS and VIX."""
from __future__ import annotations

from copy import deepcopy
from datetime import datetime
import json
from pathlib import Path
import sys

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from viking_v2.backtest.data import HistoricalDataStore, VN_TZ
from viking_v2.backtest.engine import BacktestEngine
from viking_v2.backtest.models import BacktestConfig
from viking_v2.backtest.replay import ReplayDataStore
from viking_v2.backtest.report import export_run_excel
from viking_v2.storage import AtomicJSONStore


BASE = ROOT / "viking_v2/runtime/backtest"
REFERENCE = BASE / "runs/BT-20260904-102513-2EFF47.json"
SPECS = {
    "CTS": {"start": "2026-03-14", "end": "2026-08-20", "resolution": "1"},
    "VIX": {"start": "2026-03-02", "end": "2026-09-03", "resolution": "2"},
}


def trade_rows(result, variant: str) -> list[list]:
    rows = []
    for trade in result.trades:
        rows.append([
            trade.symbol, variant, trade.trade_id, trade.opened_date,
            trade.closed_date or "ĐANG GIỮ", trade.avg_entry_price,
            trade.avg_exit_price or None, trade.entry_quantity,
            " + ".join(trade.exit_events) or "ĐANG GIỮ", trade.outcome,
            trade.pnl_pct / 100, trade.net_pnl,
        ])
    return rows


def write_summary(batch: Path, records: list[dict], trades: list[list]) -> Path:
    path = batch / "TỔNG HỢP CHI TIẾT · CTS VIX · GỐC VS MUA TỪ 14H.xlsx"
    book = Workbook()
    summary = book.active
    summary.title = "SO SÁNH"
    summary.append([
        "MÃ", "KỊCH BẢN", "MUA", "THẮNG", "THUA", "ĐANG GIỮ",
        "LÃI/LỖ", "% TRÊN 600 TRIỆU", "% TRÊN NAV 1 TỶ", "RUN ID", "FILE CHI TIẾT",
    ])
    for record in records:
        summary.append([
            record["symbol"], record["variant"], record["buys"], record["wins"],
            record["losses"], record["open_positions"], record["net_pnl"],
            record["return_600m_pct"] / 100, record["return_nav_pct"] / 100,
            record["run_id"], Path(record["workbook"]).name,
        ])
    detail = book.create_sheet("CHI TIẾT LỆNH")
    detail.append([
        "MÃ", "KỊCH BẢN", "MÃ LỆNH", "VÀO", "RA", "GIÁ VÀO", "GIÁ RA",
        "KHỐI LƯỢNG", "THOÁT BỞI", "KẾT QUẢ", "%", "LÃI/LỖ",
    ])
    for row in trades:
        detail.append(row)
    fill = PatternFill("solid", fgColor="17365D")
    for sheet in (summary, detail):
        for cell in sheet[1]:
            cell.fill = fill
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(horizontal="center")
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
    for row in range(2, summary.max_row + 1):
        summary.cell(row, 7).number_format = '#,##0;[Red]-#,##0'
        for column in (8, 9):
            summary.cell(row, column).number_format = '0.00%;[Red]-0.00%'
    for row in range(2, detail.max_row + 1):
        detail.cell(row, 11).number_format = '0.00%;[Red]-0.00%'
        detail.cell(row, 12).number_format = '#,##0;[Red]-#,##0'
    for sheet in (summary, detail):
        for column in sheet.columns:
            letter = column[0].column_letter
            sheet.column_dimensions[letter].width = min(70, max(12, max(len(str(cell.value or "")) for cell in column) + 2))
    book.save(path)
    return path


def main() -> None:
    stamp = datetime.now(VN_TZ).strftime("%Y%m%d-%H%M%S")
    batch = BASE / "runs/exports" / f"VA-CTS-VIX-E-ONLY-{stamp}"
    batch.mkdir(parents=True, exist_ok=False)
    raw_reference = json.loads(REFERENCE.read_text(encoding="utf-8-sig"))["config"]
    engine = BacktestEngine(HistoricalDataStore(root=BASE), ReplayDataStore(BASE / "replay"))
    records: list[dict] = []
    trades: list[list] = []
    for symbol, spec in SPECS.items():
        for window in (False, True):
            raw = deepcopy(raw_reference)
            raw.update(
                symbols=[symbol], start_date=spec["start"], end_date=spec["end"],
                execution_resolution=spec["resolution"],
                run_name=f"{symbol} · E ONLY · {'MUA TỪ 14H' if window else 'GỐC'}",
            )
            raw["rule_parameters"].update(
                buy_confirmation_enabled=False,
                buy_window_enabled=window,
                buy_window_start="14:00",
            )
            config = BacktestConfig.from_dict(raw)
            result = engine.run(config, save=True)
            variant = "MUA TỪ 14:00" if window else "GỐC"
            workbook = export_run_excel(result, batch, mode="MODE 2 - E ONLY", stamp=result.run_id)
            book = load_workbook(workbook, read_only=True, data_only=True)
            assert {"TÍN HIỆU", "KHỚP LỆNH", "THÔNG TIN"} <= set(book.sheetnames)
            book.close()
            assert result.data_quality["fallback_count"] == 0
            assert result.data_quality["indicator_warmup_source"][symbol] == "REPLAY_INTRADAY_AGGREGATED"
            if window:
                assert all(
                    datetime.fromisoformat(event.fill_time).astimezone(VN_TZ).hour >= 14
                    for event in result.events if event.side == "BUY"
                )
            record = {
                "symbol": symbol, "variant": variant, "run_id": result.run_id,
                "workbook": str(workbook.resolve()), "buys": result.buy_count,
                "sells": result.sell_count, "wins": result.win_count,
                "losses": result.loss_count,
                "open_positions": sum(trade.outcome == "OPEN" for trade in result.trades),
                "net_pnl": result.net_pnl, "return_nav_pct": result.return_pct,
                "return_600m_pct": result.net_pnl / 600_000_000 * 100,
                "resolution": spec["resolution"],
            }
            records.append(record)
            trades.extend(trade_rows(result, variant))
            print(json.dumps(record, ensure_ascii=False), flush=True)
    summary = write_summary(batch, records, trades)
    manifest = {"created_at": datetime.now(VN_TZ).isoformat(), "results": records,
                "summary_workbook": str(summary.resolve())}
    AtomicJSONStore(batch / "manifest.json").write(manifest)
    print("BATCH=" + str(batch.resolve()), flush=True)


if __name__ == "__main__":
    main()
